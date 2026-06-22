from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.factory.models import Plant
from apps.inventory.models import BulkTransaction, InventoryBulk, InventoryLocation, InventoryRoll, Vendor
from apps.materials.models import InventoryMaterial
from apps.users.models import Role, User


class SmartGrnStockFormTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.role, _ = Role.objects.get_or_create(
            code="STORE",
            defaults={"name": "Store", "default_permissions": ["inventory.view", "inventory.manage"]},
        )
        self.user = User.objects.create_user(username="smart-grn-stock-form", password="pass12345", role=self.role)
        self.client.force_authenticate(self.user)
        self.plant = Plant.objects.create(name="GRN Plant", code="GRN-P")
        self.location = InventoryLocation.objects.create(plant=self.plant, code="GRN-RM", name="RM", type="RM")
        self.vendor = Vendor.objects.create(name="GRN Vendor", code="GRN-V", type="RM", status="ACTIVE")
        self.granule = InventoryMaterial.objects.create(code="GRN-LLDPE", name="LLDPE", category="GRANULE", base_uom="KG")
        self.granule_b = InventoryMaterial.objects.create(code="GRN-HDPE", name="HDPE", category="GRANULE", base_uom="KG")
        self.film = InventoryMaterial.objects.create(code="GRN-BOPP-ROLL", name="BOPP roll", category="FILM_VARIANT", base_uom="KG")

    def test_unified_bulk_grn_posts_multiple_lines_with_same_invoice(self):
        response = self.client.post(
            "/api/inventory/grn/create/",
            {
                "klass": "BULK",
                "source_type": "DIRECT",
                "vendor_id": str(self.vendor.id),
                "vendor_invoice_no": "MULTI-INV-1",
                "warehouse_id": str(self.location.id),
                "lines": [
                    {"material_code": self.granule.code, "qty": "3125", "uom": "KG", "rate_per_uom": "100", "location_id": str(self.location.id)},
                    {"material_code": self.granule_b.code, "qty": "31725", "uom": "KG", "rate_per_uom": "101", "location_id": str(self.location.id)},
                ],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(BulkTransaction.objects.filter(vendor=self.vendor, vendor_invoice_no="MULTI-INV-1", type="INWARD").count(), 2)
        self.assertEqual(InventoryBulk.objects.get(material=self.granule, location=self.location).qty_kg, Decimal("3125.0000"))
        self.assertEqual(InventoryBulk.objects.get(material=self.granule_b, location=self.location).qty_kg, Decimal("31725.0000"))

    def test_unified_bulk_grn_rolls_back_all_lines_when_later_line_fails(self):
        response = self.client.post(
            "/api/inventory/grn/create/",
            {
                "klass": "BULK",
                "source_type": "DIRECT",
                "vendor_id": str(self.vendor.id),
                "vendor_invoice_no": "MULTI-ROLLBACK-1",
                "warehouse_id": str(self.location.id),
                "lines": [
                    {"material_code": self.granule.code, "qty": "3125", "uom": "KG", "rate_per_uom": "100", "location_id": str(self.location.id)},
                    {"material_code": "NO-SUCH-MATERIAL", "qty": "31725", "uom": "KG", "rate_per_uom": "101", "location_id": str(self.location.id)},
                ],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400, response.data)
        self.assertFalse(BulkTransaction.objects.filter(vendor=self.vendor, vendor_invoice_no="MULTI-ROLLBACK-1", type="INWARD").exists())
        self.assertFalse(InventoryBulk.objects.filter(material=self.granule, location=self.location).exists())

    def test_unified_grn_roll_carries_stock_form_and_width_basis(self):
        response = self.client.post(
            "/api/inventory/grn/create/",
            {
                "klass": "ROLL",
                "source_type": "DIRECT",
                "vendor_id": str(self.vendor.id),
                "vendor_invoice_no": "ROLL-TUBE-INV-1",
                "warehouse_id": str(self.location.id),
                "lines": [
                    {
                        "material_code": self.film.code,
                        "qty": "100",
                        "net_weight_kg": "100",
                        "width_mm": "300",
                        "thickness_um": "50",
                        "stock_form": "LAYFLAT_TUBE",
                        "width_basis": "LAYFLAT_WIDTH",
                        "location_id": str(self.location.id),
                    }
                ],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)
        roll = InventoryRoll.objects.get(vendor_invoice_no="ROLL-TUBE-INV-1")
        self.assertEqual(roll.stock_form, "LAYFLAT_TUBE")
        self.assertEqual(roll.width_basis, "LAYFLAT_WIDTH")
        self.assertEqual(roll.width_mm, Decimal("300.00"))

    def test_roll_upload_template_omits_supplier_roll_and_upload_auto_numbers(self):
        template_response = self.client.get("/api/inventory/grn/roll-upload-template/")
        self.assertEqual(template_response.status_code, 200, template_response.content[:200])

        workbook = load_workbook(BytesIO(template_response.content))
        roll_lines = workbook["Roll Lines"]
        headers = [cell.value for cell in roll_lines[1]]
        self.assertNotIn("Supplier Roll No", headers)
        self.assertNotIn("ERP Roll Label", headers)
        self.assertEqual(headers[:4], ["Line No", "Material Type", "Material Code", "Material Name"])

        header = workbook["GRN Header"]
        header["B2"] = self.vendor.code
        header["B3"] = self.plant.code
        header["B4"] = self.location.code
        header["B5"] = "AUTO-ROLL-UPLOAD-1"

        header_index = {name: idx + 1 for idx, name in enumerate(headers)}
        roll_lines.cell(row=2, column=header_index["Material Code"]).value = self.film.code
        roll_lines.cell(row=2, column=header_index["Gross Wt KG"]).value = 12
        roll_lines.cell(row=2, column=header_index["Tare/Core KG"]).value = 1
        roll_lines.cell(row=2, column=header_index["Net Wt KG"]).value = 11
        roll_lines.cell(row=2, column=header_index["Width MM"]).value = 570
        roll_lines.cell(row=2, column=header_index["Thickness Micron"]).value = 50
        roll_lines.cell(row=2, column=header_index["Location Code"]).value = self.location.code

        payload = BytesIO()
        workbook.save(payload)
        upload = SimpleUploadedFile(
            "roll-upload.xlsx",
            payload.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post("/api/inventory/grn/upload-rolls/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        roll = InventoryRoll.objects.get(material=self.film, weight_kg=Decimal("11.0000"))
        self.assertTrue(roll.label_id.startswith("ROLL-"))
        self.assertTrue(roll.batch_no.startswith("AUTO-"))
        self.assertEqual(roll.meta_json.get("vendor_roll_label"), "")
