from __future__ import annotations

import csv
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO, TextIOWrapper
from typing import Any, Dict, Iterable, List, Optional

from django.core.exceptions import ValidationError
from django.db import models, transaction
from django.db.models import Sum
from django.utils import timezone

from apps.factory.models import Plant
from apps.inventory.models import (
    BulkTransaction,
    DeliveryChallan,
    InventoryAlert,
    InventoryAuditBatch,
    InventoryAuditLine,
    InventoryBulk,
    InventoryFinancialPeriod,
    InventoryLocation,
    InventoryRoll,
    JobWorkOrder,
    PackagingStock,
    PackagingTransaction,
    RollMovement,
)
from apps.inventory.services.bulk_service import BulkService
from apps.inventory.services.packaging_service import PackagingService
from apps.inventory.services.roll_service import RollService
from apps.inventory.services.wac import apply_wac, display_value, q4, signed_value
from apps.materials.models import GranuleQualityCode, InventoryMaterial
from apps.users.models import PermissionAuditLog

try:
    from apps.costing.models import MaterialCostSnapshot
except Exception:  # pragma: no cover - costing app may be unavailable in isolated test settings
    MaterialCostSnapshot = None

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except Exception:  # pragma: no cover
    Workbook = None
    load_workbook = None
    Alignment = None
    Font = None
    PatternFill = None


def _dec(value: Any, fallback: str = "0") -> Decimal:
    if value in (None, ""):
        return Decimal(fallback)
    return Decimal(str(value))


def current_indian_financial_year(today: Optional[date] = None) -> str:
    today = today or timezone.localdate()
    start_year = today.year if today.month >= 4 else today.year - 1
    return f"{start_year}-{start_year + 1}"


def financial_year_dates(financial_year: str) -> tuple[date, date]:
    try:
        start_year = int(str(financial_year).split("-", 1)[0])
    except Exception as exc:
        raise ValidationError("Financial year must be like 2026-2027.") from exc
    return date(start_year, 4, 1), date(start_year + 1, 3, 31)


def _latest_material_rate(material: Optional[InventoryMaterial]) -> tuple[Decimal, str]:
    if not material:
        return Decimal("0.0000"), "ZERO"
    if MaterialCostSnapshot is not None:
        snapshot = MaterialCostSnapshot.objects.filter(material=material).order_by("-effective_date").first()
        if snapshot and _dec(snapshot.avg_rate_per_kg) > 0:
            return q4(snapshot.avg_rate_per_kg), "MATERIAL_COST_SNAPSHOT"
    for attr in ("standard_cost", "avg_cost", "cost_price"):
        value = getattr(material, attr, None)
        if value not in (None, "") and _dec(value) > 0:
            return q4(value), "MATERIAL_STANDARD"
    return Decimal("0.0000"), "ZERO"


def _roll_rate(roll: InventoryRoll) -> tuple[Decimal, str]:
    meta = roll.meta_json or {}
    for key in ("unit_cost_per_kg", "unit_cost", "rate_per_kg", "rate_per_uom"):
        value = meta.get(key)
        if value not in (None, "") and _dec(value) > 0:
            return q4(value), "ROLL_GRN_META"
    return _latest_material_rate(roll.material)


class InventoryAuditService:
    AUDIT_REFERENCE_PREFIXES = ("OPENING_STOCK:", "PHYSICAL_COUNT:", "FY_CLOSE:", "FY_CORRECTION:")
    MOVEMENT_TYPES = {"INWARD", "PRODUCE", "CONSUME", "TRANSFER", "ADJUST", "COUNT_SHORT", "COUNT_EXCESS", "FY_CORRECTION"}
    CUTOVER_OPENING_MODES = {"CUTOVER_OPENING", "ADDITIVE_CUTOVER", "JUNE_CUTOVER"}

    @classmethod
    def _batch_workflow(cls, batch: InventoryAuditBatch) -> Dict[str, Any]:
        summary = batch.summary_json if isinstance(batch.summary_json, dict) else {}
        workflow = summary.get("workflow") if isinstance(summary.get("workflow"), dict) else {}
        return workflow

    @classmethod
    def _is_cutover_opening(cls, batch: InventoryAuditBatch) -> bool:
        if batch.type != "OPENING_STOCK":
            return False
        workflow = cls._batch_workflow(batch)
        mode = str(workflow.get("mode") or workflow.get("opening_mode") or "").upper()
        return mode in cls.CUTOVER_OPENING_MODES

    @classmethod
    def ensure_default_period(cls) -> InventoryFinancialPeriod:
        fy = current_indian_financial_year()
        start_date, end_date = financial_year_dates(fy)
        period, _ = InventoryFinancialPeriod.objects.get_or_create(
            financial_year=fy,
            defaults={"start_date": start_date, "end_date": end_date, "status": "OPEN"},
        )
        return period

    @classmethod
    def start_period(cls, *, financial_year: str, user=None) -> InventoryFinancialPeriod:
        start_date, end_date = financial_year_dates(financial_year)
        open_period = InventoryFinancialPeriod.objects.filter(status="OPEN").exclude(financial_year=financial_year).first()
        if open_period:
            raise ValidationError(f"Close {open_period.financial_year} before starting another open period.")
        period, created = InventoryFinancialPeriod.objects.get_or_create(
            financial_year=financial_year,
            defaults={"start_date": start_date, "end_date": end_date, "status": "OPEN"},
        )
        if not created and period.status == "CLOSED":
            raise ValidationError("Cannot reopen a closed financial year from this screen.")
        if period.status != "OPEN":
            period.status = "OPEN"
            period.save(update_fields=["status", "updated_at"])
        return period

    @classmethod
    def assert_period_allows_cutoff(cls, financial_year: str, cutoff_at) -> None:
        period = InventoryFinancialPeriod.objects.filter(financial_year=financial_year).first()
        if period and period.status == "CLOSED":
            raise ValidationError(
                f"Financial year {financial_year} is closed. Use an FY correction batch with approval trail."
            )
        if period and cutoff_at:
            cutoff_date = cutoff_at.date() if hasattr(cutoff_at, "date") else cutoff_at
            if cutoff_date < period.start_date or cutoff_date > period.end_date:
                raise ValidationError("Cutoff date must fall inside the selected financial year.")

    @classmethod
    def require_closed_period_for_correction(cls, financial_year: str) -> InventoryFinancialPeriod:
        period = InventoryFinancialPeriod.objects.filter(financial_year=financial_year).first()
        if not period or period.status != "CLOSED":
            raise ValidationError("FY correction can only be created for a closed financial year. Use Stock Count for an open year.")
        return period

    @classmethod
    def create_batch(cls, *, payload: Dict[str, Any], user=None) -> InventoryAuditBatch:
        batch_type = str(payload.get("type") or "").upper()
        if batch_type not in dict(InventoryAuditBatch.TYPE_CHOICES):
            raise ValidationError("Invalid audit batch type.")
        financial_year = str(payload.get("financial_year") or current_indian_financial_year()).strip()
        plant_id = payload.get("plant")
        if not plant_id:
            raise ValidationError("Plant is required.")
        cutoff_value = payload.get("cutoff_at")
        cutoff_at = cls._parse_cutoff(cutoff_value) if cutoff_value else timezone.now()
        if batch_type != "FY_CORRECTION":
            cls.assert_period_allows_cutoff(financial_year, cutoff_at)
        else:
            period = cls.require_closed_period_for_correction(financial_year)
            if not str(payload.get("notes") or "").strip():
                raise ValidationError("FY correction requires a reason or reference note.")
            if not cutoff_value:
                cutoff_at = timezone.make_aware(datetime.combine(period.end_date, time.max))
            else:
                cutoff_date = cutoff_at.date() if hasattr(cutoff_at, "date") else cutoff_at
                if cutoff_date < period.start_date or cutoff_date > period.end_date:
                    raise ValidationError("FY correction cutoff must fall inside the closed financial year.")
        summary_json = payload.get("summary_json") if isinstance(payload.get("summary_json"), dict) else {}
        workflow = payload.get("_v36_workflow") if isinstance(payload.get("_v36_workflow"), dict) else None
        if workflow:
            summary_json = dict(summary_json or {})
            summary_json["workflow"] = {**dict(summary_json.get("workflow") or {}), **workflow}
        batch = InventoryAuditBatch.objects.create(
            type=batch_type,
            plant_id=plant_id,
            financial_year=financial_year,
            cutoff_at=cutoff_at,
            notes=payload.get("notes") or "",
            source_file_name=payload.get("source_file_name") or "",
            summary_json=summary_json,
            created_by=user if getattr(user, "is_authenticated", False) else None,
        )
        cls._emit_batch_event(batch=batch, user=user, action="INVENTORY_BATCH_CREATED", before=None, after={"status": batch.status})
        return batch

    @classmethod
    def _parse_cutoff(cls, value):
        if isinstance(value, datetime):
            return value if timezone.is_aware(value) else timezone.make_aware(value)
        if isinstance(value, date):
            return timezone.make_aware(datetime.combine(value, time.max))
        raw = str(value or "").strip()
        if not raw:
            return timezone.now()
        parsed = None
        try:
            parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        except ValueError:
            try:
                parsed = datetime.combine(date.fromisoformat(raw), time.max)
            except ValueError as exc:
                raise ValidationError("cutoff_at must be an ISO date or datetime.") from exc
        return parsed if timezone.is_aware(parsed) else timezone.make_aware(parsed)

    @classmethod
    @transaction.atomic
    def import_lines(cls, *, batch: InventoryAuditBatch, rows: Iterable[Dict[str, Any]]) -> List[InventoryAuditLine]:
        if batch.status != "DRAFT":
            raise ValidationError("Only draft batches can be edited.")
        created = []
        for row in rows:
            created.append(cls.upsert_line(batch=batch, payload=row))
        cls.refresh_batch_summary(batch)
        return created

    @classmethod
    @transaction.atomic
    def replace_lines(cls, *, batch: InventoryAuditBatch, rows: Iterable[Dict[str, Any]]) -> List[InventoryAuditLine]:
        if batch.status != "DRAFT":
            raise ValidationError("Only draft batches can be edited.")
        batch.lines.all().delete()
        return cls.import_lines(batch=batch, rows=rows)

    @classmethod
    def import_file(cls, *, batch: InventoryAuditBatch, uploaded_file, default_stock_class: Optional[str] = None) -> List[InventoryAuditLine]:
        rows = cls._parse_upload_rows(uploaded_file, default_stock_class=default_stock_class)
        return cls.import_lines(batch=batch, rows=rows)

    @classmethod
    def upsert_line(cls, *, batch: InventoryAuditBatch, payload: Dict[str, Any]) -> InventoryAuditLine:
        material = InventoryMaterial.objects.get(id=payload.get("material"))
        location = InventoryLocation.objects.select_related("plant").get(id=payload.get("location"))
        if str(location.plant_id) != str(batch.plant_id):
            raise ValidationError("Line location must belong to the batch plant.")

        line = InventoryAuditLine(
            batch=batch,
            stock_class=str(payload.get("stock_class") or "").upper(),
            material=material,
            granule_code_id=payload.get("granule_code") or payload.get("granule_code_id") or None,
            grade_id=payload.get("grade") or payload.get("grade_id") or None,
            plant=batch.plant,
            location=location,
            uom=payload.get("uom") or getattr(material, "base_uom", "KG") or "KG",
            system_qty=_dec(payload.get("system_qty")),
            counted_qty=_dec(payload.get("counted_qty")) if payload.get("counted_qty") not in (None, "") else None,
            opening_qty=_dec(payload.get("opening_qty") if payload.get("opening_qty") not in (None, "") else payload.get("quantity")),
            rate=_dec(payload.get("rate")) if payload.get("rate") not in (None, "") else None,
            label_id=str(payload.get("label_id") or "").strip(),
            batch_no=str(payload.get("batch_no") or "").strip(),
            width_mm=_dec(payload.get("width_mm")) if payload.get("width_mm") not in (None, "") else None,
            thickness_micron=_dec(payload.get("thickness_micron")) if payload.get("thickness_micron") not in (None, "") else None,
            length_m=_dec(payload.get("length_m")) if payload.get("length_m") not in (None, "") else None,
            is_fg=bool(payload.get("is_fg", False)),
            stage_index=int(payload.get("stage_index") or 0),
            status=str(payload.get("status") or "AVAILABLE").upper(),
            packaging_kind=str(payload.get("packaging_kind") or getattr(material, "packaging_kind", "") or ""),
            base_uom=str(payload.get("base_uom") or getattr(material, "base_uom", "") or ""),
        )
        refs: Dict[str, Any] = {}
        for key in ("stock_form", "width_basis", "count_reason_code", "count_reason_note", "counted_at", "entry_at"):
            if payload.get(key) not in (None, ""):
                refs[key] = payload.get(key)
        if refs:
            line.posted_reference_json = refs
        if batch.type in {"PHYSICAL_COUNT", "FY_CORRECTION"}:
            line.system_qty = cls.resolve_system_qty(line)
            if line.counted_qty is None:
                line.counted_qty = line.system_qty
        line.row_errors = cls.validate_line(line)
        cls._apply_resolved_rate(line)
        line.save()
        return line

    @classmethod
    def validate_line(cls, line: InventoryAuditLine) -> List[str]:
        errors: List[str] = []
        stock_class = str(line.stock_class or "").upper()
        category = str(getattr(line.material, "category", "") or "").upper()
        if stock_class not in {"BULK", "ROLL", "PACKAGING"}:
            errors.append("Stock class must be BULK, ROLL, or PACKAGING.")
        if not line.material_id:
            errors.append("Material is required.")
        if not line.location_id:
            errors.append("Location is required.")
        if line.material_id and str(getattr(line.material, "status", "ACTIVE") or "").upper() != "ACTIVE":
            errors.append(f"E-OS-07: material {getattr(line.material, 'code', '')} is inactive.")
        if line.material_id and line.uom and getattr(line.material, "base_uom", None):
            if str(line.uom or "").upper() != str(line.material.base_uom or "").upper():
                errors.append(f"E-OS-01: UOM '{line.uom}' does not match {line.material.code} base UOM '{line.material.base_uom}'.")
        if line.rate is not None and _dec(line.rate) <= 0:
            errors.append("E-OS-06: rate must be greater than zero, or leave it blank for the system fallback rate.")
        if line.granule_code_id:
            if category != "GRANULE":
                errors.append("Granule code can only be used with granule materials.")
            elif str(line.granule_code.granule_id) != str(line.material_id):
                errors.append("Granule code does not belong to this granule.")
        if stock_class == "PACKAGING" and category != "PACKAGING":
            errors.append("Packaging audit lines require a PACKAGING material.")
        if stock_class == "ROLL":
            if category not in {"FILM_VARIANT", "POD"}:
                errors.append("Roll audit lines require a FILM_VARIANT or POD material.")
            if _dec(line.opening_qty) <= 0 and line.batch.type == "OPENING_STOCK":
                errors.append("Roll opening weight is required.")
            if _dec(line.width_mm) <= 0:
                errors.append("Roll width is required.")
            if _dec(line.thickness_micron) <= 0:
                errors.append("Roll thickness is required.")
            if getattr(line.material, "is_extrudable", False) and not line.grade_id:
                errors.append("Grade is required for extrudable roll stock.")
        if stock_class in {"BULK", "PACKAGING"} and line.batch.type == "OPENING_STOCK" and _dec(line.opening_qty) <= 0:
            errors.append("Opening quantity is required.")
        if line.batch.type == "OPENING_STOCK" and not errors:
            if cls._is_cutover_opening(line.batch):
                duplicate = cls._opening_duplicate_roll_label(line)
                if duplicate:
                    errors.append(duplicate)
            else:
                prior = cls._opening_prior_movement(line)
                if prior:
                    errors.append(prior)
        if line.batch.type in {"PHYSICAL_COUNT", "FY_CORRECTION"} and line.counted_qty is None:
            errors.append("Counted quantity is required.")
        if line.batch.type == "FY_CORRECTION" and not str(line.batch.notes or "").strip():
            errors.append("FY correction requires a written reason in sheet notes.")
        if not errors and line.stock_class in {"BULK", "PACKAGING", "ROLL"}:
            resolved_rate, source = cls.resolve_line_rate(line)
            if resolved_rate <= 0:
                errors.append(f"W-RATE-01: no rate supplied and no fallback rate found; value will post at 0 ({source}).")
        return errors

    @classmethod
    def _blocking_errors(cls, row_errors: Iterable[str]) -> List[str]:
        return [str(error) for error in (row_errors or []) if not str(error).startswith("W-")]

    @classmethod
    def _opening_duplicate_roll_label(cls, line: InventoryAuditLine) -> str:
        if line.stock_class == "ROLL" and line.label_id and InventoryRoll.objects.filter(label_id=line.label_id).exclude(meta_json__inventory_audit_batch=str(line.batch_id)).exists():
            return f"E-OS-02: roll label {line.label_id} already exists."
        return ""

    @classmethod
    def resolve_line_rate(cls, line: InventoryAuditLine) -> tuple[Decimal, str]:
        if line.rate is not None and _dec(line.rate) > 0:
            return q4(line.rate), "USER"

        stock_class = str(line.stock_class or "").upper()
        if stock_class == "BULK":
            stock_qs = InventoryBulk.objects.filter(material=line.material, plant=line.plant, location=line.location)
            if line.granule_code_id:
                stock_qs = stock_qs.filter(granule_code=line.granule_code)
            stock = stock_qs.order_by("-updated_at").first()
            if stock and _dec(stock.avg_cost) > 0:
                return q4(stock.avg_cost), "MATERIAL_AVG"

            start_date, end_date = financial_year_dates(line.batch.financial_year)
            tx_qs = BulkTransaction.objects.filter(
                material=line.material,
                location=line.location,
                type="INWARD",
                avg_cost__gt=0,
                created_at__date__gte=start_date,
                created_at__date__lte=end_date,
            )
            if line.granule_code_id:
                tx_qs = tx_qs.filter(granule_code=line.granule_code)
            tx = tx_qs.order_by("-created_at").first()
            if tx:
                return q4(tx.avg_cost), "LAST_GRN"

        elif stock_class == "PACKAGING":
            stock = PackagingStock.objects.filter(material=line.material, plant=line.plant, location=line.location).order_by("-updated_at").first()
            if stock and _dec(stock.avg_cost) > 0:
                return q4(stock.avg_cost), "MATERIAL_AVG"
            start_date, end_date = financial_year_dates(line.batch.financial_year)
            tx = PackagingTransaction.objects.filter(
                material=line.material,
                location=line.location,
                type="INWARD",
                avg_cost__gt=0,
                created_at__date__gte=start_date,
                created_at__date__lte=end_date,
            ).order_by("-created_at").first()
            if tx:
                return q4(tx.avg_cost), "LAST_GRN"

        return _latest_material_rate(line.material)

    @classmethod
    def _apply_resolved_rate(cls, line: InventoryAuditLine) -> tuple[Decimal, str]:
        rate, source = cls.resolve_line_rate(line)
        if line.rate is None and rate > 0:
            line.rate = rate
        refs = dict(line.posted_reference_json or {})
        refs["rate_source"] = source
        refs["resolved_rate"] = float(rate)
        line.posted_reference_json = refs
        return rate, source

    @classmethod
    def _opening_prior_movement(cls, line: InventoryAuditLine) -> str:
        start_date, end_date = financial_year_dates(line.batch.financial_year)
        reference_prefix = f"OPENING_STOCK:{line.batch.financial_year}:"
        if line.stock_class == "BULK":
            qs = BulkTransaction.objects.filter(
                material=line.material,
                location=line.location,
                created_at__date__gte=start_date,
                created_at__date__lte=end_date,
            )
            if line.granule_code_id:
                qs = qs.filter(granule_code=line.granule_code)
            qs = qs.exclude(type="OPENING_BALANCE").exclude(reference__startswith=reference_prefix)
            count = qs.count()
            if count:
                return f"E-OS-05: opening blocked because {count} bulk movement(s) already exist for this material/location in {line.batch.financial_year}."
        elif line.stock_class == "PACKAGING":
            count = PackagingTransaction.objects.filter(
                material=line.material,
                location=line.location,
                created_at__date__gte=start_date,
                created_at__date__lte=end_date,
            ).exclude(type="OPENING_BALANCE").exclude(reference__startswith=reference_prefix).count()
            if count:
                return f"E-OS-05: opening blocked because {count} packaging movement(s) already exist for this material/location in {line.batch.financial_year}."
        elif line.stock_class == "ROLL":
            duplicate = cls._opening_duplicate_roll_label(line)
            if duplicate:
                return duplicate
            count = RollMovement.objects.filter(
                roll__material=line.material,
                to_location=line.location,
                timestamp__date__gte=start_date,
                timestamp__date__lte=end_date,
            ).exclude(reason="OPENING_BALANCE").exclude(reason_note__startswith=reference_prefix).count()
            if count:
                return f"E-OS-05: opening blocked because {count} roll movement(s) already exist for this material/location in {line.batch.financial_year}."
        return ""

    @classmethod
    def _parse_upload_rows(cls, uploaded_file, default_stock_class: Optional[str] = None) -> List[Dict[str, Any]]:
        name = str(getattr(uploaded_file, "name", "") or "").lower()
        if name.endswith(".xlsx"):
            return cls._parse_xlsx_rows(uploaded_file, default_stock_class=default_stock_class)
        return cls._parse_csv_rows(uploaded_file, default_stock_class=default_stock_class)

    @classmethod
    def _parse_csv_rows(cls, uploaded_file, default_stock_class: Optional[str] = None) -> List[Dict[str, Any]]:
        uploaded_file.seek(0)
        wrapper = TextIOWrapper(uploaded_file.file, encoding="utf-8-sig")
        try:
            reader = csv.DictReader(wrapper)
            rows = []
            for row in reader:
                normalized = cls._normalize_import_row(row, default_stock_class=default_stock_class)
                if normalized:
                    rows.append(normalized)
            return rows
        finally:
            wrapper.detach()

    @classmethod
    def _parse_xlsx_rows(cls, uploaded_file, default_stock_class: Optional[str] = None) -> List[Dict[str, Any]]:
        if load_workbook is None:
            raise ValidationError("Excel import is unavailable because openpyxl is not installed.")
        uploaded_file.seek(0)
        workbook = load_workbook(uploaded_file, data_only=True)
        worksheet = workbook.active
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            return []
        headers = [str(value or "").strip() for value in header_row]
        rows = []
        for raw_row in worksheet.iter_rows(min_row=2, values_only=True):
            payload = {headers[index]: raw_row[index] for index in range(len(headers))}
            normalized = cls._normalize_import_row(payload, default_stock_class=default_stock_class)
            if normalized:
                rows.append(normalized)
        return rows

    @classmethod
    def _normalize_import_row(cls, row: Dict[str, Any], default_stock_class: Optional[str] = None) -> Optional[Dict[str, Any]]:
        payload = {str(key or "").strip().lower(): value for key, value in (row or {}).items() if str(key or "").strip()}
        if not any(str(value or "").strip() for value in payload.values()):
            return None
        stock_class = str(payload.get("stock_class") or default_stock_class or "").upper().strip()
        normalized = {
            "stock_class": stock_class,
            "material": cls._clean_cell(payload.get("material")),
            "granule_code": cls._clean_cell(payload.get("granule_code") or payload.get("granule_code_id")),
            "grade": cls._clean_cell(payload.get("grade") or payload.get("grade_id")),
            "location": cls._clean_cell(payload.get("location")),
            "uom": cls._clean_cell(payload.get("uom")),
            "quantity": cls._clean_cell(payload.get("quantity") or payload.get("opening_qty")),
            "counted_qty": cls._clean_cell(payload.get("counted_qty")),
            "rate": cls._clean_cell(payload.get("rate")),
            "label_id": cls._clean_cell(payload.get("label_id")),
            "batch_no": cls._clean_cell(payload.get("batch_no")),
            "width_mm": cls._clean_cell(payload.get("width_mm")),
            "thickness_micron": cls._clean_cell(payload.get("thickness_micron")),
            "length_m": cls._clean_cell(payload.get("length_m")),
            "is_fg": cls._to_bool(payload.get("is_fg")),
            "stage_index": cls._clean_cell(payload.get("stage_index")),
            "status": cls._clean_cell(payload.get("status")),
            "packaging_kind": cls._clean_cell(payload.get("packaging_kind")),
            "base_uom": cls._clean_cell(payload.get("base_uom")),
            "stock_form": cls._clean_cell(payload.get("stock_form")),
            "width_basis": cls._clean_cell(payload.get("width_basis")),
        }
        return {key: value for key, value in normalized.items() if value not in (None, "")}

    @staticmethod
    def _clean_cell(value):
        if value in (None, ""):
            return None
        if isinstance(value, Decimal):
            return str(value)
        if isinstance(value, (int, float)):
            if int(value) == value:
                return str(int(value))
            return str(value)
        return str(value).strip()

    @staticmethod
    def _to_bool(value) -> bool:
        if isinstance(value, bool):
            return value
        token = str(value or "").strip().lower()
        return token in {"1", "true", "yes", "y", "fg"}

    @classmethod
    def validate_batch(cls, *, batch: InventoryAuditBatch) -> Dict[str, Any]:
        errors = []
        for line in batch.lines.select_related("material", "location", "granule_code"):
            if batch.type in {"PHYSICAL_COUNT", "FY_CORRECTION"}:
                line.system_qty = cls.resolve_system_qty(line)
            line.row_errors = cls.validate_line(line)
            cls._apply_resolved_rate(line)
            line.save(update_fields=["row_errors", "system_qty", "rate", "variance_qty", "value", "posted_reference_json", "updated_at"])
            blockers = cls._blocking_errors(line.row_errors)
            if blockers:
                errors.append({"line_id": str(line.id), "errors": blockers})
        cls.refresh_batch_summary(batch)
        return {"ok": not errors, "errors": errors, "summary": batch.summary_json}

    @classmethod
    def preview_batch(cls, *, batch: InventoryAuditBatch) -> Dict[str, Any]:
        validation = cls.validate_batch(batch=batch)
        summary = batch.summary_json or {}
        line_rows = []
        transaction_count = 0
        for line in batch.lines.select_related("material", "location").order_by("created_at"):
            if cls._blocking_errors(line.row_errors):
                continue
            if batch.type == "OPENING_STOCK":
                delta_qty = _dec(line.opening_qty) if cls._is_cutover_opening(batch) else _dec(line.opening_qty) - cls.resolve_system_qty(line)
                effect_qty = _dec(line.opening_qty)
            elif batch.type in {"PHYSICAL_COUNT", "FY_CORRECTION"}:
                delta_qty = _dec(line.counted_qty) - _dec(line.system_qty)
                effect_qty = delta_qty
            else:
                delta_qty = _dec(line.counted_qty or line.system_qty or 0)
                effect_qty = delta_qty
            if effect_qty != 0:
                transaction_count += 1
            line_rows.append(
                {
                    "line_id": str(line.id),
                    "stock_class": line.stock_class,
                    "material_code": getattr(line.material, "code", ""),
                    "material_name": getattr(line.material, "name", ""),
                    "location_name": getattr(line.location, "name", ""),
                    "system_qty": float(line.system_qty or 0),
                    "entered_qty": float(line.opening_qty if batch.type == "OPENING_STOCK" else (line.counted_qty or 0)),
                    "delta_qty": float(delta_qty),
                    "value": float(line.value or 0),
                    "errors": line.row_errors or [],
                }
            )
        return {
            "batch": {
                "id": str(batch.id),
                "batch_no": batch.batch_no,
                "type": batch.type,
                "status": batch.status,
                "financial_year": batch.financial_year,
                "workflow": cls._batch_workflow(batch),
            },
            "ok": validation["ok"],
            "summary": summary,
            "transaction_count": transaction_count,
            "blockers": validation["errors"],
            "rows": line_rows[:500],
            "impact": {
                "bulk_kg": summary.get("bulk_kg", 0),
                "roll_kg": summary.get("roll_kg", 0),
                "packaging_qty": summary.get("packaging_qty", 0),
                "value": summary.get("value", 0),
            },
            "dual_fy_impact": cls._preview_dual_fy_impact(batch=batch) if batch.type == "FY_CORRECTION" else None,
        }

    @classmethod
    def _next_financial_year(cls, financial_year: str) -> str:
        start, _ = financial_year_dates(financial_year)
        return f"{start.year + 1}-{start.year + 2}"

    @classmethod
    def _preview_dual_fy_impact(cls, *, batch: InventoryAuditBatch) -> Dict[str, Any]:
        next_fy = cls._next_financial_year(batch.financial_year)
        period = InventoryFinancialPeriod.objects.filter(financial_year=batch.financial_year).first()
        next_opening = getattr(period, "opening_batch_next_year", None) if period else None
        rows = []
        total_before = Decimal("0")
        total_after = Decimal("0")
        total_next_before = Decimal("0")
        total_next_after = Decimal("0")
        total_delta = Decimal("0")

        for line in batch.lines.select_related("material", "location", "granule_code", "grade").order_by("created_at"):
            if cls._blocking_errors(line.row_errors):
                continue
            before_qty = q4(_dec(line.system_qty))
            corrected_qty = q4(_dec(line.counted_qty))
            delta_qty = q4(corrected_qty - before_qty)
            next_opening_qty = Decimal("0")
            if next_opening:
                qs = next_opening.lines.filter(
                    stock_class=line.stock_class,
                    material=line.material,
                    plant=line.plant,
                    location=line.location,
                )
                if line.stock_class == "BULK":
                    qs = qs.filter(granule_code=line.granule_code)
                elif line.stock_class == "ROLL" and line.label_id:
                    qs = qs.filter(label_id=line.label_id)
                elif line.stock_class == "ROLL" and line.grade_id:
                    qs = qs.filter(grade=line.grade)
                next_opening_qty = q4(sum((_dec(row.opening_qty) for row in qs), Decimal("0")))
            next_after = q4(next_opening_qty + delta_qty)
            total_before += before_qty
            total_after += corrected_qty
            total_next_before += next_opening_qty
            total_next_after += next_after
            total_delta += delta_qty
            rows.append(
                {
                    "line_id": str(line.id),
                    "stock_class": line.stock_class,
                    "material_code": getattr(line.material, "code", ""),
                    "material_name": getattr(line.material, "name", ""),
                    "location_name": getattr(line.location, "name", ""),
                    "closed_fy_before": float(before_qty),
                    "closed_fy_after": float(corrected_qty),
                    "next_fy_before": float(next_opening_qty),
                    "next_fy_after": float(next_after),
                    "delta_qty": float(delta_qty),
                    "rate": float(line.rate or 0),
                    "value": float(line.value or 0),
                }
            )

        return {
            "closed_financial_year": batch.financial_year,
            "next_financial_year": next_fy,
            "next_opening_batch_no": getattr(next_opening, "batch_no", "") if next_opening else "",
            "totals": {
                "closed_before": float(q4(total_before)),
                "closed_after": float(q4(total_after)),
                "next_before": float(q4(total_next_before)),
                "next_after": float(q4(total_next_after)),
                "delta_qty": float(q4(total_delta)),
            },
            "rows": rows[:200],
        }

    @classmethod
    @transaction.atomic
    def submit_batch(cls, *, batch: InventoryAuditBatch, user=None) -> InventoryAuditBatch:
        batch = InventoryAuditBatch.objects.select_for_update().get(id=batch.id)
        if batch.status != "DRAFT":
            raise ValidationError("Only draft batches can be submitted.")
        validation = cls.validate_batch(batch=batch)
        if not validation["ok"]:
            raise ValidationError({"lines": validation["errors"]})
        if not batch.lines.exists():
            raise ValidationError("Add at least one audit line before submitting.")
        before = {"status": batch.status}
        workflow = dict((batch.summary_json or {}).get("workflow") or {})
        workflow.update(
            {
                "submitted_by": str(getattr(user, "id", "")) if getattr(user, "is_authenticated", False) else None,
                "submitted_by_name": getattr(user, "username", "") if getattr(user, "is_authenticated", False) else "",
                "submitted_at": timezone.now().isoformat(),
            }
        )
        summary = dict(batch.summary_json or {})
        summary["workflow"] = workflow
        batch.status = "SUBMITTED"
        batch.summary_json = summary
        batch.save(update_fields=["status", "summary_json", "updated_at"])
        cls._emit_batch_event(batch=batch, user=user, action="INVENTORY_BATCH_SUBMITTED", before=before, after={"status": batch.status, "workflow": workflow})
        return batch

    @classmethod
    @transaction.atomic
    def approve_batch(cls, *, batch: InventoryAuditBatch, user=None) -> InventoryAuditBatch:
        batch = InventoryAuditBatch.objects.select_for_update().get(id=batch.id)
        if batch.status != "SUBMITTED":
            raise ValidationError("Only submitted batches can be approved.")
        workflow = dict((batch.summary_json or {}).get("workflow") or {})
        submitted_by = workflow.get("submitted_by")
        actor_id = str(getattr(user, "id", "")) if getattr(user, "is_authenticated", False) else None
        if submitted_by and actor_id and submitted_by == actor_id and not (getattr(user, "is_owner", False) or getattr(user, "is_superuser", False)):
            raise ValidationError("Maker and checker must be different users.")
        before = {"status": batch.status, "workflow": workflow}
        workflow.update(
            {
                "approved_by": actor_id,
                "approved_by_name": getattr(user, "username", "") if getattr(user, "is_authenticated", False) else "",
                "approved_at": timezone.now().isoformat(),
            }
        )
        summary = dict(batch.summary_json or {})
        summary["workflow"] = workflow
        batch.status = "APPROVED"
        batch.summary_json = summary
        batch.save(update_fields=["status", "summary_json", "updated_at"])
        cls._emit_batch_event(batch=batch, user=user, action="INVENTORY_BATCH_APPROVED", before=before, after={"status": batch.status, "workflow": workflow})
        return batch

    @classmethod
    @transaction.atomic
    def cancel_batch(cls, *, batch: InventoryAuditBatch, user=None, reason: Optional[str] = None) -> InventoryAuditBatch:
        batch = InventoryAuditBatch.objects.select_for_update().get(id=batch.id)
        if batch.status in {"POSTED", "LOCKED"}:
            raise ValidationError("Posted or locked batches cannot be cancelled.")
        before = {"status": batch.status}
        batch.status = "CANCELLED"
        if reason:
            batch.notes = f"{batch.notes}\nCANCELLED: {reason}".strip()
        batch.save(update_fields=["status", "notes", "updated_at"])
        cls._emit_batch_event(batch=batch, user=user, action="INVENTORY_BATCH_CANCELLED", before=before, after={"status": batch.status, "reason": reason or ""})
        return batch

    @classmethod
    def resolve_system_qty(cls, line: InventoryAuditLine) -> Decimal:
        stock_class = str(line.stock_class or "").upper()
        if stock_class == "BULK":
            qs = InventoryBulk.objects.filter(material=line.material, plant=line.plant, location=line.location)
            if line.granule_code_id:
                qs = qs.filter(granule_code=line.granule_code)
            return qs.aggregate(total=Sum("qty_kg")).get("total") or Decimal("0")
        if stock_class == "PACKAGING":
            stock = PackagingStock.objects.filter(material=line.material, plant=line.plant, location=line.location).first()
            return _dec(getattr(stock, "qty", 0))
        if stock_class == "ROLL":
            qs = InventoryRoll.objects.filter(material=line.material, location=line.location)
            if line.label_id:
                qs = qs.filter(label_id=line.label_id)
            if line.grade_id:
                qs = qs.filter(grade=line.grade)
            return qs.aggregate(total=Sum("weight_kg")).get("total") or Decimal("0")
        return Decimal("0")

    @classmethod
    @transaction.atomic
    def post_batch(cls, *, batch: InventoryAuditBatch, user=None) -> InventoryAuditBatch:
        batch = InventoryAuditBatch.objects.select_for_update().get(id=batch.id)
        if batch.status not in {"DRAFT", "APPROVED"}:
            raise ValidationError("Only approved batches can be posted. Draft posting is kept only for legacy API callers.")
        if batch.type != "FY_CORRECTION":
            cls.assert_period_allows_cutoff(batch.financial_year, batch.cutoff_at)
        validation = cls.validate_batch(batch=batch)
        if not validation["ok"]:
            raise ValidationError({"lines": validation["errors"]})
        if not batch.lines.exists():
            raise ValidationError("Add at least one audit line before posting.")

        before = {"status": batch.status, "summary": batch.summary_json}
        for line in batch.lines.select_related("material", "location", "grade", "granule_code").order_by("created_at"):
            existing_refs = dict(line.posted_reference_json or {})
            refs = cls._post_line(line=line, user=user)
            resolved_refs = dict(line.posted_reference_json or {})
            line.posted_reference_json = {**existing_refs, **resolved_refs, **refs}
            line.save(update_fields=["posted_reference_json", "updated_at"])

        batch.status = "POSTED"
        batch.posted_by = user if getattr(user, "is_authenticated", False) else None
        batch.posted_at = timezone.now()
        cls.refresh_batch_summary(batch, save=False)
        batch.save(update_fields=["status", "posted_by", "posted_at", "summary_json", "updated_at"])
        cls._emit_batch_event(batch=batch, user=user, action="INVENTORY_BATCH_POSTED", before=before, after={"status": batch.status, "summary": batch.summary_json})
        return batch

    @classmethod
    def _post_line(cls, *, line: InventoryAuditLine, user=None) -> Dict[str, Any]:
        if line.batch.type == "OPENING_STOCK":
            return cls._post_opening_line(line=line, user=user)
        elif line.batch.type in {"PHYSICAL_COUNT", "FY_CORRECTION"}:
            cls._apply_resolved_rate(line)
            qty = q4(_dec(line.counted_qty) - _dec(line.system_qty))
        elif line.batch.type == "FY_CLOSE":
            return {"snapshot": True}
        else:
            raise ValidationError("Unsupported audit batch type.")

        if qty == 0:
            return {"skipped": "zero_variance"}
        reference = f"{line.batch.type}:{line.batch.financial_year}:{line.batch.batch_no}"
        stock_class = str(line.stock_class).upper()
        rollforward_refs: Dict[str, Any] = {}
        if stock_class == "BULK":
            tx_type = cls._delta_tx_type(batch_type=line.batch.type, qty=qty)
            tx = cls._adjust_bulk(line=line, qty=qty, reference=reference, tx_type=tx_type)
            refs = {"bulk_transaction_id": str(tx.id)}
        elif stock_class == "PACKAGING":
            tx_type = cls._delta_tx_type(batch_type=line.batch.type, qty=qty)
            tx = cls._adjust_packaging(line=line, qty=qty, reference=reference, tx_type=tx_type)
            refs = {"packaging_transaction_id": str(tx.id)}
        elif stock_class == "ROLL":
            refs = cls._adjust_roll(line=line, qty=qty, reference=reference, user=user)
        else:
            raise ValidationError("Unsupported stock class.")
        if line.batch.type == "FY_CORRECTION":
            rollforward_refs = cls._sync_next_opening_after_correction(line=line, delta_qty=qty, user=user)
        return {**refs, **rollforward_refs}

    @classmethod
    def _post_opening_line(cls, *, line: InventoryAuditLine, user=None) -> Dict[str, Any]:
        rate, rate_source = cls._apply_resolved_rate(line)
        qty = q4(line.opening_qty)
        reference = f"{line.batch.type}:{line.batch.financial_year}:{line.batch.batch_no}"
        stock_class = str(line.stock_class).upper()
        additive_cutover = cls._is_cutover_opening(line.batch)
        if qty == 0:
            return {"skipped": "zero_opening"}
        if additive_cutover and stock_class == "BULK":
            before = cls.resolve_system_qty(line)
            tx = BulkService.add_bulk(
                str(line.material_id),
                qty,
                str(line.plant_id),
                str(line.location_id),
                cost=rate,
                reference=reference,
                tx_type="OPENING_BALANCE",
                granule_code_id=str(line.granule_code_id) if line.granule_code_id else None,
                qty_uom=line.uom or getattr(line.material, "base_uom", None),
            )
            after = q4(before + qty)
            return {
                "bulk_transaction_id": str(tx.id),
                "before_qty": float(before),
                "after_qty": float(after),
                "additive_cutover": True,
                "source": "CUTOVER_OPENING_ADDITIVE",
                "rate_source": rate_source,
            }
        if additive_cutover and stock_class == "PACKAGING":
            before = cls.resolve_system_qty(line)
            tx = PackagingService.add_packaging_stock(
                material_id=str(line.material_id),
                qty=qty,
                location_id=str(line.location_id),
                cost=rate,
                reference=reference,
                tx_type="OPENING_BALANCE",
                input_uom=line.uom or getattr(line.material, "base_uom", None),
                meta_json={
                    "inventory_audit_batch": str(line.batch_id),
                    "inventory_audit_line": str(line.id),
                    "additive_cutover": True,
                    "source": "CUTOVER_OPENING_ADDITIVE",
                    "rate_source": rate_source,
                },
            )
            after = q4(before + qty)
            return {
                "packaging_transaction_id": str(tx.id),
                "before_qty": float(before),
                "after_qty": float(after),
                "additive_cutover": True,
                "source": "CUTOVER_OPENING_ADDITIVE",
                "rate_source": rate_source,
            }
        if stock_class == "BULK":
            bulk, _ = InventoryBulk.objects.select_for_update().get_or_create(
                material=line.material,
                granule_code=line.granule_code,
                plant=line.plant,
                location=line.location,
                defaults={"qty_kg": Decimal("0"), "avg_cost": Decimal("0")},
            )
            before = q4(bulk.qty_kg)
            bulk.qty_kg = qty
            bulk.avg_cost = rate
            bulk.save(update_fields=["qty_kg", "avg_cost", "updated_at"])
            tx = BulkTransaction.objects.create(
                material=line.material,
                granule_code=line.granule_code,
                location=line.location,
                type="OPENING_BALANCE",
                qty_kg=qty,
                avg_cost=rate,
                reference=reference,
            )
            return {"bulk_transaction_id": str(tx.id), "before_qty": float(before), "after_qty": float(qty), "absolute": True, "rate_source": rate_source}
        if stock_class == "PACKAGING":
            stock, _ = PackagingStock.objects.select_for_update().get_or_create(
                material=line.material,
                plant=line.plant,
                location=line.location,
                defaults={"qty": Decimal("0"), "avg_cost": Decimal("0")},
            )
            before = q4(stock.qty)
            stock.qty = qty
            stock.avg_cost = rate
            stock.save(update_fields=["qty", "avg_cost", "updated_at"])
            tx = PackagingTransaction.objects.create(
                type="OPENING_BALANCE",
                material=line.material,
                location=line.location,
                qty=qty,
                avg_cost=rate,
                reference=reference,
                meta_json={"inventory_audit_batch": str(line.batch_id), "inventory_audit_line": str(line.id), "absolute": True, "rate_source": rate_source},
            )
            return {"packaging_transaction_id": str(tx.id), "before_qty": float(before), "after_qty": float(qty), "absolute": True, "rate_source": rate_source}
        if stock_class == "ROLL":
            refs = dict(line.posted_reference_json or {})
            roll = RollService.create_roll(
                material=line.material,
                weight_kg=qty,
                location=line.location,
                width_mm=line.width_mm,
                thickness_micron=line.thickness_micron,
                batch_no=line.batch_no or None,
                length_m=line.length_m or Decimal("0"),
                grade=line.grade,
                is_fg=line.is_fg,
                plant=line.plant,
                user=user if getattr(user, "is_authenticated", False) else None,
                notes=reference,
                stock_form=refs.get("stock_form") or "OPEN_WEB",
                width_basis=refs.get("width_basis") or "",
            )
            if line.label_id and line.label_id != roll.label_id:
                roll.label_id = line.label_id
            meta = dict(roll.meta_json or {})
            meta.update({
                "inventory_audit_batch": str(line.batch_id),
                "inventory_audit_line": str(line.id),
                "source": "CUTOVER_OPENING_ADDITIVE" if additive_cutover else "OPENING_STOCK",
                "stock_form": roll.stock_form,
                "width_basis": roll.width_basis,
            })
            roll.meta_json = meta
            roll.save(update_fields=["label_id", "meta_json"])
            RollMovement.objects.filter(roll=roll, reason_note=reference).update(reason="OPENING_BALANCE")
            return {"roll_id": str(roll.id), "roll_label": roll.label_id, "absolute": not additive_cutover, "additive_cutover": additive_cutover, "stock_form": roll.stock_form, "width_basis": roll.width_basis}
        raise ValidationError("Unsupported stock class.")

    @classmethod
    def _delta_tx_type(cls, *, batch_type: str, qty: Decimal) -> str:
        if batch_type == "FY_CORRECTION":
            return "FY_CORRECTION"
        return "COUNT_EXCESS" if qty > 0 else "COUNT_SHORT"

    @classmethod
    def _sync_next_opening_after_correction(cls, *, line: InventoryAuditLine, delta_qty: Decimal, user=None) -> Dict[str, Any]:
        period = cls.require_closed_period_for_correction(line.batch.financial_year)
        next_opening = period.opening_batch_next_year
        if not next_opening:
            next_start = period.end_date + timedelta(days=1)
            next_fy = f"{next_start.year}-{next_start.year + 1}"
            next_opening = InventoryAuditBatch.objects.filter(
                type="OPENING_STOCK",
                plant=line.plant,
                financial_year=next_fy,
                status__in=["POSTED", "LOCKED"],
            ).order_by("cutoff_at", "created_at").first()
            if not next_opening:
                next_opening = InventoryAuditBatch.objects.create(
                    type="OPENING_STOCK",
                    plant=line.plant,
                    financial_year=next_fy,
                    cutoff_at=timezone.make_aware(datetime.combine(next_start, time.min)),
                    status="POSTED",
                    posted_by=user if getattr(user, "is_authenticated", False) else None,
                    posted_at=timezone.now(),
                    notes=f"System generated to carry FY correction from {line.batch.financial_year}.",
                )
            period.opening_batch_next_year = next_opening
            period.save(update_fields=["opening_batch_next_year", "updated_at"])

        qs = next_opening.lines.select_for_update().filter(
            stock_class=line.stock_class,
            material=line.material,
            plant=line.plant,
            location=line.location,
        )
        if line.stock_class == "BULK":
            qs = qs.filter(granule_code=line.granule_code)
        elif line.stock_class == "ROLL" and line.label_id:
            qs = qs.filter(label_id=line.label_id)
        elif line.stock_class == "ROLL" and line.grade_id:
            qs = qs.filter(grade=line.grade)

        before_qty = q4(sum((_dec(row.opening_qty) for row in qs), Decimal("0")))
        delta_qty = q4(delta_qty)
        after_qty = q4(before_qty + delta_qty)
        if after_qty < 0:
            raise ValidationError("FY correction would make next year opening stock negative.")
        if delta_qty == 0:
            return {"next_opening_sync": "skipped_zero_delta"}

        adjustment_line = InventoryAuditLine.objects.create(
            batch=next_opening,
            stock_class=line.stock_class,
            material=line.material,
            granule_code=line.granule_code,
            grade=line.grade,
            plant=line.plant,
            location=line.location,
            uom=line.uom or getattr(line.material, "base_uom", "") or "KG",
            system_qty=delta_qty,
            counted_qty=delta_qty,
            opening_qty=delta_qty,
            rate=line.rate,
            label_id=line.label_id,
            batch_no=line.batch_no,
            width_mm=line.width_mm,
            thickness_micron=line.thickness_micron,
            length_m=line.length_m,
            is_fg=line.is_fg,
            stage_index=line.stage_index,
            status=line.status or "AVAILABLE",
            packaging_kind=line.packaging_kind,
            base_uom=line.base_uom,
            posted_reference_json={
                "source": "OPENING_BALANCE_ADJUST",
                "correction_batch_id": str(line.batch_id),
                "correction_batch_no": line.batch.batch_no,
                "correction_line_id": str(line.id),
                "delta_qty": float(delta_qty),
                "before_opening_qty": float(before_qty),
                "after_opening_qty": float(after_qty),
                "actor": getattr(user, "username", "") if getattr(user, "is_authenticated", False) else "",
                "at": timezone.now().isoformat(),
            },
        )
        cls.refresh_batch_summary(next_opening)
        cls._emit_batch_event(
            batch=next_opening,
            user=user,
            action="INVENTORY_FY_CORRECTION_ROLLFORWARD_SYNCED",
            before={"opening_qty": float(before_qty), "source_batch": line.batch.batch_no},
            after={"opening_qty": float(after_qty), "source_batch": line.batch.batch_no},
        )
        return {"next_opening_batch_id": str(next_opening.id), "next_opening_line_id": str(adjustment_line.id), "next_opening_delta_qty": float(delta_qty)}

    @classmethod
    def _adjust_bulk(cls, *, line: InventoryAuditLine, qty: Decimal, reference: str, tx_type: str) -> BulkTransaction:
        if qty > 0:
            return BulkService.add_bulk(
                str(line.material_id),
                qty,
                str(line.plant_id),
                str(line.location_id),
                cost=line.rate or 0,
                reference=reference,
                tx_type=tx_type,
                granule_code_id=str(line.granule_code_id) if line.granule_code_id else None,
                qty_uom=line.uom or getattr(line.material, "base_uom", None),
            )
        tx = BulkService.consume_bulk(
            str(line.material_id),
            abs(qty),
            str(line.location_id),
            reference=reference,
            granule_code_id=str(line.granule_code_id) if line.granule_code_id else None,
            qty_uom=line.uom or getattr(line.material, "base_uom", None),
        )
        tx.type = tx_type
        tx.save(update_fields=["type"])
        return tx

    @classmethod
    def _adjust_packaging(cls, *, line: InventoryAuditLine, qty: Decimal, reference: str, tx_type: str) -> PackagingTransaction:
        if qty > 0:
            return PackagingService.add_packaging_stock(
                material_id=str(line.material_id),
                qty=qty,
                location_id=str(line.location_id),
                cost=line.rate or 0,
                reference=reference,
                tx_type=tx_type,
                input_uom=line.uom or getattr(line.material, "base_uom", None),
                meta_json={"inventory_audit_batch": str(line.batch_id), "inventory_audit_line": str(line.id)},
            )
        tx = PackagingService.consume_packaging_stock(
            material_id=str(line.material_id),
            qty=abs(qty),
            location_id=str(line.location_id),
            reference=reference,
            input_uom=line.uom or getattr(line.material, "base_uom", None),
            meta_json={"inventory_audit_batch": str(line.batch_id), "inventory_audit_line": str(line.id)},
        )
        tx.type = tx_type
        tx.save(update_fields=["type"])
        return tx

    @classmethod
    def _adjust_roll(cls, *, line: InventoryAuditLine, qty: Decimal, reference: str, user=None) -> Dict[str, Any]:
        if line.label_id:
            roll = (
                InventoryRoll.objects.select_for_update()
                .filter(material=line.material, location=line.location, label_id=line.label_id)
                .first()
            )
            if roll:
                before_weight = _dec(roll.weight_kg)
                new_weight = before_weight + qty
                if new_weight < 0:
                    raise ValidationError(f"Roll adjustment would make {roll.label_id} negative.")
                roll.weight_kg = new_weight
                if new_weight == 0:
                    roll.status = "MISSING" if line.batch.type == "PHYSICAL_COUNT" else "SCRAPPED"
                elif roll.status in {"MISSING", "SCRAPPED"}:
                    roll.status = "AVAILABLE"
                roll.save(update_fields=["weight_kg", "status"])
                RollMovement.objects.create(
                    roll=roll,
                    from_location=roll.location,
                    to_location=roll.location,
                    reason=cls._delta_tx_type(batch_type=line.batch.type, qty=qty),
                    reason_note=reference,
                    moved_by=user if getattr(user, "is_authenticated", False) else None,
                )
                return {"roll_id": str(roll.id), "roll_label": roll.label_id, "before_qty": float(before_weight), "after_qty": float(new_weight)}

        if qty > 0:
            refs = dict(line.posted_reference_json or {})
            roll = RollService.create_roll(
                material=line.material,
                weight_kg=qty,
                location=line.location,
                width_mm=line.width_mm,
                thickness_micron=line.thickness_micron,
                batch_no=line.batch_no or None,
                length_m=line.length_m or Decimal("0"),
                grade=line.grade,
                is_fg=line.is_fg,
                plant=line.plant,
                user=user if getattr(user, "is_authenticated", False) else None,
                notes=reference,
                stock_form=refs.get("stock_form") or "OPEN_WEB",
                width_basis=refs.get("width_basis") or "",
            )
            movement_reason = cls._delta_tx_type(batch_type=line.batch.type, qty=qty)
            RollMovement.objects.filter(roll=roll, reason_note=reference).update(reason=movement_reason)
            if line.label_id and line.label_id != roll.label_id:
                roll.label_id = line.label_id
                roll.save(update_fields=["label_id"])
            return {"roll_id": str(roll.id), "roll_label": roll.label_id}

        qs = InventoryRoll.objects.select_for_update().filter(material=line.material, location=line.location)
        if line.label_id:
            qs = qs.filter(label_id=line.label_id)
        roll = qs.order_by("created_at").first()
        if not roll:
            raise ValidationError("No roll exists for the negative roll adjustment.")
        new_weight = _dec(roll.weight_kg) + qty
        if new_weight < 0:
            raise ValidationError(f"Roll adjustment would make {roll.label_id} negative.")
        roll.weight_kg = new_weight
        if new_weight == 0:
            roll.status = "MISSING" if line.batch.type == "PHYSICAL_COUNT" else "SCRAPPED"
        roll.save(update_fields=["weight_kg", "status"])
        movement = RollMovement.objects.create(
            roll=roll,
            from_location=roll.location,
            to_location=roll.location,
            reason=cls._delta_tx_type(batch_type=line.batch.type, qty=qty),
            reason_note=reference,
            moved_by=user if getattr(user, "is_authenticated", False) else None,
        )
        return {"roll_id": str(roll.id), "roll_movement_id": str(movement.id)}

    @classmethod
    def refresh_batch_summary(cls, batch: InventoryAuditBatch, save: bool = True) -> Dict[str, Any]:
        existing_summary = batch.summary_json if isinstance(batch.summary_json, dict) else {}
        existing_workflow = existing_summary.get("workflow") if isinstance(existing_summary.get("workflow"), dict) else None
        summary = {
            "lines": batch.lines.count(),
            "bulk_kg": 0.0,
            "roll_kg": 0.0,
            "packaging_qty": 0.0,
            "wip_kg": 0.0,
            "fg_kg": 0.0,
            "scrap_kg": 0.0,
            "value": 0.0,
            "missing_rates": 0,
            "errors": 0,
        }
        for line in batch.lines.all():
            qty = _dec(line.opening_qty or line.counted_qty or line.system_qty)
            if batch.type in {"PHYSICAL_COUNT", "FY_CORRECTION"}:
                qty = abs(_dec(line.variance_qty))
            if line.stock_class == "BULK":
                summary["bulk_kg"] += float(qty)
            elif line.stock_class == "ROLL":
                summary["roll_kg"] += float(qty)
                if line.is_fg:
                    summary["fg_kg"] += float(qty)
                else:
                    summary["wip_kg"] += float(qty)
            elif line.stock_class == "PACKAGING":
                summary["packaging_qty"] += float(qty)
            summary["value"] += float(line.value or 0)
            if line.rate in (None, Decimal("0")):
                summary["missing_rates"] += 1
            if cls._blocking_errors(line.row_errors):
                summary["errors"] += 1
        if existing_workflow:
            summary["workflow"] = existing_workflow
        batch.summary_json = summary
        if save:
            batch.save(update_fields=["summary_json", "updated_at"])
        return summary

    @classmethod
    def closing_preview(cls, *, plant_id: Optional[str], financial_year: Optional[str]) -> Dict[str, Any]:
        if not financial_year:
            financial_year = current_indian_financial_year()
        plant = Plant.objects.filter(id=plant_id).first() if plant_id else Plant.objects.first()
        if not plant:
            return {"financial_year": financial_year, "plant": None, "rows": [], "totals": {}, "blockers": []}
        rows = cls.current_stock_snapshot_rows(plant=plant)
        totals = cls._snapshot_totals(rows)
        movements = cls.period_movement_summary(plant=plant, financial_year=financial_year)
        blockers = cls.close_blockers(plant=plant)
        return {
            "financial_year": financial_year,
            "plant": {"id": str(plant.id), "name": plant.name, "code": plant.code},
            "rows": rows,
            "totals": totals,
            "movements": movements,
            "blockers": blockers,
        }

    @classmethod
    def stock_snapshot(
        cls,
        *,
        plant_id: Optional[str],
        stock_class: Optional[str] = None,
        location_id: Optional[str] = None,
        material_id: Optional[str] = None,
        query: Optional[str] = None,
    ) -> Dict[str, Any]:
        plant = Plant.objects.filter(id=plant_id).first() if plant_id else Plant.objects.first()
        if not plant:
            return {"plant": None, "rows": [], "totals": {}}
        rows = cls.current_stock_snapshot_rows(plant=plant)
        stock_class = str(stock_class or "").upper().strip()
        if stock_class:
            rows = [row for row in rows if row.get("stock_class") == stock_class]
        if location_id:
            rows = [row for row in rows if str(row.get("location")) == str(location_id)]
        if material_id:
            rows = [row for row in rows if str(row.get("material")) == str(material_id)]
        if query:
            needle = str(query).strip().lower()
            rows = [
                row
                for row in rows
                if needle in str(row.get("material_code") or "").lower()
                or needle in str(row.get("material_name") or "").lower()
                or needle in str(row.get("label_id") or "").lower()
                or needle in str(row.get("location_name") or "").lower()
            ]
        return {
            "plant": {"id": str(plant.id), "name": plant.name, "code": plant.code},
            "rows": rows,
            "totals": cls._snapshot_totals(rows),
        }

    @classmethod
    @transaction.atomic
    def load_batch_from_snapshot(
        cls,
        *,
        batch: InventoryAuditBatch,
        stock_class: Optional[str] = None,
        location_id: Optional[str] = None,
        material_id: Optional[str] = None,
        query: Optional[str] = None,
        replace_existing: bool = True,
    ) -> InventoryAuditBatch:
        if batch.type not in {"PHYSICAL_COUNT", "FY_CORRECTION"}:
            raise ValidationError("System stock preload is only available for stock count and FY correction batches.")
        snapshot = cls.stock_snapshot(
            plant_id=str(batch.plant_id),
            stock_class=stock_class,
            location_id=location_id,
            material_id=material_id,
            query=query,
        )
        rows = []
        for row in snapshot["rows"]:
            rows.append(
                {
                    "stock_class": row.get("stock_class"),
                    "material": row.get("material"),
                    "granule_code": row.get("granule_code"),
                    "grade": row.get("grade"),
                    "location": row.get("location"),
                    "uom": row.get("uom"),
                    "counted_qty": row.get("qty"),
                    "label_id": row.get("label_id"),
                    "batch_no": row.get("batch_no"),
                    "width_mm": row.get("width_mm"),
                    "thickness_micron": row.get("thickness_micron"),
                    "length_m": row.get("length_m"),
                    "is_fg": row.get("is_fg"),
                    "stage_index": row.get("stage_index"),
                    "status": row.get("status"),
                    "packaging_kind": row.get("packaging_kind"),
                    "base_uom": row.get("base_uom"),
                    "rate": row.get("rate") if row.get("rate") not in (None, "", 0, 0.0) else None,
                }
            )
        if replace_existing:
            cls.replace_lines(batch=batch, rows=rows)
        else:
            cls.import_lines(batch=batch, rows=rows)
        batch.refresh_from_db()
        return batch

    @classmethod
    def current_stock_snapshot_rows(cls, *, plant: Plant) -> List[Dict[str, Any]]:
        rows: List[Dict[str, Any]] = []
        for bulk in InventoryBulk.objects.select_related("material", "granule_code", "location").filter(plant=plant, qty_kg__gt=0):
            rows.append({
                "stock_class": "BULK",
                "material": str(bulk.material_id),
                "material_category": bulk.material.category,
                "category": bulk.material.category,
                "material_code": bulk.material.code,
                "material_name": bulk.material.name,
                "granule_code": str(bulk.granule_code_id) if bulk.granule_code_id else None,
                "granule_code_label": bulk.granule_code.code if bulk.granule_code else None,
                "location": str(bulk.location_id),
                "location_name": bulk.location.name,
                "qty": float(bulk.qty_kg or 0),
                "uom": bulk.material.base_uom or "KG",
                "rate": float(bulk.avg_cost or 0),
                "rate_source": "MATERIAL_AVG" if bulk.avg_cost else "ZERO",
                "rate_missing": not bool(bulk.avg_cost),
            })
        for roll in InventoryRoll.objects.select_related("material", "grade", "location").filter(location__plant=plant).exclude(status__in=["CONSUMED", "SCRAPPED"]):
            rate, rate_source = _roll_rate(roll)
            rows.append({
                "stock_class": "ROLL",
                "material": str(roll.material_id),
                "material_category": roll.material.category if roll.material else "",
                "category": roll.material.category if roll.material else "",
                "material_code": roll.material.code if roll.material else "",
                "material_name": roll.material.name if roll.material else "",
                "grade": str(roll.grade_id) if roll.grade_id else None,
                "grade_name": roll.grade.name if roll.grade else None,
                "location": str(roll.location_id),
                "location_name": roll.location.name,
                "qty": float(roll.weight_kg or 0),
                "uom": "KG",
                "label_id": roll.label_id,
                "batch_no": roll.batch_no,
                "width_mm": float(roll.width_mm or 0),
                "thickness_micron": float(roll.thickness_micron or 0),
                "length_m": float(roll.length_m or 0),
                "stock_form": roll.stock_form,
                "width_basis": roll.width_basis,
                "is_fg": bool(roll.is_fg),
                "stage_index": roll.stage_index,
                "status": roll.status,
                "rate": float(rate),
                "rate_source": rate_source,
                "rate_missing": rate <= 0,
            })
        for stock in PackagingStock.objects.select_related("material", "location").filter(plant=plant, qty__gt=0):
            rows.append({
                "stock_class": "PACKAGING",
                "material": str(stock.material_id),
                "material_category": stock.material.category,
                "category": stock.material.category,
                "material_code": stock.material.code,
                "material_name": stock.material.name,
                "location": str(stock.location_id),
                "location_name": stock.location.name,
                "qty": float(stock.qty or 0),
                "uom": stock.material.base_uom,
                "packaging_kind": stock.material.packaging_kind,
                "base_uom": stock.material.base_uom,
                "rate": float(stock.avg_cost or 0),
                "rate_source": "MATERIAL_AVG" if stock.avg_cost else "ZERO",
                "rate_missing": not bool(stock.avg_cost),
            })
        return rows

    @classmethod
    def _snapshot_totals(cls, rows: Iterable[Dict[str, Any]]) -> Dict[str, Any]:
        totals = defaultdict(float)
        for row in rows:
            stock_class = row["stock_class"]
            qty = float(row.get("qty") or 0)
            if stock_class == "BULK":
                totals["bulk_kg"] += qty
            elif stock_class == "ROLL":
                totals["roll_kg"] += qty
                totals["fg_kg" if row.get("is_fg") else "wip_kg"] += qty
            elif stock_class == "PACKAGING":
                totals["packaging_qty"] += qty
            totals["value"] += qty * float(row.get("rate") or 0)
        totals["rows"] = len(list(rows)) if not isinstance(rows, list) else len(rows)
        return dict(totals)

    @classmethod
    def period_movement_summary(cls, *, plant: Plant, financial_year: str) -> Dict[str, float]:
        start_date, end_date = financial_year_dates(financial_year)
        bulk_qs = BulkTransaction.objects.filter(location__plant=plant, created_at__date__gte=start_date, created_at__date__lte=end_date)
        packaging_qs = PackagingTransaction.objects.filter(location__plant=plant, created_at__date__gte=start_date, created_at__date__lte=end_date)
        roll_qs = RollMovement.objects.filter(to_location__plant=plant, timestamp__date__gte=start_date, timestamp__date__lte=end_date)
        movement = defaultdict(float)
        for tx in bulk_qs:
            movement[f"bulk_{tx.type.lower()}_kg"] += float(tx.qty_kg or 0)
        for tx in packaging_qs:
            movement[f"packaging_{tx.type.lower()}"] += float(tx.qty or 0)
        for mv in roll_qs.select_related("roll"):
            movement[f"roll_{mv.reason.lower()}_kg"] += float(getattr(mv.roll, "weight_kg", 0) or 0)
        return dict(movement)

    @classmethod
    def close_blockers(cls, *, plant: Plant) -> List[Dict[str, Any]]:
        blockers = []
        negative_bulk = InventoryBulk.objects.filter(plant=plant, qty_kg__lt=0).count()
        if negative_bulk:
            blockers.append({"code": "NEGATIVE_BULK", "label": "Negative bulk stock", "count": negative_bulk})
        open_challans = DeliveryChallan.objects.filter(from_plant=plant, status__in=["DRAFT", "APPROVED", "IN_TRANSIT"]).count()
        if open_challans:
            blockers.append({"code": "OPEN_INTERPLANT", "label": "Open inter-plant challans", "count": open_challans})
        open_jobwork = JobWorkOrder.objects.filter(plant=plant, status__in=["DRAFT", "SENT", "PARTIAL"]).count()
        if open_jobwork:
            blockers.append({"code": "OPEN_JOBWORK", "label": "Open jobwork orders", "count": open_jobwork})
        critical_alerts = InventoryAlert.objects.filter(plant=plant, resolved=False, severity="CRITICAL").count()
        if critical_alerts:
            blockers.append({"code": "CRITICAL_ALERTS", "label": "Unresolved critical inventory alerts", "count": critical_alerts})
        draft_batches = InventoryAuditBatch.objects.filter(plant=plant, status="DRAFT").count()
        if draft_batches:
            blockers.append({"code": "DRAFT_AUDIT_BATCHES", "label": "Draft inventory audit batches", "count": draft_batches})
        return blockers

    @classmethod
    @transaction.atomic
    def begin_close(cls, *, period: InventoryFinancialPeriod) -> InventoryFinancialPeriod:
        if period.status == "CLOSED":
            raise ValidationError("Financial year is already closed.")
        period.status = "CLOSING_IN_PROGRESS"
        period.save(update_fields=["status", "updated_at"])
        return period

    @classmethod
    @transaction.atomic
    def close_period(cls, *, period: InventoryFinancialPeriod, plant_id: str, user=None) -> InventoryFinancialPeriod:
        period = InventoryFinancialPeriod.objects.select_for_update().get(id=period.id)
        if period.status == "CLOSED":
            raise ValidationError("Financial year is already closed.")
        plant = Plant.objects.get(id=plant_id)
        blockers = cls.close_blockers(plant=plant)
        if blockers:
            raise ValidationError({"blockers": blockers})

        cutoff_at = timezone.make_aware(datetime.combine(period.end_date, time.max))
        close_batch = InventoryAuditBatch.objects.create(
            type="FY_CLOSE",
            plant=plant,
            financial_year=period.financial_year,
            cutoff_at=cutoff_at,
            status="LOCKED",
            posted_by=user if getattr(user, "is_authenticated", False) else None,
            posted_at=timezone.now(),
            locked_by=user if getattr(user, "is_authenticated", False) else None,
            locked_at=timezone.now(),
            notes="System generated closing stock snapshot.",
        )
        snapshot_rows = cls.current_stock_snapshot_rows(plant=plant)
        for row in snapshot_rows:
            line = cls._line_from_snapshot(close_batch, row)
            line.system_qty = _dec(row["qty"])
            line.counted_qty = _dec(row["qty"])
            line.opening_qty = _dec(row["qty"])
            line.save()
        cls.refresh_batch_summary(close_batch)

        next_start = period.end_date + timedelta(days=1)
        next_fy = f"{next_start.year}-{next_start.year + 1}"
        next_opening = InventoryAuditBatch.objects.create(
            type="OPENING_STOCK",
            plant=plant,
            financial_year=next_fy,
            cutoff_at=timezone.make_aware(datetime.combine(next_start, time.min)),
            status="POSTED",
            posted_by=user if getattr(user, "is_authenticated", False) else None,
            posted_at=timezone.now(),
            notes=f"System generated from closing batch {close_batch.batch_no}.",
        )
        for row in snapshot_rows:
            line = cls._line_from_snapshot(next_opening, row)
            line.opening_qty = _dec(row["qty"])
            line.posted_reference_json = {"generated_from_closing_batch": str(close_batch.id), "no_physical_repost": True}
            line.save()
        cls.refresh_batch_summary(next_opening)

        period.status = "CLOSED"
        period.closed_by = user if getattr(user, "is_authenticated", False) else None
        period.closed_at = timezone.now()
        period.closing_batch = close_batch
        period.opening_batch_next_year = next_opening
        period.save(update_fields=["status", "closed_by", "closed_at", "closing_batch", "opening_batch_next_year", "updated_at"])

        next_start_date, next_end_date = financial_year_dates(next_fy)
        next_period = InventoryFinancialPeriod.objects.filter(financial_year=next_fy).first()
        if not next_period and not InventoryFinancialPeriod.objects.filter(status="OPEN").exists():
            InventoryFinancialPeriod.objects.create(
                financial_year=next_fy,
                start_date=next_start_date,
                end_date=next_end_date,
                status="OPEN",
            )
        return period

    @classmethod
    def _line_from_snapshot(cls, batch: InventoryAuditBatch, row: Dict[str, Any]) -> InventoryAuditLine:
        return InventoryAuditLine(
            batch=batch,
            stock_class=row["stock_class"],
            material_id=row["material"],
            granule_code_id=row.get("granule_code"),
            grade_id=row.get("grade"),
            plant=batch.plant,
            location_id=row["location"],
            uom=row.get("uom") or "KG",
            rate=_dec(row.get("rate")) if row.get("rate") not in (None, "") else None,
            label_id=row.get("label_id") or "",
            batch_no=row.get("batch_no") or "",
            width_mm=_dec(row.get("width_mm")) if row.get("width_mm") not in (None, "") else None,
            thickness_micron=_dec(row.get("thickness_micron")) if row.get("thickness_micron") not in (None, "") else None,
            length_m=_dec(row.get("length_m")) if row.get("length_m") not in (None, "") else None,
            is_fg=bool(row.get("is_fg")),
            stage_index=int(row.get("stage_index") or 0),
            status=row.get("status") or "AVAILABLE",
            packaging_kind=row.get("packaging_kind") or "",
            base_uom=row.get("base_uom") or "",
            posted_reference_json={
                key: row.get(key)
                for key in ("stock_form", "width_basis")
                if row.get(key) not in (None, "")
            },
        )

    @classmethod
    def stock_card(
        cls,
        *,
        material_id: Optional[str] = None,
        plant_id: Optional[str] = None,
        location_id: Optional[str] = None,
        financial_year: Optional[str] = None,
        date_from=None,
        date_to=None,
    ) -> Dict[str, Any]:
        auto_financial_year_window = bool(financial_year and not date_from and not date_to)
        if financial_year and not date_from and not date_to:
            date_from, date_to = financial_year_dates(financial_year)
        entries: List[Dict[str, Any]] = []

        batch_lines = InventoryAuditLine.objects.select_related("batch", "material", "location").filter(batch__status__in=["POSTED", "LOCKED"])
        if material_id:
            batch_lines = batch_lines.filter(material_id=material_id)
        if plant_id:
            batch_lines = batch_lines.filter(plant_id=plant_id)
        if location_id:
            batch_lines = batch_lines.filter(location_id=location_id)
        for line in batch_lines.order_by("batch__cutoff_at", "created_at"):
            if line.batch.type == "FY_CLOSE":
                continue
            if financial_year and line.batch.financial_year != financial_year:
                continue
            if not auto_financial_year_window and date_from and line.batch.cutoff_at.date() < date_from:
                continue
            if not auto_financial_year_window and date_to and line.batch.cutoff_at.date() > date_to:
                continue
            if line.batch.type == "OPENING_STOCK":
                qty = _dec(line.opening_qty)
            elif line.batch.type in {"PHYSICAL_COUNT", "FY_CORRECTION"}:
                qty = _dec(line.variance_qty)
            else:
                qty = _dec(line.counted_qty or line.system_qty or 0)
            source = "CUTOVER_OPENING" if cls._is_cutover_opening(line.batch) else line.batch.type
            if (line.posted_reference_json or {}).get("source") == "OPENING_BALANCE_ADJUST":
                source = "OPENING_BALANCE_ADJUST"
            entries.append(cls._stock_card_entry(
                at=line.batch.cutoff_at,
                source=source,
                reference=line.batch.batch_no,
                material=line.material,
                location=line.location,
                qty=qty,
                uom=line.uom,
                rate=line.rate,
                meta={"stock_class": line.stock_class, "line_id": str(line.id), "source_doc": "inventory_audit_batch"},
            ))

        bulk = BulkTransaction.objects.select_related("material", "location").all()
        if material_id:
            bulk = bulk.filter(material_id=material_id)
        if plant_id:
            bulk = bulk.filter(location__plant_id=plant_id)
        if location_id:
            bulk = bulk.filter(location_id=location_id)
        if date_from:
            bulk = bulk.filter(created_at__date__gte=date_from)
        if date_to:
            bulk = bulk.filter(created_at__date__lte=date_to)
        for tx in bulk:
            if str(tx.reference or "").startswith(cls.AUDIT_REFERENCE_PREFIXES):
                continue
            qty = _dec(tx.qty_kg)
            entries.append(cls._stock_card_entry(tx.created_at, f"BULK_{tx.type}", tx.reference or str(tx.id), tx.material, tx.location, qty, "KG", tx.avg_cost, {"transaction_id": str(tx.id), "source_doc": "bulk_transaction"}))

        packaging = PackagingTransaction.objects.select_related("material", "location").all()
        if material_id:
            packaging = packaging.filter(material_id=material_id)
        if plant_id:
            packaging = packaging.filter(location__plant_id=plant_id)
        if location_id:
            packaging = packaging.filter(location_id=location_id)
        if date_from:
            packaging = packaging.filter(created_at__date__gte=date_from)
        if date_to:
            packaging = packaging.filter(created_at__date__lte=date_to)
        for tx in packaging:
            if str(tx.reference or "").startswith(cls.AUDIT_REFERENCE_PREFIXES):
                continue
            qty = _dec(tx.qty)
            entries.append(cls._stock_card_entry(tx.created_at, f"PACKAGING_{tx.type}", tx.reference or str(tx.id), tx.material, tx.location, qty, getattr(tx.material, "base_uom", ""), tx.avg_cost, {"transaction_id": str(tx.id), "source_doc": "packaging_transaction"}))

        rolls = RollMovement.objects.select_related("roll__material", "from_location", "to_location").all()
        if material_id:
            rolls = rolls.filter(roll__material_id=material_id)
        if location_id:
            rolls = rolls.filter(models.Q(from_location_id=location_id) | models.Q(to_location_id=location_id))
        elif plant_id:
            rolls = rolls.filter(models.Q(from_location__plant_id=plant_id) | models.Q(to_location__plant_id=plant_id))
        if date_from:
            rolls = rolls.filter(timestamp__date__gte=date_from)
        if date_to:
            rolls = rolls.filter(timestamp__date__lte=date_to)
        for mv in rolls:
            if str(mv.reason_note or "").startswith(cls.AUDIT_REFERENCE_PREFIXES):
                continue
            qty = Decimal("0")
            weight = _dec(getattr(mv.roll, "weight_kg", 0))
            if location_id:
                if str(mv.to_location_id or "") == str(location_id) and str(mv.from_location_id or "") != str(location_id):
                    qty = weight
                elif str(mv.from_location_id or "") == str(location_id) and str(mv.to_location_id or "") != str(location_id):
                    qty = -weight
            elif plant_id:
                from_same = bool(mv.from_location_id and str(getattr(mv.from_location, "plant_id", "")) == str(plant_id))
                to_same = bool(mv.to_location_id and str(getattr(mv.to_location, "plant_id", "")) == str(plant_id))
                if to_same and not from_same:
                    qty = weight
                elif from_same and not to_same:
                    qty = -weight
            if qty == 0:
                continue
            entries.append(
                cls._stock_card_entry(
                    mv.timestamp,
                    f"ROLL_{mv.reason}",
                    mv.reason_note or str(mv.id),
                    mv.roll.material if mv.roll_id else None,
                    mv.to_location if qty > 0 else mv.from_location,
                    qty,
                    "KG",
                    None,
                    {"movement_id": str(mv.id), "roll_label": getattr(mv.roll, "label_id", ""), "source_doc": "roll_movement"},
                )
            )

        entries.sort(key=lambda row: (row["at"], row["reference"], row["source"]))
        rows = []
        balance = Decimal("0")
        balance_rate = Decimal("0")
        opening_total = Decimal("0")
        opening_value = Decimal("0")
        movement_total = Decimal("0")
        movement_value = Decimal("0")
        for entry in entries:
            qty = q4(entry["qty"])
            entry_rate = _dec(entry.get("rate")) if entry.get("rate") not in (None, "") else balance_rate
            try:
                balance, balance_rate = apply_wac(
                    balance_qty=balance,
                    balance_rate=balance_rate,
                    delta_qty=qty,
                    posting_rate=entry_rate,
                )
            except ValueError:
                balance = q4(balance + qty)
            transaction_value = signed_value(qty, entry_rate)
            if entry["source"] in {"OPENING_STOCK", "CUTOVER_OPENING", "OPENING_BALANCE_ADJUST"}:
                opening_total += qty
                opening_value += transaction_value
            else:
                movement_total += qty
                movement_value += transaction_value
            rows.append(cls._stock_card_row_from_entry(entry, balance, balance_rate, transaction_value))
        return {
            "opening_qty": float(opening_total),
            "movement_qty": float(movement_total),
            "closing_qty": float(balance),
            "opening_value": float(opening_value),
            "movement_value": float(movement_value),
            "closing_rate": float(balance_rate),
            "closing_value": float(signed_value(balance, balance_rate)),
            "rows": rows,
        }

    @classmethod
    def build_batch_workbook(cls, *, batch: InventoryAuditBatch):
        workbook = cls._base_workbook()
        sheet = workbook.active
        sheet.title = "Audit Register"
        metadata = [
            ["Batch No", batch.batch_no],
            ["Type", batch.type],
            ["Plant", getattr(batch.plant, "name", "")],
            ["Financial Year", batch.financial_year],
            ["Cutoff", timezone.localtime(batch.cutoff_at).strftime("%d-%b-%Y %H:%M") if batch.cutoff_at else ""],
            ["Status", batch.status],
            ["Notes", batch.notes or ""],
        ]
        for row in metadata:
            sheet.append(row)
        sheet.append([])
        headers = [
            "stock_class",
            "material_code",
            "material_name",
            "location_name",
            "uom",
            "system_qty",
            "opening_qty",
            "counted_qty",
            "variance_qty",
            "rate",
            "value",
            "granule_code",
            "grade",
            "label_id",
            "batch_no",
            "width_mm",
            "thickness_micron",
            "length_m",
            "is_fg",
            "stage_index",
            "status",
            "row_errors",
        ]
        sheet.append(headers)
        cls._style_header_row(sheet[sheet.max_row])
        for line in batch.lines.select_related("material", "location", "granule_code", "grade").all():
            sheet.append(
                [
                    line.stock_class,
                    getattr(line.material, "code", ""),
                    getattr(line.material, "name", ""),
                    getattr(line.location, "name", ""),
                    line.uom,
                    float(line.system_qty or 0),
                    float(line.opening_qty or 0),
                    float(line.counted_qty or 0) if line.counted_qty is not None else "",
                    float(line.variance_qty or 0),
                    float(line.rate or 0) if line.rate is not None else "",
                    float(line.value or 0),
                    getattr(line.granule_code, "code", ""),
                    getattr(line.grade, "name", ""),
                    line.label_id,
                    line.batch_no,
                    float(line.width_mm or 0) if line.width_mm is not None else "",
                    float(line.thickness_micron or 0) if line.thickness_micron is not None else "",
                    float(line.length_m or 0) if line.length_m is not None else "",
                    "YES" if line.is_fg else "NO",
                    line.stage_index,
                    line.status,
                    "; ".join(line.row_errors or []),
                ]
            )
        cls._autosize_columns(sheet)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue(), f"{batch.batch_no.lower()}-register.xlsx"

    @classmethod
    def build_sample_template(cls, *, batch_type: str, stock_class: Optional[str] = None):
        workbook = cls._base_workbook()
        sheet = workbook.active
        sheet.title = "Sample"
        headers = [
            "stock_class",
            "material",
            "location",
            "uom",
            "quantity",
            "counted_qty",
            "rate",
            "granule_code",
            "grade",
            "label_id",
            "batch_no",
            "width_mm",
            "thickness_micron",
            "length_m",
            "is_fg",
            "stage_index",
            "status",
        ]
        sheet.append(headers)
        cls._style_header_row(sheet[1])
        stock_class = str(stock_class or "BULK").upper()
        sample_rows = {
            "BULK": [stock_class, "MATERIAL_UUID", "LOCATION_UUID", "KG", "125.5", "125.5" if batch_type != "OPENING_STOCK" else "", "82.25", "GRANULE_CODE_UUID", "", "", "", "", "", "", "", "", "AVAILABLE"],
            "ROLL": [stock_class, "FILM_VARIANT_UUID", "LOCATION_UUID", "KG", "50", "50" if batch_type != "OPENING_STOCK" else "", "", "", "GRADE_UUID", "OPEN-ROLL-001", "LOT-001", "500", "50", "1200", "YES", "5", "AVAILABLE"],
            "PACKAGING": [stock_class, "PACKAGING_UUID", "LOCATION_UUID", "PCS", "2500", "2500" if batch_type != "OPENING_STOCK" else "", "1.25", "", "", "", "", "", "", "", "", "", "AVAILABLE"],
        }
        sheet.append(sample_rows.get(stock_class, sample_rows["BULK"]))
        notes = workbook.create_sheet("Read Me")
        notes.append(["Field", "Meaning"])
        cls._style_header_row(notes[1])
        for field, meaning in [
            ("material", "Use the UUID from the master record; the picker in the UI still helps for manual entry."),
            ("location", "Use the UUID of the plant location selected for this audit batch."),
            ("quantity", "Opening quantity for OPENING_STOCK imports."),
            ("counted_qty", "Physical quantity for PHYSICAL_COUNT and FY_CORRECTION imports."),
            ("granule_code", "Only for granule bulk stock when code-level tracking is needed."),
            ("grade", "Only for rolls that require grade-level trace."),
            ("width_mm / thickness_micron", "Required for roll imports."),
        ]:
            notes.append([field, meaning])
        cls._autosize_columns(sheet)
        cls._autosize_columns(notes)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue(), f"{str(batch_type).lower()}-{stock_class.lower()}-sample.xlsx"

    @classmethod
    def build_closing_preview_workbook(cls, *, preview: Dict[str, Any], financial_year: str):
        workbook = cls._base_workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["Financial Year", financial_year])
        summary.append(["Plant", preview.get("plant", {}).get("name", "") if preview.get("plant") else ""])
        summary.append([])
        summary.append(["Metric", "Value"])
        cls._style_header_row(summary[4])
        for key, value in (preview.get("totals") or {}).items():
            summary.append([key, value])
        movements = workbook.create_sheet("Movements")
        movements.append(["Metric", "Value"])
        cls._style_header_row(movements[1])
        for key, value in (preview.get("movements") or {}).items():
            movements.append([key, value])
        blockers = workbook.create_sheet("Blockers")
        blockers.append(["Code", "Label", "Count"])
        cls._style_header_row(blockers[1])
        for blocker in preview.get("blockers") or []:
            blockers.append([blocker.get("code"), blocker.get("label"), blocker.get("count")])
        stock = workbook.create_sheet("Closing Stock")
        headers = ["stock_class", "material_code", "material_name", "location_name", "qty", "uom", "grade_name", "granule_code_label", "label_id", "status"]
        stock.append(headers)
        cls._style_header_row(stock[1])
        for row in preview.get("rows") or []:
            stock.append([row.get("stock_class"), row.get("material_code"), row.get("material_name"), row.get("location_name"), row.get("qty"), row.get("uom"), row.get("grade_name"), row.get("granule_code_label"), row.get("label_id"), row.get("status")])
        for ws in workbook.worksheets:
            cls._autosize_columns(ws)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue(), f"fy-close-{financial_year}.xlsx"

    @classmethod
    def build_stock_card_workbook(cls, *, payload: Dict[str, Any]):
        workbook = cls._base_workbook()
        summary = workbook.active
        summary.title = "Ledger Summary"
        summary.append(["Opening Qty", payload.get("opening_qty", 0)])
        summary.append(["Movement Qty", payload.get("movement_qty", 0)])
        summary.append(["Closing Qty", payload.get("closing_qty", 0)])
        rows_sheet = workbook.create_sheet("Stock Card")
        headers = ["at", "source", "reference", "material_code", "material_name", "location_name", "in_qty", "out_qty", "balance_qty", "tx_rate", "balance_rate", "balance_value", "uom"]
        rows_sheet.append(headers)
        cls._style_header_row(rows_sheet[1])
        for row in payload.get("rows") or []:
            rows_sheet.append([row.get("at"), row.get("source"), row.get("reference"), row.get("material_code"), row.get("material_name"), row.get("location_name"), row.get("in_qty"), row.get("out_qty"), row.get("balance_qty"), row.get("rate"), row.get("balance_rate"), row.get("value"), row.get("uom")])
        cls._autosize_columns(summary)
        cls._autosize_columns(rows_sheet)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue(), "stock-card.xlsx"

    @classmethod
    def _emit_batch_event(cls, *, batch: InventoryAuditBatch, user=None, action: str, before: Optional[Dict[str, Any]], after: Dict[str, Any]) -> None:
        try:
            PermissionAuditLog.objects.create(
                user=user if getattr(user, "is_authenticated", False) else None,
                action=action[:40],
                method="POST",
                path=f"/api/inventory/audit/batches/{batch.id}/",
                required_permission="inventory.audit.manage",
                effective_role=getattr(getattr(user, "role", None), "code", "") if user else "",
                details={
                    "batch_id": str(batch.id),
                    "batch_no": batch.batch_no,
                    "type": batch.type,
                    "financial_year": batch.financial_year,
                    "plant_id": str(batch.plant_id),
                    "before": before,
                    "after": after,
                },
            )
        except Exception:
            # Audit mirroring must never break the inventory posting transaction.
            pass

    @classmethod
    def _base_workbook(cls):
        if Workbook is None:
            raise ValidationError("Excel export is unavailable because openpyxl is not installed.")
        return Workbook()

    @classmethod
    def _style_header_row(cls, cells):
        if not (Font and PatternFill and Alignment):
            return
        for cell in cells:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="0F172A")
            cell.alignment = Alignment(horizontal="center")

    @classmethod
    def _autosize_columns(cls, worksheet):
        for column in worksheet.columns:
            values = [str(cell.value or "") for cell in column[: min(len(column), 40)]]
            width = max((len(value) for value in values), default=10) + 2
            worksheet.column_dimensions[column[0].column_letter].width = min(max(width, 12), 36)

    @classmethod
    def _stock_card_entry(cls, at, source: str, reference: str, material, location, qty: Decimal, uom: str, rate: Any, meta: Dict[str, Any]) -> Dict[str, Any]:
        return {
            "at": at,
            "source": source,
            "reference": reference,
            "material_id": str(material.id) if material else None,
            "material_code": getattr(material, "code", ""),
            "material_name": getattr(material, "name", ""),
            "location_id": str(location.id) if location else None,
            "location_name": getattr(location, "name", ""),
            "qty": float(_dec(qty)),
            "uom": uom,
            "rate": float(_dec(rate)) if rate not in (None, "") else None,
            "meta": meta,
        }

    @classmethod
    def _stock_card_row_from_entry(
        cls,
        entry: Dict[str, Any],
        balance: Decimal,
        balance_rate: Decimal,
        transaction_value: Decimal,
    ) -> Dict[str, Any]:
        qty = _dec(entry.get("qty"))
        rate = entry.get("rate")
        rate_dec = _dec(rate) if rate not in (None, "") else None
        value = signed_value(balance, balance_rate)
        return {
            **entry,
            "at": entry["at"].isoformat() if hasattr(entry["at"], "isoformat") else str(entry["at"]),
            "qty": float(qty),
            "in_qty": float(qty) if qty > 0 else 0,
            "out_qty": float(abs(qty)) if qty < 0 else 0,
            "balance_qty": float(balance),
            "rate": float(rate_dec) if rate_dec is not None else None,
            "transaction_value": float(transaction_value),
            "balance_rate": float(balance_rate),
            "value": float(value),
            "display_value": float(display_value(balance, balance_rate)),
        }
