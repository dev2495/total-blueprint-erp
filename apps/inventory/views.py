from rest_framework import viewsets, generics, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.http import FileResponse, HttpResponse
from django.views.decorators.clickjacking import xframe_options_exempt
from django.utils.dateparse import parse_date
from django.utils import timezone
from decimal import Decimal
from datetime import date, datetime, timedelta
from io import BytesIO
import uuid
import csv
from collections import defaultdict

from .models import (
    InventoryLocation,
    JobWorkOrder,
    DeliveryChallan,
    InventoryRoll,
    RollMovement,
    RollConsumption,
    InventoryBulk,
    BulkTransaction,
    PackagingStock,
    PackagingTransaction,
    InventoryReservation,
    InventorySavedView,
    Vendor,
)
from .serializers import (
    InventoryLocationSerializer, JobWorkOrderSerializer,
    BulkGRNSerializer, RollGRNSerializer, PackagingGRNSerializer, JobWorkDispatchSerializer, JobWorkReceiveSerializer,
    DeliveryChallanSerializer, ChallanDispatchSerializer, ChallanReceiveSerializer,
    InventoryRollSerializer, RollDetailSerializer, RollMovementSerializer, RollConsumptionSerializer,
    RollReserveSerializer, RollReleaseSerializer, RollMoveSerializer, RollConsumeSerializer,
    InventoryBulkSerializer, BulkTransactionSerializer,
    PackagingStockSerializer, PackagingTransactionSerializer,
    resolve_roll_role, resolve_roll_stage_name,
)
from .services.grn import GRNService
from .services.grn_history import GRNHistoryService
from .services.stock import StockService
from .services.job_work import JobWorkService
from .services.inter_plant import InterPlantService
from .services.challan_pdf import ChallanPDFService
from .services.roll_service import RollService
from .services.stock_form_conversion import StockFormConversionService
from .services.bulk_service import BulkService
from .services.packaging_service import PackagingService
from .services.audit import InventoryAuditService as StockLifecycleService, current_indian_financial_year
from .services.roll_naming import (
    build_roll_naming_payload,
    build_variant_key,
    stock_strategy_label,
)
from apps.materials.models import InventoryMaterial
from apps.production.models import ProductionJob
from apps.factory.models import Process, Machine, Plant
from django_filters.rest_framework import DjangoFilterBackend

try:
    from apps.costing.models import MaterialCostSnapshot
except Exception:  # pragma: no cover - costing can be disabled in isolated settings
    MaterialCostSnapshot = None


def _inventory_roll_base_queryset():
    return InventoryRoll.objects.select_related(
        'material',
        'material__parent_family',
        'material__commercial_family',
        'material__parent_family__commercial_family',
        'location',
        'location__plant',
        'grade',
        'plant',
        'template',
        'template__commercial_family',
        'created_process',
        'created_by_job',
        'created_by_job__mts_order',
        'production_job',
        'parent_roll__created_process',
    )


def _roll_query_rows(params):
    plant_id = params.get('plant')
    stage_filter = (params.get('stage') or '').strip().lower()
    role_filter = (params.get('roll_role') or '').strip().upper()
    origin_type_filter = (params.get('origin_type') or '').strip().upper()
    stock_strategy_filter = (params.get('stock_strategy') or '').strip().upper()
    status_filter = params.get('status')
    material_id = params.get('material')
    grade_id = params.get('grade')
    location_id = params.get('location')
    family_id = params.get('family')
    job_number = (params.get('job_number') or '').strip()
    date_from = parse_date(params.get('date_from')) if params.get('date_from') else None
    date_to = parse_date(params.get('date_to')) if params.get('date_to') else None
    weight_min = Decimal(str(params.get('weight_min'))) if params.get('weight_min') else None
    weight_max = Decimal(str(params.get('weight_max'))) if params.get('weight_max') else None

    qs = _inventory_roll_base_queryset()
    if plant_id:
        qs = qs.filter(location__plant_id=plant_id)
    if material_id:
        qs = qs.filter(material_id=material_id)
    if grade_id:
        qs = qs.filter(grade_id=grade_id)
    if location_id:
        qs = qs.filter(location_id=location_id)
    if family_id:
        qs = qs.filter(
            Q(template__commercial_family_id=family_id)
            | Q(material__commercial_family_id=family_id)
            | Q(material__parent_family__commercial_family_id=family_id)
        )
    if status_filter:
        statuses = [s.strip().upper() for s in status_filter.split(',') if s.strip()]
        if statuses:
            qs = qs.filter(status__in=statuses)
    if job_number:
        qs = qs.filter(
            Q(created_by_job__job_number__icontains=job_number) |
            Q(production_job__job_number__icontains=job_number)
        )
    if date_from:
        qs = qs.filter(created_at__date__gte=date_from)
    if date_to:
        qs = qs.filter(created_at__date__lte=date_to)
    if weight_min is not None:
        qs = qs.filter(weight_kg__gte=weight_min)
    if weight_max is not None:
        qs = qs.filter(weight_kg__lte=weight_max)

    rows = []
    for roll in qs.order_by('-created_at')[:2000]:
        role = resolve_roll_role(roll)
        stage_name = resolve_roll_stage_name(roll) or 'Raw Material'
        bucket_label = 'Remainder / Freed' if role == 'REMAINDER' else stage_name
        parent = getattr(roll, 'parent_roll', None)
        source_stage = resolve_roll_stage_name(parent) if parent else None
        naming = build_roll_naming_payload(roll, role=role, stage_name=stage_name)
        row = {
            'id': str(roll.id),
            'label_id': roll.label_id,
            'roll_role': role,
            'origin_type': naming['origin_type'],
            'origin_label': naming['origin_label'],
            'stock_strategy': naming['stock_strategy'],
            'stock_strategy_label': naming['stock_strategy_label'],
            'family_display_name': naming['family_display_name'],
            'display_name': naming['variant_display_name'],
            'variant_display_name': naming['variant_display_name'],
            'size_line': naming['size_line'],
            'form_label': naming['form_label'],
            'process_state_label': naming['process_state_label'],
            'availability_label': naming['availability_label'],
            'print_status': naming['print_status'],
            'lamination_status': naming['lamination_status'],
            'reporting_group': naming['reporting_group'],
            'variant_summary': " • ".join(
                part for part in [
                    naming['size_line'],
                    stage_name,
                    roll.location.name if roll.location else '',
                ] if str(part or '').strip()
            ),
            'consumability_mode': (
                'Direct consume'
                if naming['stock_strategy'] == 'FINAL_STOCK'
                else ('Packaging only' if naming['stock_strategy'] == 'PACKAGING_STOCK' else 'Continue from WIP')
            ),
            'is_quarantined': bool((roll.meta_json or {}).get('is_quarantined')),
            'status': roll.status,
            'stage_name': stage_name,
            'bucket_label': bucket_label,
            'source_stage_name': source_stage,
            'material_id': str(roll.material_id) if roll.material_id else None,
            'material_name': roll.material.name if roll.material else None,
            'grade_id': str(roll.grade_id) if roll.grade_id else None,
            'grade_name': roll.grade.name if roll.grade else None,
            'weight_kg': float(roll.weight_kg or 0),
            'width_mm': float(roll.width_mm or 0),
            'stock_form': roll.stock_form,
            'width_basis': roll.width_basis,
            'thickness_micron': float(roll.thickness_micron or 0),
            'location_id': str(roll.location_id) if roll.location_id else None,
            'location_name': roll.location.name if roll.location else None,
            'plant_id': str(roll.location.plant_id) if roll.location_id and roll.location and roll.location.plant_id else (str(roll.plant_id) if roll.plant_id else None),
            'plant_name': roll.location.plant.name if roll.location_id and roll.location and roll.location.plant else (roll.plant.name if roll.plant else None),
            'created_job_id': str(roll.created_by_job_id) if roll.created_by_job_id else None,
            'created_job_number': roll.created_by_job.job_number if roll.created_by_job else None,
            'production_job_id': str(roll.production_job_id) if roll.production_job_id else None,
            'production_job_number': roll.production_job.job_number if roll.production_job else None,
            'created_at': roll.created_at.isoformat() if roll.created_at else None,
        }
        rows.append(row)

    if stage_filter:
        rows = [r for r in rows if r.get('bucket_label', '').lower() == stage_filter or r.get('stage_name', '').lower() == stage_filter]
    if role_filter:
        rows = [r for r in rows if str(r.get('roll_role') or '').upper() == role_filter]
    if origin_type_filter:
        rows = [r for r in rows if str(r.get('origin_type') or '').upper() == origin_type_filter]
    if stock_strategy_filter:
        rows = [r for r in rows if str(r.get('stock_strategy') or '').upper() == stock_strategy_filter]
    return rows


def _group_rows_by_bucket(rows):
    buckets = {}
    totals_rolls = 0
    totals_weight = Decimal('0')
    remainder_rolls = 0
    remainder_weight = Decimal('0')
    for row in rows:
        key = row['bucket_label']
        if key not in buckets:
            buckets[key] = {
                'key': key,
                'label': key,
                'roll_count': 0,
                'weight_kg': 0.0,
                'rolls': [],
            }
        buckets[key]['roll_count'] += 1
        buckets[key]['weight_kg'] = round(float(Decimal(str(buckets[key]['weight_kg'])) + Decimal(str(row['weight_kg']))), 3)
        buckets[key]['rolls'].append(row)
        totals_rolls += 1
        totals_weight += Decimal(str(row['weight_kg']))
        if row.get('roll_role') == 'REMAINDER':
            remainder_rolls += 1
            remainder_weight += Decimal(str(row['weight_kg']))

    bucket_rows = list(buckets.values())
    bucket_rows.sort(key=lambda b: (b['label'] != 'Remainder / Freed', b['label']))
    return {
        'totals': {
            'roll_count': totals_rolls,
            'weight_kg': round(float(totals_weight), 3),
            'remainder_roll_count': remainder_rolls,
            'remainder_weight_kg': round(float(remainder_weight), 3),
        },
        'buckets': bucket_rows,
    }


def _group_rows_by_variant(rows):
    family_map: dict[tuple, dict] = {}
    for row in rows:
        family_key = (
            str(row.get('family_display_name') or '').strip(),
            str(row.get('form_label') or '').strip(),
            str(row.get('reporting_group') or '').strip(),
        )
        family_bucket = family_map.setdefault(
            family_key,
            {
                'family_key': "|".join(family_key),
                'family_display_name': row.get('family_display_name') or 'Material',
                'form_label': row.get('form_label') or 'Roll',
                'reporting_group': row.get('reporting_group') or 'OTHER',
                'total_roll_count': 0,
                'total_available_kg': 0.0,
                'total_reserved_kg': 0.0,
                'total_blocked_kg': 0.0,
                'oldest_age_days': 0,
                'variants': {},
            },
        )
        family_bucket['total_roll_count'] += 1
        if str(row.get('status') or '').upper() == 'AVAILABLE':
            family_bucket['total_available_kg'] += float(row.get('weight_kg') or 0)
        elif str(row.get('status') or '').upper() == 'RESERVED':
            family_bucket['total_reserved_kg'] += float(row.get('weight_kg') or 0)
        else:
            family_bucket['total_blocked_kg'] += float(row.get('weight_kg') or 0)
        age_days = 0
        try:
            from django.utils.dateparse import parse_datetime
            created_at = parse_datetime(str(row.get('created_at') or '')) if row.get('created_at') else None
            if created_at:
                age_days = max((timezone.now().date() - created_at.date()).days, 0)
        except Exception:
            age_days = 0
        family_bucket['oldest_age_days'] = max(family_bucket['oldest_age_days'], age_days)

        variant_payload = dict(row)
        variant_payload['blocked_kg'] = float(row.get('weight_kg') or 0) if str(row.get('status') or '').upper() not in {'AVAILABLE', 'RESERVED'} else 0.0
        variant_key = build_variant_key(variant_payload)
        variant_bucket = family_bucket['variants'].setdefault(
            variant_key,
            {
                'variant_key': "|".join(str(bit) for bit in variant_key),
                'variant_display_name': row.get('variant_display_name') or row.get('display_name') or 'Variant',
                'size_line': row.get('size_line') or '',
                'form_label': row.get('form_label') or 'Roll',
                'stage_name': row.get('stage_name') or '',
                'print_status': row.get('print_status') or '',
                'lamination_status': row.get('lamination_status') or '',
                'stock_strategy': row.get('stock_strategy') or '',
                'stock_strategy_label': row.get('stock_strategy_label') or stock_strategy_label(row.get('stock_strategy')),
                'roll_count': 0,
                'available_kg': 0.0,
                'reserved_kg': 0.0,
                'blocked_kg': 0.0,
                'oldest_age_days': 0,
                'plants': defaultdict(lambda: {'plant_name': '', 'locations': set(), 'available_kg': 0.0, 'reserved_kg': 0.0, 'blocked_kg': 0.0}),
                'rolls': [],
            },
        )
        variant_bucket['roll_count'] += 1
        weight = float(row.get('weight_kg') or 0)
        status = str(row.get('status') or '').upper()
        plant_key = str(row.get('plant_id') or row.get('plant_name') or '')
        plant_bucket = variant_bucket['plants'][plant_key]
        plant_bucket['plant_name'] = row.get('plant_name') or '-'
        if row.get('location_name'):
            plant_bucket['locations'].add(row.get('location_name'))
        if status == 'AVAILABLE':
            variant_bucket['available_kg'] += weight
            plant_bucket['available_kg'] += weight
        elif status == 'RESERVED':
            variant_bucket['reserved_kg'] += weight
            plant_bucket['reserved_kg'] += weight
        else:
            variant_bucket['blocked_kg'] += weight
            plant_bucket['blocked_kg'] += weight
        variant_bucket['oldest_age_days'] = max(variant_bucket['oldest_age_days'], age_days)
        variant_bucket['rolls'].append(row)

    families = []
    for family_bucket in family_map.values():
        variants = []
        for variant_bucket in family_bucket['variants'].values():
            variant_bucket['plant_summary'] = [
                {
                    'plant_name': plant['plant_name'],
                    'locations': sorted(location for location in plant['locations'] if location),
                    'available_kg': round(float(plant['available_kg']), 3),
                    'reserved_kg': round(float(plant['reserved_kg']), 3),
                    'blocked_kg': round(float(plant['blocked_kg']), 3),
                }
                for plant in variant_bucket['plants'].values()
            ]
            variant_bucket.pop('plants', None)
            variants.append(variant_bucket)
        variants.sort(key=lambda row: (row['variant_display_name'], row['stage_name'], row['stock_strategy']))
        family_bucket['variants'] = variants
        families.append(family_bucket)
    families.sort(key=lambda row: (row['family_display_name'], row['reporting_group']))
    return families


def _api_error_message(exc):
    detail = getattr(exc, "message_dict", None) or getattr(exc, "messages", None) or str(exc)
    if isinstance(detail, (list, tuple)):
        return " ".join(str(part) for part in detail)
    if isinstance(detail, dict):
        return detail
    return str(detail)


def _as_decimal(value, default="0"):
    if value in (None, ""):
        return Decimal(default)
    return Decimal(str(value))


def _latest_material_rate(material):
    if not material:
        return Decimal("0"), "ZERO"
    if MaterialCostSnapshot is not None:
        snapshot = MaterialCostSnapshot.objects.filter(material=material).order_by("-effective_date").first()
        if snapshot and _as_decimal(snapshot.avg_rate_per_kg) > 0:
            return _as_decimal(snapshot.avg_rate_per_kg), "MATERIAL_COST_SNAPSHOT"
    for attr in ("standard_cost", "avg_cost", "cost_price"):
        value = getattr(material, attr, None)
        if value not in (None, "") and _as_decimal(value) > 0:
            return _as_decimal(value), "MATERIAL_STANDARD"
    return Decimal("0"), "ZERO"


def _roll_rate(roll):
    meta = roll.meta_json or {}
    for key in ("unit_cost_per_kg", "unit_cost", "rate_per_kg", "rate_per_uom"):
        value = meta.get(key)
        if value not in (None, "") and _as_decimal(value) > 0:
            return _as_decimal(value), "ROLL_GRN_META"
    return _latest_material_rate(roll.material)


def _resolve_material_from_payload(payload):
    material_id = payload.get("material_id") or payload.get("product_id") or payload.get("variant_id")
    if material_id:
        return InventoryMaterial.objects.get(id=material_id)
    code = str(payload.get("material_code") or payload.get("product_code") or payload.get("variant_code") or "").strip()
    name = str(payload.get("material_name") or payload.get("product_name") or payload.get("variant_name") or "").strip()
    if not code and not name:
        raise ValidationError("Every GRN line needs material_id/product_id/variant_id, material_code, or exact material_name.")
    material = InventoryMaterial.objects.filter(code__iexact=code).first() if code else None
    if not material and name:
        material = InventoryMaterial.objects.filter(name__iexact=name).first()
    if not material:
        ref = code or name
        raise ValidationError(f"Material {ref} was not found.")
    return material


RECEIPT_UOM_ALIASES = {
    "KG": "KG",
    "KGS": "KG",
    "KILOGRAM": "KG",
    "KILOGRAMS": "KG",
    "PCS": "PCS",
    "PC": "PCS",
    "PIECE": "PCS",
    "PIECES": "PCS",
    "METER": "METER",
    "METERS": "METER",
    "METRE": "METER",
    "METRES": "METER",
    "M": "METER",
    "MTR": "METER",
    "MTRS": "METER",
    "MTS": "METER",
}
VALID_RECEIPT_UOMS = {"KG", "PCS", "METER"}


def _normalize_receipt_uom(value):
    compact = "".join(ch for ch in str(value or "").strip().upper() if ch.isalnum())
    return RECEIPT_UOM_ALIASES.get(compact, compact)


def _material_receipt_uom(material):
    category = str(getattr(material, "category", "") or "").upper()
    raw_uom = getattr(material, "addon_purchase_uom", None) if category == "ADDON" else None
    expected = _normalize_receipt_uom(raw_uom or getattr(material, "base_uom", None))
    if expected not in VALID_RECEIPT_UOMS:
        code = getattr(material, "code", material)
        raise ValidationError(
            f"Material {code} has unsupported master UOM '{getattr(material, 'base_uom', '')}'. "
            "GRN supports KG, PCS, and METER only."
        )
    return expected


def _validate_receipt_line_uom(line, material):
    expected = _material_receipt_uom(material)
    incoming = _normalize_receipt_uom(line.get("uom") or expected)
    if incoming != expected:
        code = getattr(material, "code", material)
        raise ValidationError(f"Material {code} must be received in its master UOM {expected}; got {incoming or 'blank'}.")
    return expected


def _resolve_location_from_payload(payload, fallback_id=None):
    location_id = payload.get("location_id") or payload.get("location") or payload.get("store_location_id") or fallback_id
    if not location_id:
        raise ValidationError("Receiving location is required.")
    return InventoryLocation.objects.select_related("plant").get(id=location_id)


def _normalise_upload_key(value):
    return "".join(ch for ch in str(value or "").strip().lower() if ch.isalnum())


def _upload_cell(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def _upload_decimal(value, *, default=None):
    if value in (None, ""):
        if default is None:
            return None
        return Decimal(str(default))
    return Decimal(str(value))


def _resolve_upload_vendor(ref):
    ref = str(ref or "").strip()
    if not ref:
        raise ValidationError("GRN upload needs Vendor code in the GRN Header sheet.")
    qs = Vendor.objects.filter(Q(code__iexact=ref) | Q(name__iexact=ref))
    if _looks_like_uuid(ref):
        qs = Vendor.objects.filter(Q(id=ref) | Q(code__iexact=ref) | Q(name__iexact=ref))
    vendor = qs.first()
    if not vendor:
        raise ValidationError(f"Vendor {ref} was not found.")
    return vendor


def _resolve_upload_location(ref, *, plant_ref=""):
    ref = str(ref or "").strip()
    if not ref:
        raise ValidationError("GRN upload needs Receiving warehouse in the GRN Header sheet or Location Code per row.")
    qs = InventoryLocation.objects.select_related("plant")
    if plant_ref:
        plant_ref = str(plant_ref).strip()
        plant = Plant.objects.filter(Q(code__iexact=plant_ref) | Q(name__iexact=plant_ref)).first()
        if _looks_like_uuid(plant_ref):
            plant = Plant.objects.filter(Q(id=plant_ref) | Q(code__iexact=plant_ref) | Q(name__iexact=plant_ref)).first()
        if not plant:
            raise ValidationError(f"Receiving plant {plant_ref} was not found.")
        qs = qs.filter(plant=plant)
    if _looks_like_uuid(ref):
        location = qs.filter(id=ref).first()
    else:
        location = qs.filter(Q(code__iexact=ref) | Q(name__iexact=ref)).first()
    if not location:
        raise ValidationError(f"Receiving location {ref} was not found.")
    return location


def _resolve_upload_grade(ref):
    ref = str(ref or "").strip()
    if not ref:
        return None
    from apps.recipes.models import RecipeGrade
    if _looks_like_uuid(ref):
        grade = RecipeGrade.objects.filter(id=ref, is_active=True).first()
    else:
        grade = RecipeGrade.objects.filter(name__iexact=ref, is_active=True).first()
    if not grade:
        raise ValidationError(f"Recipe grade {ref} was not found.")
    return grade


def _looks_like_uuid(value):
    try:
        uuid.UUID(str(value))
        return True
    except Exception:
        return False


def _active_reservation_weight_by_roll():
    reserved = {}
    for row in InventoryReservation.objects.filter(status="ACTIVE", roll_id__isnull=False).values("roll_id").annotate(total=Sum("quantity")):
        reserved[str(row["roll_id"])] = float(row["total"] or 0)
    return reserved


def _active_reservation_qty_by_material():
    reserved = {}
    for row in InventoryReservation.objects.filter(status="ACTIVE", roll_id__isnull=True).values("material_id").annotate(total=Sum("quantity")):
        reserved[str(row["material_id"])] = float(row["total"] or 0)
    return reserved


def _roll_age_days(roll):
    if not roll.created_at:
        return 0
    return max((timezone.now().date() - roll.created_at.date()).days, 0)


def _material_class(material):
    category = str(getattr(material, "category", "") or "").upper()
    if category in {"GRANULE", "FILM_FAMILY", "FILM_VARIANT"}:
        return "granule" if category == "GRANULE" else "film"
    if category in {"INK", "ADHESIVE", "SOLVENT", "CHEMICAL", "ADDON", "POD"}:
        return category.lower()
    if category == "PACKAGING":
        return "packaging"
    return category.lower() or "other"


def _facet_counts(items, keys):
    facets = {}
    for key in keys:
        counts = defaultdict(int)
        for item in items:
            value = item.get(key)
            if value in (None, ""):
                value = "Unassigned"
            counts[str(value)] += 1
        facets[key] = [
            {"value": value, "count": count}
            for value, count in sorted(counts.items(), key=lambda pair: (pair[0] == "Unassigned", pair[0]))
        ]
    return facets


def _paged_payload(items, request, *, facet_keys=()):
    try:
        limit = min(max(int(request.query_params.get("limit") or 250), 1), 1000)
    except Exception:
        limit = 250
    try:
        offset = max(int(request.query_params.get("cursor") or request.query_params.get("offset") or 0), 0)
    except Exception:
        offset = 0
    search = str(request.query_params.get("search") or request.query_params.get("q") or "").strip().lower()
    if search:
        items = [
            row for row in items
            if search in " ".join(str(value or "").lower() for value in row.values())
        ]
    next_offset = offset + limit
    return {
        "items": items[offset:next_offset],
        "total": len(items),
        "next_cursor": str(next_offset) if next_offset < len(items) else None,
        "facets": _facet_counts(items, facet_keys),
    }


def _bulk_reference_vendor(reference):
    text = str(reference or "")
    if not text.startswith("VENDOR:"):
        return {"vendor_code": "", "vendor_name": ""}
    label = text.split("|", 1)[0].replace("VENDOR:", "").strip()
    return {"vendor_code": label, "vendor_name": label}


def _stock_snapshot_payload(request):
    plant_id = request.query_params.get("plant_id") or request.query_params.get("plant")
    as_of = timezone.now()
    roll_reservations = _active_reservation_weight_by_roll()
    material_reservations = _active_reservation_qty_by_material()

    rolls_qs = _inventory_roll_base_queryset().exclude(status__in=["CONSUMED", "SCRAPPED", "MISSING"])
    bulk_qs = InventoryBulk.objects.select_related("material", "granule_code", "location", "plant").filter(qty_kg__gt=0)
    packaging_qs = PackagingStock.objects.select_related("material", "location", "plant").filter(qty__gt=0)
    if plant_id:
        rolls_qs = rolls_qs.filter(Q(location__plant_id=plant_id) | Q(plant_id=plant_id))
        bulk_qs = bulk_qs.filter(plant_id=plant_id)
        packaging_qs = packaging_qs.filter(plant_id=plant_id)

    roll_rows = []
    ageing_counts = {"0-30": 0, "31-60": 0, "61-90": 0, "90+": 0}
    total_roll_kg = Decimal("0")
    reservation_kg = Decimal("0")
    for roll in rolls_qs.order_by("-created_at")[:2000]:
        kg = _as_decimal(roll.net_weight_kg if roll.net_weight_kg is not None else roll.weight_kg)
        rate, rate_source = _roll_rate(roll)
        reserved_kg = Decimal(str(roll_reservations.get(str(roll.id), 0)))
        if roll.status == "RESERVED" and reserved_kg <= 0:
            reserved_kg = kg
        age_days = _roll_age_days(roll)
        if age_days <= 30:
            ageing_counts["0-30"] += 1
        elif age_days <= 60:
            ageing_counts["31-60"] += 1
        elif age_days <= 90:
            ageing_counts["61-90"] += 1
        else:
            ageing_counts["90+"] += 1
        total_roll_kg += kg
        reservation_kg += min(reserved_kg, kg)
        roll_rows.append({
            "id": str(roll.id),
            "label": roll.label_id,
            "label_id": roll.label_id,
            "product_name": roll.material.name if roll.material else "",
            "material_name": roll.material.name if roll.material else "",
            "material_code": roll.material.code if roll.material else "",
            "material_category": roll.material.category if roll.material else "",
            "category": roll.material.category if roll.material else "",
            "variant_code": roll.material.code if roll.material else "",
            "thickness_um": float(roll.thickness_micron or 0),
            "thickness_micron": float(roll.thickness_micron or 0),
            "width_mm": float(roll.width_mm or 0),
            "stock_form": roll.stock_form,
            "width_basis": roll.width_basis,
            "length_m": float(roll.length_m or 0),
            "net_weight_kg": float(kg),
            "weight_kg": float(roll.weight_kg or 0),
            "location": str(roll.location_id),
            "location_code": roll.location.code if roll.location else "",
            "location_name": roll.location.name if roll.location else "",
            "reserved_qty": float(reserved_kg),
            "free_qty": float(max(Decimal("0"), kg - reserved_kg)),
            "reserved_for_so_id": str(roll.sales_order_item_id) if roll.sales_order_item_id else None,
            "age_days": age_days,
            "status": roll.status,
            "roll_role": resolve_roll_role(roll),
            "rate": float(rate),
            "avg_cost": float(rate),
            "rate_source": rate_source,
            "rate_missing": rate <= 0,
        })

    bulk_stocks = list(bulk_qs.order_by("material__code", "location__code")[:2000])
    latest_bulk_tx_by_key = {}
    if bulk_stocks:
        material_ids = {stock.material_id for stock in bulk_stocks}
        location_ids = {stock.location_id for stock in bulk_stocks}
        for tx in BulkTransaction.objects.select_related("vendor").filter(
            type="INWARD",
            material_id__in=material_ids,
            location_id__in=location_ids,
        ).order_by("-created_at"):
            key = (tx.material_id, tx.location_id, tx.granule_code_id)
            if key not in latest_bulk_tx_by_key:
                latest_bulk_tx_by_key[key] = tx

    bulk_rows = []
    total_bulk_kg = Decimal("0")
    for stock in bulk_stocks:
        qty = _as_decimal(stock.qty_kg)
        reserved = Decimal(str(material_reservations.get(str(stock.material_id), 0)))
        latest_tx = latest_bulk_tx_by_key.get((stock.material_id, stock.location_id, stock.granule_code_id))
        uom = stock.material.base_uom or "KG"
        reference_vendor = _bulk_reference_vendor(latest_tx.reference if latest_tx else "")
        vendor_code = latest_tx.vendor.code if latest_tx and latest_tx.vendor else reference_vendor.get("vendor_code", "")
        vendor_name = latest_tx.vendor.name if latest_tx and latest_tx.vendor else reference_vendor.get("vendor_name", "")
        total_bulk_kg += qty
        reservation_kg += min(reserved, qty)
        bulk_rows.append({
            "id": str(stock.id),
            "product_name": stock.material.name,
            "material": str(stock.material_id),
            "material_name": stock.material.name,
            "material_code": stock.material.code,
            "material_category": stock.material.category,
            "material_class": _material_class(stock.material),
            "category": stock.material.category,
            "is_addon": stock.material.category == "ADDON",
            "addon_is_purchased": bool(getattr(stock.material, "addon_is_purchased", False)),
            "addon_purchase_uom": getattr(stock.material, "addon_purchase_uom", "") or "",
            "purchased_addon": stock.material.category == "ADDON" and bool(getattr(stock.material, "addon_is_purchased", False)),
            "lot_no": stock.granule_code.code if stock.granule_code else "",
            "qty": float(qty),
            "qty_kg": float(qty),
            "quantity": float(qty),
            "uom": uom,
            "stock_uom": uom,
            "base_uom": uom,
            "reserved_qty": float(min(reserved, qty)),
            "free_qty": float(max(Decimal("0"), qty - reserved)),
            "plant": str(stock.plant_id) if stock.plant_id else "",
            "plant_name": stock.plant.name if stock.plant else "",
            "location": str(stock.location_id),
            "location_code": stock.location.code,
            "location_name": stock.location.name,
            "age_days": max((as_of.date() - stock.updated_at.date()).days, 0) if stock.updated_at else 0,
            "avg_cost": float(stock.avg_cost or 0),
            "vendor": str(latest_tx.vendor_id) if latest_tx and latest_tx.vendor_id else "",
            "vendor_code": vendor_code,
            "vendor_name": vendor_name,
            "vendor_invoice_no": latest_tx.vendor_invoice_no if latest_tx else "",
            "manual_po_ref": latest_tx.manual_po_ref if latest_tx else "",
            "last_grn_no": (latest_tx.vendor_invoice_no or latest_tx.reference or "") if latest_tx else "",
            "last_grn_ref": latest_tx.reference if latest_tx else "",
            "last_grn_at": latest_tx.created_at.isoformat() if latest_tx and latest_tx.created_at else None,
        })

    packaging_rows = []
    packaging_qty = Decimal("0")
    for stock in packaging_qs.order_by("material__code", "location__code")[:2000]:
        qty = _as_decimal(stock.qty)
        reserved = Decimal(str(material_reservations.get(str(stock.material_id), 0)))
        packaging_qty += qty
        packaging_rows.append({
            "id": str(stock.id),
            "product_name": stock.material.name,
            "material": str(stock.material_id),
            "material_name": stock.material.name,
            "material_code": stock.material.code,
            "code": stock.material.code,
            "name": stock.material.name,
            "packaging_kind": stock.material.packaging_kind,
            "material_class": _material_class(stock.material),
            "packaging_supply_mode": getattr(stock.material, "packaging_supply_mode", "") or "",
            "supply_mode": getattr(stock.material, "packaging_supply_mode", "") or "",
            "qty": float(qty),
            "uom": stock.material.base_uom or "PCS",
            "base_uom": stock.material.base_uom or "PCS",
            "reserved_qty": float(min(reserved, qty)),
            "free_qty": float(max(Decimal("0"), qty - reserved)),
            "plant": str(stock.plant_id) if stock.plant_id else "",
            "plant_name": stock.plant.name if stock.plant else "",
            "location": str(stock.location_id),
            "location_code": stock.location.code,
            "location_name": stock.location.name,
            "avg_cost": float(stock.avg_cost or 0),
        })

    total_rows = max(len(roll_rows), 1)
    ageing_buckets = {key: round((value / total_rows) * 100, 2) for key, value in ageing_counts.items()}
    total_kg = total_roll_kg + total_bulk_kg
    free_kg = max(Decimal("0"), total_kg - reservation_kg)
    total_value = sum(Decimal(str(row.get("qty_kg", 0))) * Decimal(str(row.get("avg_cost", 0))) for row in bulk_rows)
    total_value += sum(Decimal(str(row.get("net_weight_kg", row.get("weight_kg", 0)))) * Decimal(str(row.get("avg_cost", 0))) for row in roll_rows)
    total_value += sum(Decimal(str(row.get("qty", 0))) * Decimal(str(row.get("avg_cost", 0))) for row in packaging_rows)

    return {
        "as_of": as_of.isoformat(),
        "plant_id": str(plant_id) if plant_id else None,
        "kpi": {
            "total_value_inr": float(total_value),
            "total_kg": float(total_kg),
            "rolls_count": len(roll_rows),
            "bulk_lots": len(bulk_rows),
            "packaging_skus": len(packaging_rows),
            "packaging_qty": float(packaging_qty),
            "reservation_kg": float(reservation_kg),
            "free_kg": float(free_kg),
            "reservation_pct": float((reservation_kg / total_kg * 100) if total_kg > 0 else 0),
            "ageing_buckets": ageing_buckets,
        },
        "rolls": roll_rows,
        "bulk": bulk_rows,
        "packaging": packaging_rows,
    }


class InventoryV36SnapshotView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response(_stock_snapshot_payload(request))


class InventoryV36RollMatrixView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        snapshot = _stock_snapshot_payload(request)
        rows = {}
        for roll in snapshot["rolls"]:
            variant_id = roll.get("material_code") or "UNKNOWN"
            thickness = int(round(float(roll.get("thickness_micron") or 0)))
            if not thickness:
                continue
            bucket = rows.setdefault(variant_id, {"variant_id": variant_id, "variant_code": variant_id, "cells": {}})
            cell = bucket["cells"].setdefault(thickness, {"thickness_um": thickness, "kg": 0.0, "count": 0})
            cell["kg"] += float(roll.get("net_weight_kg") or roll.get("weight_kg") or 0)
            cell["count"] += 1
        payload = []
        for row in rows.values():
            row["cells"] = sorted(row["cells"].values(), key=lambda item: item["thickness_um"])
            payload.append(row)
        return Response({"rows": sorted(payload, key=lambda item: item["variant_code"])})


class InventoryV36AnomaliesView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        kind_filter = str(request.query_params.get("kind") or "").upper().strip()
        snapshot = _stock_snapshot_payload(request)
        items = []

        def add(kind, severity, ref_type, ref_id, label, message, suggested_action):
            if kind_filter and kind != kind_filter:
                return
            items.append({
                "kind": kind,
                "severity": severity,
                "ref_type": ref_type,
                "ref_id": ref_id,
                "label": label,
                "message": message,
                "suggested_action": suggested_action,
            })

        for roll in snapshot["rolls"]:
            if int(roll.get("age_days") or 0) > 90:
                add("AGEING", "warn", "ROLL", roll["id"], f"Roll {roll['label_id']}", f"Aged {roll['age_days']} days, exceeds 90-day target.", "Review for clearance, conversion priority, or write-off.")
            if float(roll.get("reserved_qty") or 0) > float(roll.get("net_weight_kg") or 0):
                add("OVER_RESERVED", "block", "ROLL", roll["id"], f"Roll {roll['label_id']}", "Reserved quantity is greater than available roll weight.", "Release or correct the sales/job reservation.")
        for row in snapshot["bulk"]:
            if float(row.get("qty_kg") or 0) < 0:
                add("NEGATIVE", "block", "BULK", row["id"], row["material_code"], "Bulk stock is negative.", "Post stock correction before period close.")
            if float(row.get("reserved_qty") or 0) > float(row.get("qty_kg") or 0):
                add("OVER_RESERVED", "block", "BULK", row["id"], row["material_code"], "Reserved quantity exceeds on-hand stock.", "Release stale reservations or inward stock.")
        for row in snapshot["packaging"]:
            if float(row.get("reserved_qty") or 0) > float(row.get("qty") or 0):
                add("OVER_RESERVED", "block", "PACKAGING", row["id"], row["material_code"], "Reserved packaging exceeds on-hand stock.", "Release stale reservations or inward packaging.")
        return Response({"items": items[:500]})


class InventoryV36ReservationsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        ref_id = request.query_params.get("ref_id")
        ref_type = str(request.query_params.get("ref_type") or "").upper().strip()
        qs = InventoryReservation.objects.select_related("job", "roll", "material").filter(status="ACTIVE")
        if ref_id and ref_type == "ROLL":
            qs = qs.filter(roll_id=ref_id)
        elif ref_id:
            if ref_type == "BULK":
                material_id = InventoryBulk.objects.filter(id=ref_id).values_list("material_id", flat=True).first() or ref_id
            elif ref_type == "PACKAGING":
                material_id = PackagingStock.objects.filter(id=ref_id).values_list("material_id", flat=True).first() or ref_id
            else:
                material_id = ref_id
            qs = qs.filter(material_id=material_id, roll_id__isnull=True)
        rows = []
        for reservation in qs.order_by("-created_at")[:500]:
            job = reservation.job
            rows.append({
                "so_id": str(getattr(job, "sales_order_id", "") or getattr(job, "sales_order_item_id", "") or ""),
                "so_no": getattr(getattr(job, "sales_order", None), "order_no", "") or getattr(job, "job_number", ""),
                "customer_name": getattr(getattr(getattr(job, "sales_order", None), "customer", None), "name", "") or "",
                "job_id": str(job.id) if job else "",
                "job_no": getattr(job, "job_number", "") if job else "",
                "qty": float(reservation.quantity or 0),
                "uom": reservation.uom,
                "reserved_at": reservation.created_at.isoformat() if reservation.created_at else None,
                "promise_date": str(getattr(job, "due_date", "") or ""),
            })
        return Response({"items": rows})


class InventoryV36SavedViewsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        workspace = str(request.query_params.get("workspace") or "").strip().lower()
        qs = InventorySavedView.objects.filter(Q(user=request.user) | Q(shared_team=True))
        if workspace:
            qs = qs.filter(workspace=workspace)
        return Response({"items": [self._serialize(row) for row in qs[:200]]})

    def post(self, request):
        workspace = str(request.data.get("workspace") or "").strip().lower()
        name = str(request.data.get("name") or "").strip()
        if workspace not in dict(InventorySavedView.WORKSPACE_CHOICES):
            return Response({"error": "workspace must be rolls, bulk, packaging, addons, or home."}, status=status.HTTP_400_BAD_REQUEST)
        if not name:
            return Response({"error": "name is required."}, status=status.HTTP_400_BAD_REQUEST)
        row = InventorySavedView.objects.create(
            workspace=workspace,
            name=name,
            icon=str(request.data.get("icon") or "")[:16],
            pinned=bool(request.data.get("pinned")),
            state=request.data.get("state") if isinstance(request.data.get("state"), dict) else {},
            shared_team=bool(request.data.get("shared_team")),
            user=request.user if request.user.is_authenticated else None,
        )
        return Response(self._serialize(row), status=status.HTTP_201_CREATED)

    @staticmethod
    def _serialize(row):
        return {
            "id": str(row.id),
            "workspace": row.workspace,
            "name": row.name,
            "icon": row.icon,
            "pinned": row.pinned,
            "state": row.state or {},
            "shared_team": row.shared_team,
            "created_at": row.created_at.isoformat() if row.created_at else None,
            "updated_at": row.updated_at.isoformat() if row.updated_at else None,
        }


class InventoryV36SavedViewDetail(APIView):
    permission_classes = [IsAuthenticated]

    def patch(self, request, pk):
        row = InventorySavedView.objects.filter(id=pk).filter(Q(user=request.user) | Q(shared_team=True)).first()
        if not row:
            return Response({"error": "Saved view was not found."}, status=status.HTTP_404_NOT_FOUND)
        for field in ("name", "icon"):
            if field in request.data:
                setattr(row, field, str(request.data.get(field) or ""))
        if "pinned" in request.data:
            row.pinned = bool(request.data.get("pinned"))
        if "state" in request.data and isinstance(request.data.get("state"), dict):
            row.state = request.data.get("state")
        row.save()
        return Response(InventoryV36SavedViewsView._serialize(row))

    def delete(self, request, pk):
        deleted, _ = InventorySavedView.objects.filter(id=pk, user=request.user).delete()
        if not deleted:
            return Response({"error": "Saved view was not found or is shared."}, status=status.HTTP_404_NOT_FOUND)
        return Response(status=status.HTTP_204_NO_CONTENT)


class InventoryV36ClassSnapshotView(APIView):
    permission_classes = [IsAuthenticated]
    klass = ""

    def get(self, request):
        snapshot = _stock_snapshot_payload(request)
        klass = self.klass or str(request.resolver_match.kwargs.get("klass") or "").lower()
        if klass == "rolls":
            items = snapshot["rolls"]
            facets = ("status", "roll_role", "plant_name", "location_name", "width_mm", "thickness_micron")
        elif klass == "bulk":
            items = snapshot["bulk"]
            facets = ("material_class", "material_category", "plant_name", "location_name")
        elif klass == "packaging":
            items = snapshot["packaging"]
            facets = ("packaging_kind", "supply_mode", "plant_name", "location_name")
        elif klass == "addons":
            items = [
                row for row in snapshot["bulk"]
                if str(row.get("material_category") or "").upper() in {"ADDON", "INK", "ADHESIVE", "SOLVENT", "CHEMICAL"}
            ]
            facets = ("material_class", "material_category", "plant_name", "location_name", "uom")
        else:
            return Response({"error": "Unknown inventory class."}, status=status.HTTP_400_BAD_REQUEST)
        return Response(_paged_payload(items, request, facet_keys=facets))


class InventoryV36CoverageView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        snapshot = _stock_snapshot_payload(request)
        rows = []
        for row in snapshot["bulk"]:
            free = float(row.get("free_qty") or 0)
            reserved = float(row.get("reserved_qty") or 0)
            rows.append({
                "stock_class": "BULK",
                "material_code": row.get("material_code"),
                "material_name": row.get("material_name"),
                "material_class": row.get("material_class"),
                "on_hand": row.get("qty"),
                "reserved": reserved,
                "free": free,
                "uom": row.get("uom") or "KG",
                "days_cover": None,
                "reorder_health": "blocked" if free <= 0 and reserved > 0 else ("low" if free <= 0 else "ok"),
                "suggested_action": "Inward stock or release reservations." if free <= 0 else "No action.",
            })
        for row in snapshot["packaging"]:
            free = float(row.get("free_qty") or 0)
            reserved = float(row.get("reserved_qty") or 0)
            rows.append({
                "stock_class": "PACKAGING",
                "material_code": row.get("material_code"),
                "material_name": row.get("material_name"),
                "material_class": row.get("material_class"),
                "on_hand": row.get("qty"),
                "reserved": reserved,
                "free": free,
                "uom": row.get("uom") or "PCS",
                "days_cover": None,
                "reorder_health": "blocked" if free <= 0 and reserved > 0 else ("low" if free <= 0 else "ok"),
                "suggested_action": "Inward packaging or release reservations." if free <= 0 else "No action.",
            })
        return Response({"items": rows[:1000]})


class InventoryV36SnapshotTrendView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        days = min(max(int(request.query_params.get("days") or 14), 1), 90)
        start_date = timezone.now().date() - timedelta(days=days - 1)
        buckets = {
            (start_date + timedelta(days=idx)).isoformat(): {
                "date": (start_date + timedelta(days=idx)).isoformat(),
                "bulk_in_kg": 0.0,
                "bulk_out_kg": 0.0,
                "packaging_in": 0.0,
                "packaging_out": 0.0,
                "rolls_created": 0,
            }
            for idx in range(days)
        }
        for tx in BulkTransaction.objects.filter(created_at__date__gte=start_date):
            key = tx.created_at.date().isoformat()
            qty = float(abs(tx.qty_kg or 0))
            if Decimal(str(tx.qty_kg or 0)) >= 0:
                buckets[key]["bulk_in_kg"] += qty
            else:
                buckets[key]["bulk_out_kg"] += qty
        for tx in PackagingTransaction.objects.filter(created_at__date__gte=start_date):
            key = tx.created_at.date().isoformat()
            qty = float(abs(tx.qty or 0))
            if Decimal(str(tx.qty or 0)) >= 0:
                buckets[key]["packaging_in"] += qty
            else:
                buckets[key]["packaging_out"] += qty
        for row in InventoryRoll.objects.filter(created_at__date__gte=start_date).values("created_at__date").annotate(count=Count("id")):
            key = row["created_at__date"].isoformat()
            buckets[key]["rolls_created"] = int(row["count"] or 0)
        return Response({"items": list(buckets.values())})


class InventoryV36ExportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, klass):
        snapshot = _stock_snapshot_payload(request)
        normalized = str(klass or "").lower()
        if normalized == "rolls":
            rows = snapshot["rolls"]
        elif normalized == "bulk":
            rows = snapshot["bulk"]
        elif normalized == "packaging":
            rows = snapshot["packaging"]
        elif normalized == "addons":
            rows = [
                row for row in snapshot["bulk"]
                if str(row.get("material_category") or "").upper() in {"ADDON", "INK", "ADHESIVE", "SOLVENT", "CHEMICAL"}
            ]
        else:
            return Response({"error": "Unknown export class."}, status=status.HTTP_400_BAD_REQUEST)
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="inventory-{normalized}-{timezone.now().date().isoformat()}.csv"'
        writer = csv.writer(response)
        columns = sorted({key for row in rows for key in row.keys()})
        writer.writerow(columns)
        for row in rows:
            writer.writerow([row.get(key, "") for key in columns])
        return response


class LocationViewSet(viewsets.ModelViewSet):
    queryset = InventoryLocation.objects.all()
    serializer_class = InventoryLocationSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['plant', 'type', 'is_active']

class StockViewSet(viewsets.ViewSet):
    permission_classes = [IsAuthenticated]

    @action(detail=False, methods=['get'])
    def bulk(self, request):
        plant_id = request.query_params.get('plant')
        location_id = request.query_params.get('location')
        if not plant_id:
            return Response({"error": "plant param required"}, status=400)
        
        data = StockService.get_bulk_stock(plant_id, location_id)
        return Response(data)

    @action(detail=False, methods=['get'])
    def rolls(self, request):
        plant_id = request.query_params.get('plant')
        location_id = request.query_params.get('location')
        if not plant_id:
            return Response({"error": "plant param required"}, status=400)
        
        data = StockService.get_roll_stock(plant_id, location_id)
        return Response(data)

    @action(detail=False, methods=['get'], url_path='bulk-availability')
    def bulk_availability(self, request):
        """
        Return bulk availability grouped by plant + location for a material.
        Used by WCM to request inter-plant DC transfers.
        """
        material_id = request.query_params.get('material_id')
        exclude_plant_id = request.query_params.get('exclude_plant')
        if not material_id:
            return Response({"error": "material_id param required"}, status=400)

        from .models import InventoryBulk
        qs = InventoryBulk.objects.filter(material_id=material_id, qty_kg__gt=0).select_related('plant', 'location', 'material')
        if exclude_plant_id:
            qs = qs.exclude(plant_id=exclude_plant_id)

        grouped = {}
        for item in qs:
            plant_id = str(item.plant_id)
            if plant_id not in grouped:
                grouped[plant_id] = {
                    'plant_id': plant_id,
                    'plant_name': item.plant.name if item.plant else None,
                    'total_qty': 0.0,
                    'locations': []
                }
            grouped[plant_id]['total_qty'] += float(item.qty_kg or 0)
            grouped[plant_id]['locations'].append({
                'location_id': str(item.location_id),
                'location_name': item.location.name if item.location else None,
                'quantity': float(item.qty_kg or 0),
                'uom': item.material.base_uom if item.material else 'KG'
            })

        return Response(list(grouped.values()))

class GRNViewSet(viewsets.ViewSet):
    permission_classes = [IsAuthenticated]

    @staticmethod
    def _vendor_invoice_exists(klass: str, vendor: Vendor, invoice_no: str) -> bool:
        invoice_no = str(invoice_no or "").strip()
        if not invoice_no:
            return False
        if klass == "BULK":
            return BulkTransaction.objects.filter(
                vendor=vendor,
                vendor_invoice_no=invoice_no,
                type="INWARD",
            ).exists()
        if klass == "PACKAGING":
            return PackagingTransaction.objects.filter(
                vendor=vendor,
                vendor_invoice_no=invoice_no,
                type="INWARD",
            ).exists()
        if klass == "ROLL":
            return InventoryRoll.objects.filter(
                vendor=vendor,
                vendor_invoice_no=invoice_no,
            ).exists()
        return False

    @action(detail=False, methods=['get'], url_path='history')
    def history(self, request):
        """Normalized inward history across bulk, roll, and packaging GRNs."""
        return Response({"results": GRNHistoryService.list_history(request.query_params)})

    @action(detail=False, methods=['get'], url_path='history/reason-codes')
    def history_reason_codes(self, request):
        """Allowed correction reason codes for GRN history corrections."""
        return Response({"results": GRNHistoryService.reason_codes()})

    @action(
        detail=False,
        methods=['post'],
        url_path=r'history/(?P<source_type>[^/.]+)/(?P<source_id>[^/.]+)/correct',
    )
    def correct_history(self, request, source_type=None, source_id=None):
        """Post an auditable correction while keeping the original GRN row immutable."""
        try:
            audit = GRNHistoryService.correct(
                source_type=source_type or "",
                source_id=source_id or "",
                payload=request.data,
                user=request.user,
            )
            return Response(
                {
                    "status": "success",
                    "audit_id": str(audit.id),
                    "delta": audit.delta_json,
                },
                status=status.HTTP_201_CREATED,
            )
        except Exception as e:
            return Response({"error": str(getattr(e, "message", "") or e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'], url_path='create')
    @transaction.atomic
    def create_unified(self, request):
        """
        V3.6 Smart GRN adapter.

        Accepts the new class-tagged payload while keeping the existing bulk,
        roll, and packaging posting services as the accounting source of truth.
        """
        klass = str(request.data.get("klass") or request.data.get("stock_class") or "").upper().strip()
        if klass not in {"BULK", "ROLL", "PACKAGING"}:
            return Response({"error": "klass must be BULK, ROLL, or PACKAGING."}, status=status.HTTP_400_BAD_REQUEST)

        lines = request.data.get("lines") or request.data.get("items") or []
        if not isinstance(lines, list) or not lines:
            return Response({"error": "At least one GRN line is required."}, status=status.HTTP_400_BAD_REQUEST)

        vendor_id = request.data.get("vendor_id") or request.data.get("vendor")
        if not vendor_id:
            return Response({"error": "Vendor is required before posting GRN."}, status=status.HTTP_400_BAD_REQUEST)
        try:
            vendor = Vendor.objects.get(id=vendor_id)
        except Vendor.DoesNotExist:
            return Response({"error": "Vendor was not found."}, status=status.HTTP_400_BAD_REQUEST)

        fallback_location_id = (
            request.data.get("store_location_id")
            or request.data.get("warehouse_id")
            or request.data.get("location_id")
            or request.data.get("location")
        )
        invoice_no = str(request.data.get("vendor_invoice_no") or request.data.get("source_ref") or request.data.get("reference") or "").strip()
        manual_po_ref = str(request.data.get("manual_po_ref") or "").strip()
        receipt_ref = invoice_no or f"GRN-{timezone.now().strftime('%Y%m%d%H%M%S')}"
        created_refs = []
        total_qty = Decimal("0")
        total_value = Decimal("0")

        try:
            StockLifecycleService.ensure_default_period()
            if invoice_no:
                Vendor.objects.select_for_update().filter(id=vendor.id).first()
                if self._vendor_invoice_exists(klass, vendor, invoice_no):
                    raise ValidationError(f"Duplicate vendor invoice '{invoice_no}' for vendor {vendor.name}.")

            for line in lines:
                line = dict(line or {})
                material = _resolve_material_from_payload(line)
                location = _resolve_location_from_payload(line, fallback_id=fallback_location_id)
                plant = location.plant
                if request.data.get("plant_id") and str(plant.id) != str(request.data.get("plant_id")):
                    raise ValidationError("Receiving location does not belong to selected plant.")

                line_ref = " | ".join(part for part in [
                    receipt_ref,
                    str(line.get("lot_no") or line.get("lot_id") or line.get("vendor_lot_ref") or "").strip(),
                ] if part)

                if klass == "BULK":
                    base_uom = _validate_receipt_line_uom(line, material)
                    qty = _as_decimal(line.get("qty") or line.get("quantity"))
                    rate = _as_decimal(line.get("rate_per_uom") or line.get("unit_cost") or line.get("cost") or request.data.get("unit_cost"))
                    tx = GRNService.create_bulk_grn(
                        material=material,
                        location=location,
                        vendor=vendor,
                        quantity=qty,
                        plant=plant,
                        cost=rate,
                        reference=line_ref,
                        granule_code_id=line.get("granule_code_id") or None,
                        granule_code=line.get("granule_code") or "",
                        vendor_invoice_no=invoice_no,
                        manual_po_ref=manual_po_ref,
                        allow_duplicate_vendor_invoice=True,
                        qty_uom=base_uom,
                    )
                    total_qty += qty
                    total_value += qty * rate
                    created_refs.append({"id": str(tx.id), "ref": tx.reference, "type": "BULK", "uom": base_uom})

                elif klass == "PACKAGING":
                    base_uom = _validate_receipt_line_uom(line, material)
                    qty = _as_decimal(line.get("qty") or line.get("quantity"))
                    rate = _as_decimal(line.get("rate_per_uom") or line.get("unit_cost") or line.get("cost"))
                    tx = PackagingService.add_packaging_stock(
                        material_id=material.id,
                        qty=qty,
                        location_id=location.id,
                        cost=rate,
                        vendor_id=vendor.id,
                        reference=line_ref or "PACKAGING_GRN",
                        input_uom=base_uom,
                        vendor_invoice_no=invoice_no,
                        manual_po_ref=manual_po_ref,
                        allow_duplicate_vendor_invoice=True,
                        meta_json={
                            "vendor_invoice_no": invoice_no,
                            "packaging_kind": line.get("packaging_kind") or getattr(material, "packaging_kind", ""),
                            "pcs_per_pack": line.get("pcs_per_pack") or None,
                            "color_variant": line.get("color_variant") or "",
                        },
                    )
                    total_qty += qty
                    total_value += qty * rate
                    created_refs.append({"id": str(tx.id), "ref": tx.reference, "type": "PACKAGING", "uom": base_uom})

                else:
                    gross = line.get("gross_weight_kg")
                    tare = line.get("tare_weight_kg")
                    net_source = line.get("net_weight_kg") or line.get("weight_kg") or line.get("qty") or line.get("quantity")
                    if net_source in (None, "") and gross not in (None, "") and tare not in (None, ""):
                        net_weight = _as_decimal(gross) - _as_decimal(tare)
                    else:
                        net_weight = _as_decimal(net_source)
                    if net_weight <= 0:
                        raise ValidationError("Roll net weight must be positive. Enter gross and tare, or net weight.")
                    if gross not in (None, "") or tare not in (None, ""):
                        gross_dec = _as_decimal(gross)
                        tare_dec = _as_decimal(tare)
                        if abs(gross_dec - net_weight - tare_dec) > Decimal("0.05"):
                            raise ValidationError("Roll gross weight must equal net + tare within 0.05 KG.")
                    label = str(line.get("roll_label") or line.get("label_id") or "").strip()
                    if label and InventoryRoll.objects.filter(label_id=label).exists():
                        raise ValidationError(f"Roll label {label} already exists.")
                    created = GRNService.create_roll_grn(
                        material=material,
                        location=location,
                        vendor=vendor,
                        plant=plant,
                        rolls_data=[{
                            "label_id": label or None,
                            "batch_no": line.get("lot_no") or line.get("batch_no") or line.get("vendor_lot_ref") or "",
                            "thickness_micron": line.get("thickness_um") or line.get("thickness_micron"),
                            "width_mm": line.get("width_mm"),
                            "stock_form": line.get("stock_form") or "OPEN_WEB",
                            "width_basis": line.get("width_basis") or "",
                            "weight_kg": net_weight,
                            "length_m": line.get("length_m") or 0,
                            "grade_id": line.get("grade_id") or line.get("grade") or None,
                            "unit_cost": line.get("rate_per_kg") or line.get("rate_per_uom") or line.get("unit_cost") or 0,
                        }],
                        reference=line_ref,
                        vendor_invoice_no=invoice_no,
                        manual_po_ref=manual_po_ref,
                        allow_duplicate_vendor_invoice=True,
                    )
                    roll = created[0]
                    update_fields = []
                    if gross not in (None, ""):
                        roll.gross_weight_kg = _as_decimal(gross)
                        update_fields.append("gross_weight_kg")
                    if tare not in (None, ""):
                        roll.tare_weight_kg = _as_decimal(tare)
                        update_fields.append("tare_weight_kg")
                    roll.net_weight_kg = net_weight
                    update_fields.append("net_weight_kg")
                    meta = dict(roll.meta_json or {})
                    meta.update({
                        "vendor_roll_label": line.get("vendor_roll_label") or "",
                        "treatment_side": line.get("treatment_side") or "",
                        "print_direction": line.get("print_direction") or "",
                        "core_size_inch": line.get("core_size_inch") or "",
                        "vendor_invoice_no": invoice_no,
                    })
                    roll.meta_json = meta
                    update_fields.append("meta_json")
                    roll.save(update_fields=update_fields)
                    rate = _as_decimal(line.get("rate_per_kg") or line.get("rate_per_uom") or line.get("unit_cost"))
                    total_qty += net_weight
                    total_value += net_weight * rate
                    created_refs.append({"id": str(roll.id), "ref": roll.label_id, "type": "ROLL"})

        except (ValidationError, InventoryMaterial.DoesNotExist, InventoryLocation.DoesNotExist) as exc:
            transaction.set_rollback(True)
            return Response({"error": _api_error_message(exc)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as exc:
            transaction.set_rollback(True)
            return Response({"error": _api_error_message(exc)}, status=status.HTTP_400_BAD_REQUEST)

        freight = _as_decimal(request.data.get("freight"))
        other_charges = _as_decimal(request.data.get("other_charges"))
        gst_percent = _as_decimal(request.data.get("gst_percent"))
        taxable_base = total_value + freight + other_charges
        gst_value = taxable_base * gst_percent / Decimal("100")
        grand_total = taxable_base + gst_value
        grn_no = f"GRN/{timezone.now().strftime('%Y/%m')}/{str(uuid.uuid4())[:5].upper()}"
        return Response(
            {
                "id": str(uuid.uuid4()),
                "grn_no": grn_no,
                "status": "POSTED",
                "klass": klass,
                "totals": {
                    "qty": float(total_qty),
                    "value": float(total_value),
                    "freight": float(freight),
                    "other_charges": float(other_charges),
                    "taxable_base": float(taxable_base),
                    "gst_percent": float(gst_percent),
                    "gst": float(gst_value),
                    "grand_total": float(grand_total),
                },
                "stock_movements": created_refs,
            },
            status=status.HTTP_201_CREATED,
        )

    @action(detail=False, methods=['get'], url_path='roll-upload-template')
    def roll_upload_template(self, request):
        """
        Generate a GRN roll-upload workbook from the current master database.

        This avoids shipping stale Excel dropdowns: Render downloads contain
        Render vendors, plants, locations, film variants, and active grades.
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.worksheet.datavalidation import DataValidation
        except Exception as exc:
            return Response({"error": f"Excel template support is unavailable: {_api_error_message(exc)}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        from apps.recipes.models import RecipeGrade

        vendors = list(
            Vendor.objects
            .filter(status="ACTIVE", type__in=["RM", "BOTH"])
            .order_by("code", "name")
        )
        locations = list(
            InventoryLocation.objects
            .select_related("plant")
            .filter(is_active=True, type__in=GRNService.ALLOWED_LOCATION_TYPES)
            .order_by("plant__code", "code", "name")
        )
        plants_by_code = {}
        for location in locations:
            if location.plant and location.plant.code not in plants_by_code:
                plants_by_code[location.plant.code] = location.plant
        plants = sorted(plants_by_code.values(), key=lambda plant: (plant.code or "", plant.name or ""))
        materials = list(
            InventoryMaterial.objects
            .filter(category="FILM_VARIANT")
            .exclude(status="INACTIVE")
            .order_by("code", "name")[:2000]
        )
        grades = list(RecipeGrade.objects.filter(is_active=True).order_by("name")[:1000])

        wb = Workbook()
        instructions = wb.active
        instructions.title = "Instructions"
        header = wb.create_sheet("GRN Header")
        lines = wb.create_sheet("Roll Lines")
        lookups = wb.create_sheet("Lookups")

        title_fill = PatternFill("solid", fgColor="1F6FEB")
        header_fill = PatternFill("solid", fgColor="EAF2FF")
        required_fill = PatternFill("solid", fgColor="FFF7ED")
        ok_fill = PatternFill("solid", fgColor="ECFDF5")
        title_font = Font(bold=True, color="FFFFFF", size=14)
        head_font = Font(bold=True, color="0F172A")
        mono_font = Font(name="Consolas", size=10)

        instructions["A1"] = "GRN bulk roll upload"
        instructions["A1"].fill = title_fill
        instructions["A1"].font = title_font
        instructions.merge_cells("A1:F1")
        instruction_rows = [
            ("1", "Fill GRN Header", "Pick Vendor Code, Receiving Plant, and Receiving Warehouse from the dropdowns."),
            ("2", "Fill Roll Lines", "One physical roll per row. Use Material Code or exact Material Name from the dropdowns."),
            ("3", "Weights", "Net Wt KG may be blank when Gross Wt KG and Tare/Core KG are filled."),
            ("4", "Validation", "Upload first validates the sheet. Stock is posted only after you click Post GRN."),
            ("5", "Master lists", "The Lookups tab is generated from the current database at download time."),
        ]
        instructions.append(["Step", "What to do", "Notes"])
        for row in instruction_rows:
            instructions.append(row)

        header_rows = [
            ("Vendor Code", vendors[0].code if vendors else "", "required", "Dropdown from active RM/BOTH vendors"),
            ("Receiving Plant", plants[0].code if plants else "", "required", "Dropdown from plants that have GRN locations"),
            ("Receiving Warehouse", locations[0].code if locations else "", "required", "Dropdown from active WAREHOUSE/QC/RM/WIP locations"),
            ("Vendor Invoice No", "", "optional", "Used as the GRN reference"),
            ("Vendor Invoice Date", date.today().isoformat(), "optional", "yyyy-mm-dd"),
        ]
        header.append(["Field", "Value", "Required", "Notes"])
        for row in header_rows:
            header.append(row)

        lookup_headers = [
            "Vendor Code", "Vendor Name",
            "Plant Code", "Plant Name",
            "Location Code", "Location Name", "Location Display",
            "Material Code", "Material Name",
            "Grade",
        ]
        lookups.append(lookup_headers)
        max_lookup_rows = max(len(vendors), len(plants), len(locations), len(materials), len(grades), 1)
        for idx in range(max_lookup_rows):
            vendor = vendors[idx] if idx < len(vendors) else None
            plant = plants[idx] if idx < len(plants) else None
            location = locations[idx] if idx < len(locations) else None
            material = materials[idx] if idx < len(materials) else None
            grade = grades[idx] if idx < len(grades) else None
            lookups.append([
                vendor.code if vendor else "",
                vendor.name if vendor else "",
                plant.code if plant else "",
                plant.name if plant else "",
                location.code if location else "",
                location.name if location else "",
                f"{location.code} · {location.name} · {location.plant.code if location.plant else ''}" if location else "",
                material.code if material else "",
                material.name if material else "",
                grade.name if grade else "",
            ])

        line_headers = [
            "Line No",
            "Supplier Roll No",
            "ERP Roll Label",
            "Material Type",
            "Material Code",
            "Material Name",
            "Vendor Lot Ref",
            "Gross Wt KG",
            "Tare/Core KG",
            "Net Wt KG",
            "Width MM",
            "Thickness Micron",
            "Length Meter",
            "Roll Form",
            "Grade",
            "Unit Cost",
            "Mfg Date",
            "Best Before",
            "QC Status",
            "Location Code",
            "Remarks",
            "Check Status",
            "Error Hint",
        ]
        lines.append(line_headers)
        for row_no in range(2, 502):
            line_index = row_no - 1
            lines.append([
                line_index,
                "",
                "",
                "FILM",
                "",
                "",
                "",
                "",
                "",
                f'=IF(AND(H{row_no}<>"",I{row_no}<>""),H{row_no}-I{row_no},"")',
                "",
                "",
                "",
                "PARENT",
                "",
                "",
                "",
                "",
                "PENDING",
                "",
                "",
                f'=IF(OR(E{row_no}<>"",F{row_no}<>""),"READY","")',
                f'=IF(AND(E{row_no}="",F{row_no}<>""),"Name-only match",IF(AND(E{row_no}<>"",F{row_no}<>""),"Code wins if code/name differ",""))',
            ])

        def add_list_validation(sheet, cell_range, values_range, allow_blank=True):
            dv = DataValidation(type="list", formula1=values_range, allow_blank=allow_blank)
            sheet.add_data_validation(dv)
            dv.add(cell_range)

        vendor_end = max(len(vendors) + 1, 2)
        plant_end = max(len(plants) + 1, 2)
        location_end = max(len(locations) + 1, 2)
        material_end = max(len(materials) + 1, 2)
        grade_end = max(len(grades) + 1, 2)
        add_list_validation(header, "B2", f"'Lookups'!$A$2:$A${vendor_end}", allow_blank=False)
        add_list_validation(header, "B3", f"'Lookups'!$C$2:$C${plant_end}", allow_blank=False)
        add_list_validation(header, "B4", f"'Lookups'!$E$2:$E${location_end}", allow_blank=False)
        add_list_validation(lines, "D2:D501", '"FILM"', allow_blank=False)
        add_list_validation(lines, "E2:E501", f"'Lookups'!$H$2:$H${material_end}")
        add_list_validation(lines, "F2:F501", f"'Lookups'!$I$2:$I${material_end}")
        add_list_validation(lines, "N2:N501", '"PARENT,SLIT,REMAINDER"', allow_blank=True)
        add_list_validation(lines, "O2:O501", f"'Lookups'!$J$2:$J${grade_end}")
        add_list_validation(lines, "S2:S501", '"PENDING,PASS,HOLD,REJECT"', allow_blank=True)
        add_list_validation(lines, "T2:T501", f"'Lookups'!$E$2:$E${location_end}")

        for sheet in (instructions, header, lines, lookups):
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = head_font
                cell.alignment = Alignment(horizontal="center")
            for row in sheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
        instructions["A1"].fill = title_fill
        instructions["A1"].font = title_font
        header["A1"].fill = header_fill
        for required_cell in ("A2", "A3", "A4", "B2", "B3", "B4"):
            header[required_cell].fill = required_fill
        for col in ("E", "F", "T"):
            for cell in lines[f"{col}2:{col}501"]:
                cell[0].fill = ok_fill
        for row in lines.iter_rows(min_row=2, max_row=501):
            row[0].font = mono_font
            row[1].font = mono_font
            row[2].font = mono_font
            row[4].font = mono_font
            row[19].font = mono_font
        lines.auto_filter.ref = "A1:W501"
        lookups.auto_filter.ref = f"A1:J{max_lookup_rows + 1}"

        widths = {
            instructions: [10, 28, 90],
            header: [26, 32, 16, 70],
            lines: [10, 20, 20, 16, 28, 42, 22, 14, 14, 14, 14, 18, 16, 16, 22, 14, 16, 16, 16, 22, 40, 18, 34],
            lookups: [22, 34, 18, 34, 22, 34, 52, 28, 52, 28],
        }
        for sheet, sizes in widths.items():
            for index, width in enumerate(sizes, start=1):
                sheet.column_dimensions[chr(64 + index)].width = width

        payload = BytesIO()
        wb.save(payload)
        payload.seek(0)
        response = HttpResponse(
            payload.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="grn_live_roll_upload_template.xlsx"'
        return response

    @action(detail=False, methods=['post'], url_path='upload-rolls')
    def upload_rolls(self, request):
        """
        Excel bulk upload for purchased film roll GRNs.

        Expected workbook tabs:
        - GRN Header: two-column key/value sheet.
        - Roll Lines: one physical roll per row.
        """
        uploaded = request.FILES.get("file") or request.FILES.get("upload")
        if not uploaded:
            return Response({"error": "Upload an .xlsx file in field 'file'."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            from openpyxl import load_workbook
            workbook = load_workbook(uploaded, data_only=True, read_only=True)
        except Exception as exc:
            return Response({"error": f"Could not read Excel file: {_api_error_message(exc)}"}, status=status.HTTP_400_BAD_REQUEST)

        if "GRN Header" not in workbook.sheetnames or "Roll Lines" not in workbook.sheetnames:
            return Response({"error": "Workbook must contain 'GRN Header' and 'Roll Lines' sheets."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            header_ws = workbook["GRN Header"]
            line_ws = workbook["Roll Lines"]
            header = {}
            for row in header_ws.iter_rows(min_row=1, max_col=2, values_only=True):
                key = _normalise_upload_key(row[0])
                if key:
                    header[key] = row[1]

            vendor_ref = request.data.get("vendor_id") or request.data.get("vendor") or header.get("vendorcode")
            vendor = _resolve_upload_vendor(vendor_ref)
            plant_ref = request.data.get("plant_id") or request.data.get("plant") or header.get("receivingplant") or ""
            default_location_ref = (
                request.data.get("store_location_id")
                or request.data.get("warehouse_id")
                or request.data.get("location_id")
                or header.get("receivingwarehouse")
                or ""
            )
            invoice_no = _upload_cell(request.data.get("vendor_invoice_no") or header.get("vendorinvoiceno")) or f"GRN-UPLOAD-{timezone.now().strftime('%Y%m%d%H%M%S')}"
            invoice_date = _upload_cell(request.data.get("vendor_invoice_date") or header.get("vendorinvoicedate"))

            raw_headers = next(line_ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not raw_headers:
                raise ValidationError("Roll Lines sheet is missing headers.")
            column_map = {_normalise_upload_key(label): idx for idx, label in enumerate(raw_headers)}

            def cell(row, label, default=""):
                idx = column_map.get(_normalise_upload_key(label))
                if idx is None or idx >= len(row):
                    return default
                value = row[idx]
                return default if value is None else value

            business_columns = [
                "Supplier Roll No",
                "ERP Roll Label",
                "Material Code",
                "Material Name",
                "Vendor Lot Ref",
                "Gross Wt KG",
                "Tare/Core KG",
                "Net Wt KG",
                "Width MM",
                "Thickness Micron",
                "Length Meter",
                "Grade",
                "Unit Cost",
                "Mfg Date",
                "Best Before",
                "Location Code",
                "Remarks",
            ]
            parsed_rows = []
            errors = []
            seen_labels = set()
            for excel_row_no, row in enumerate(line_ws.iter_rows(min_row=2, values_only=True), start=2):
                if not any(_upload_cell(cell(row, label)) for label in business_columns):
                    continue
                if len(parsed_rows) >= 500:
                    errors.append({"row": excel_row_no, "error": "Maximum 500 roll rows per upload."})
                    break

                try:
                    material_code = _upload_cell(cell(row, "Material Code"))
                    material_name = _upload_cell(cell(row, "Material Name"))
                    if not material_code and not material_name:
                        raise ValidationError("Material Code or exact Material Name is required.")
                    material = _resolve_material_from_payload({"material_code": material_code, "material_name": material_name})
                    if str(material.category or "").upper() != "FILM_VARIANT":
                        raise ValidationError(f"Roll upload requires FILM_VARIANT material. {material_code or material_name} is {material.category}.")

                    supplier_roll_no = _upload_cell(cell(row, "Supplier Roll No"))
                    label_id = _upload_cell(cell(row, "ERP Roll Label")) or supplier_roll_no
                    if not label_id:
                        raise ValidationError("Supplier Roll No or ERP Roll Label is required.")
                    label_key = label_id.upper()
                    if label_key in seen_labels:
                        raise ValidationError(f"Duplicate roll label {label_id} inside upload.")
                    seen_labels.add(label_key)
                    if InventoryRoll.objects.filter(label_id=label_id).exists():
                        raise ValidationError(f"Roll label {label_id} already exists.")

                    gross = _upload_decimal(cell(row, "Gross Wt KG"), default=0)
                    tare = _upload_decimal(cell(row, "Tare/Core KG"), default=0)
                    net = _upload_decimal(cell(row, "Net Wt KG"), default=None)
                    if net is None:
                        net = gross - tare
                    if net <= 0:
                        raise ValidationError("Net Wt KG must be positive, or Gross Wt KG must exceed Tare/Core KG.")
                    if gross and abs(gross - tare - net) > Decimal("0.05"):
                        raise ValidationError("Gross Wt KG must equal Net Wt KG + Tare/Core KG within 0.05 KG.")

                    width = _upload_decimal(cell(row, "Width MM"), default=0)
                    thickness = _upload_decimal(cell(row, "Thickness Micron"), default=0)
                    if width <= 0:
                        raise ValidationError("Width MM is required and must be positive.")
                    if thickness <= 0:
                        raise ValidationError("Thickness Micron is required and must be positive.")

                    location_ref = _upload_cell(cell(row, "Location Code")) or default_location_ref
                    location = _resolve_upload_location(location_ref, plant_ref=plant_ref)
                    grade_ref = _upload_cell(cell(row, "Grade"))
                    grade = _resolve_upload_grade(grade_ref) if getattr(material, "is_extrudable", False) else None

                    parsed_rows.append({
                        "excel_row": excel_row_no,
                        "material": material,
                        "location": location,
                        "label_id": label_id,
                        "supplier_roll_no": supplier_roll_no,
                        "batch_no": _upload_cell(cell(row, "Vendor Lot Ref")),
                        "gross_weight_kg": gross,
                        "tare_weight_kg": tare,
                        "net_weight_kg": net,
                        "width_mm": width,
                        "stock_form": _upload_cell(cell(row, "Stock Form")) or "OPEN_WEB",
                        "width_basis": _upload_cell(cell(row, "Width Basis")) or "",
                        "thickness_micron": thickness,
                        "length_m": _upload_decimal(cell(row, "Length Meter"), default=0),
                        "grade_id": str(grade.id) if grade else None,
                        "unit_cost": _upload_decimal(cell(row, "Unit Cost"), default=0),
                        "mfg_date": _upload_cell(cell(row, "Mfg Date")),
                        "best_before": _upload_cell(cell(row, "Best Before")),
                        "qc_status": _upload_cell(cell(row, "QC Status")),
                        "remarks": _upload_cell(cell(row, "Remarks")),
                    })
                except Exception as exc:
                    errors.append({"row": excel_row_no, "error": _api_error_message(exc)})

            if not parsed_rows:
                errors.append({"row": None, "error": "No roll rows found in Roll Lines sheet."})
            if errors:
                return Response({"status": "FAILED", "error": "Upload validation failed.", "errors": errors[:50]}, status=status.HTTP_400_BAD_REQUEST)

            dry_run = str(request.data.get("dry_run") or "").strip().lower() in {"1", "true", "yes"}
            created_refs = []
            total_qty = Decimal("0")
            total_value = Decimal("0")
            if not dry_run:
                with transaction.atomic():
                    for parsed in parsed_rows:
                        reference = " | ".join(part for part in [invoice_no, parsed["batch_no"]] if part)
                        created = GRNService.create_roll_grn(
                            material=parsed["material"],
                            location=parsed["location"],
                            vendor=vendor,
                            plant=parsed["location"].plant,
                            rolls_data=[{
                                "label_id": parsed["label_id"],
                                "batch_no": parsed["batch_no"],
                                "thickness_micron": parsed["thickness_micron"],
                                "width_mm": parsed["width_mm"],
                                "stock_form": parsed["stock_form"],
                                "width_basis": parsed["width_basis"],
                                "weight_kg": parsed["net_weight_kg"],
                                "length_m": parsed["length_m"],
                                "grade_id": parsed["grade_id"],
                            }],
                            reference=reference,
                        )
                        roll = created[0]
                        meta = dict(roll.meta_json or {})
                        meta.update({
                            "vendor_roll_label": parsed["supplier_roll_no"],
                            "vendor_invoice_no": invoice_no,
                            "vendor_invoice_date": invoice_date,
                            "mfg_date": parsed["mfg_date"],
                            "best_before": parsed["best_before"],
                            "qc_status": parsed["qc_status"],
                            "upload_excel_row": parsed["excel_row"],
                            "remarks": parsed["remarks"],
                        })
                        roll.gross_weight_kg = parsed["gross_weight_kg"]
                        roll.tare_weight_kg = parsed["tare_weight_kg"]
                        roll.net_weight_kg = parsed["net_weight_kg"]
                        roll.meta_json = meta
                        roll.save(update_fields=["gross_weight_kg", "tare_weight_kg", "net_weight_kg", "meta_json"])
                        total_qty += parsed["net_weight_kg"]
                        total_value += parsed["net_weight_kg"] * parsed["unit_cost"]
                        created_refs.append({"id": str(roll.id), "ref": roll.label_id, "type": "ROLL", "excel_row": parsed["excel_row"]})
            else:
                total_qty = sum((row["net_weight_kg"] for row in parsed_rows), Decimal("0"))
                total_value = sum((row["net_weight_kg"] * row["unit_cost"] for row in parsed_rows), Decimal("0"))

            grn_no = f"GRN/{timezone.now().strftime('%Y/%m')}/{str(uuid.uuid4())[:5].upper()}"
            return Response(
                {
                    "id": str(uuid.uuid4()),
                    "grn_no": grn_no,
                    "status": "VALIDATED" if dry_run else "POSTED",
                    "klass": "ROLL",
                    "dry_run": dry_run,
                    "totals": {"qty": float(total_qty), "value": float(total_value)},
                    "stock_movements": created_refs,
                    "rows": len(parsed_rows),
                    "vendor": {"id": str(vendor.id), "code": vendor.code, "name": vendor.name},
                    "vendor_invoice_no": invoice_no,
                    "vendor_invoice_date": invoice_date,
                    "review_rows": [
                        {
                            "excel_row": parsed["excel_row"],
                            "supplier_roll_no": parsed["supplier_roll_no"],
                            "label_id": parsed["label_id"],
                            "material_code": parsed["material"].code,
                            "material_name": parsed["material"].name,
                            "batch_no": parsed["batch_no"],
                            "gross_weight_kg": float(parsed["gross_weight_kg"] or 0),
                            "tare_weight_kg": float(parsed["tare_weight_kg"] or 0),
                            "net_weight_kg": float(parsed["net_weight_kg"] or 0),
                            "width_mm": float(parsed["width_mm"] or 0),
                            "thickness_micron": float(parsed["thickness_micron"] or 0),
                            "length_m": float(parsed["length_m"] or 0),
                            "grade_id": parsed["grade_id"],
                            "unit_cost": float(parsed["unit_cost"] or 0),
                            "mfg_date": parsed["mfg_date"],
                            "best_before": parsed["best_before"],
                            "qc_status": parsed["qc_status"],
                            "location_id": str(parsed["location"].id),
                            "location_code": parsed["location"].code,
                            "location_name": parsed["location"].name,
                            "plant_code": parsed["location"].plant.code if parsed["location"].plant else "",
                            "remarks": parsed["remarks"],
                        }
                        for parsed in parsed_rows
                    ],
                },
                status=status.HTTP_200_OK if dry_run else status.HTTP_201_CREATED,
            )
        except (ValidationError, InventoryMaterial.DoesNotExist, InventoryLocation.DoesNotExist) as exc:
            return Response({"error": _api_error_message(exc)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as exc:
            return Response({"error": _api_error_message(exc)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'], url_path='post-roll-review')
    def post_roll_review(self, request):
        """Post reviewed/edited roll-upload rows after the UI validation step."""
        rows = request.data.get("review_rows") or request.data.get("rows") or []
        if not isinstance(rows, list) or not rows:
            return Response({"error": "Review rows are required before posting."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            vendor_ref = request.data.get("vendor_id") or request.data.get("vendor") or request.data.get("vendor_code")
            vendor = _resolve_upload_vendor(vendor_ref)
            plant_ref = request.data.get("plant_id") or request.data.get("plant") or ""
            default_location_ref = (
                request.data.get("store_location_id")
                or request.data.get("warehouse_id")
                or request.data.get("location_id")
                or request.data.get("location")
                or ""
            )
            invoice_no = _upload_cell(request.data.get("vendor_invoice_no")) or f"GRN-UPLOAD-{timezone.now().strftime('%Y%m%d%H%M%S')}"
            invoice_date = _upload_cell(request.data.get("vendor_invoice_date"))

            parsed_rows = []
            errors = []
            seen_labels = set()
            for index, row in enumerate(rows, start=1):
                row = dict(row or {})
                excel_row = row.get("excel_row") or index
                try:
                    material_code = _upload_cell(row.get("material_code"))
                    material_name = _upload_cell(row.get("material_name"))
                    if not material_code and not material_name:
                        raise ValidationError("Material Code or exact Material Name is required.")
                    material = _resolve_material_from_payload({"material_code": material_code, "material_name": material_name})
                    if str(material.category or "").upper() != "FILM_VARIANT":
                        raise ValidationError(f"Roll upload requires FILM_VARIANT material. {material_code or material_name} is {material.category}.")

                    supplier_roll_no = _upload_cell(row.get("supplier_roll_no") or row.get("vendor_roll_label"))
                    label_id = _upload_cell(row.get("label_id") or row.get("erp_roll_label")) or supplier_roll_no
                    if not label_id:
                        raise ValidationError("Supplier Roll No or ERP Roll Label is required.")
                    label_key = label_id.upper()
                    if label_key in seen_labels:
                        raise ValidationError(f"Duplicate roll label {label_id} inside review.")
                    seen_labels.add(label_key)
                    if InventoryRoll.objects.filter(label_id=label_id).exists():
                        raise ValidationError(f"Roll label {label_id} already exists.")

                    gross = _upload_decimal(row.get("gross_weight_kg"), default=0)
                    tare = _upload_decimal(row.get("tare_weight_kg"), default=0)
                    net = _upload_decimal(row.get("net_weight_kg"), default=None)
                    if net is None:
                        net = gross - tare
                    if net <= 0:
                        raise ValidationError("Net Wt KG must be positive, or Gross Wt KG must exceed Tare/Core KG.")
                    if gross and abs(gross - tare - net) > Decimal("0.05"):
                        raise ValidationError("Gross Wt KG must equal Net Wt KG + Tare/Core KG within 0.05 KG.")

                    width = _upload_decimal(row.get("width_mm"), default=0)
                    thickness = _upload_decimal(row.get("thickness_micron") or row.get("thickness_um"), default=0)
                    if width <= 0:
                        raise ValidationError("Width MM is required and must be positive.")
                    if thickness <= 0:
                        raise ValidationError("Thickness Micron is required and must be positive.")

                    location_ref = (
                        _upload_cell(row.get("location_id"))
                        or _upload_cell(row.get("location_code"))
                        or _upload_cell(row.get("location_name"))
                        or default_location_ref
                    )
                    location = _resolve_upload_location(location_ref, plant_ref=plant_ref)
                    grade_ref = _upload_cell(row.get("grade_id") or row.get("grade") or row.get("grade_name"))
                    grade = _resolve_upload_grade(grade_ref) if getattr(material, "is_extrudable", False) else None

                    parsed_rows.append({
                        "excel_row": excel_row,
                        "material": material,
                        "location": location,
                        "label_id": label_id,
                        "supplier_roll_no": supplier_roll_no,
                        "batch_no": _upload_cell(row.get("batch_no") or row.get("vendor_lot_ref")),
                        "gross_weight_kg": gross,
                        "tare_weight_kg": tare,
                        "net_weight_kg": net,
                        "width_mm": width,
                        "thickness_micron": thickness,
                        "length_m": _upload_decimal(row.get("length_m"), default=0),
                        "grade_id": str(grade.id) if grade else None,
                        "unit_cost": _upload_decimal(row.get("unit_cost"), default=0),
                        "mfg_date": _upload_cell(row.get("mfg_date")),
                        "best_before": _upload_cell(row.get("best_before")),
                        "qc_status": _upload_cell(row.get("qc_status")),
                        "remarks": _upload_cell(row.get("remarks")),
                    })
                except Exception as exc:
                    errors.append({"row": excel_row, "error": _api_error_message(exc)})

            if errors:
                return Response({"status": "FAILED", "error": "Review validation failed.", "errors": errors[:50]}, status=status.HTTP_400_BAD_REQUEST)

            created_refs = []
            total_qty = Decimal("0")
            total_value = Decimal("0")
            with transaction.atomic():
                for parsed in parsed_rows:
                    reference = " | ".join(part for part in [invoice_no, parsed["batch_no"]] if part)
                    created = GRNService.create_roll_grn(
                        material=parsed["material"],
                        location=parsed["location"],
                        vendor=vendor,
                        plant=parsed["location"].plant,
                        rolls_data=[{
                            "label_id": parsed["label_id"],
                            "batch_no": parsed["batch_no"],
                            "thickness_micron": parsed["thickness_micron"],
                            "width_mm": parsed["width_mm"],
                            "stock_form": parsed.get("stock_form") or "OPEN_WEB",
                            "width_basis": parsed.get("width_basis") or "",
                            "weight_kg": parsed["net_weight_kg"],
                            "length_m": parsed["length_m"],
                            "grade_id": parsed["grade_id"],
                            "unit_cost": parsed["unit_cost"],
                        }],
                        reference=reference,
                    )
                    roll = created[0]
                    meta = dict(roll.meta_json or {})
                    meta.update({
                        "vendor_roll_label": parsed["supplier_roll_no"],
                        "vendor_invoice_no": invoice_no,
                        "vendor_invoice_date": invoice_date,
                        "mfg_date": parsed["mfg_date"],
                        "best_before": parsed["best_before"],
                        "qc_status": parsed["qc_status"],
                        "upload_excel_row": parsed["excel_row"],
                        "remarks": parsed["remarks"],
                    })
                    roll.gross_weight_kg = parsed["gross_weight_kg"]
                    roll.tare_weight_kg = parsed["tare_weight_kg"]
                    roll.net_weight_kg = parsed["net_weight_kg"]
                    roll.meta_json = meta
                    roll.save(update_fields=["gross_weight_kg", "tare_weight_kg", "net_weight_kg", "meta_json"])
                    total_qty += parsed["net_weight_kg"]
                    total_value += parsed["net_weight_kg"] * parsed["unit_cost"]
                    created_refs.append({"id": str(roll.id), "ref": roll.label_id, "type": "ROLL", "excel_row": parsed["excel_row"]})

            grn_no = f"GRN/{timezone.now().strftime('%Y/%m')}/{str(uuid.uuid4())[:5].upper()}"
            return Response(
                {
                    "id": str(uuid.uuid4()),
                    "grn_no": grn_no,
                    "status": "POSTED",
                    "klass": "ROLL",
                    "totals": {"qty": float(total_qty), "value": float(total_value)},
                    "stock_movements": created_refs,
                    "rows": len(parsed_rows),
                    "vendor": {"id": str(vendor.id), "code": vendor.code, "name": vendor.name},
                },
                status=status.HTTP_201_CREATED,
            )
        except (ValidationError, InventoryMaterial.DoesNotExist, InventoryLocation.DoesNotExist) as exc:
            return Response({"error": _api_error_message(exc)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as exc:
            return Response({"error": _api_error_message(exc)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'], url_path='bulk')
    def create_bulk(self, request):
        """Phase 56: Bulk GRN with cost tracking."""
        serializer = BulkGRNSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        try:
            material = InventoryMaterial.objects.get(id=data['material_id'])
            location = InventoryLocation.objects.get(id=data['location_id'])
            vendor = Vendor.objects.get(id=data['vendor_id'])
            plant = data.get('plant_id')
            if plant:
                from apps.factory.models import Plant
                plant = Plant.objects.get(id=plant)
            else:
                plant = location.plant
            
            GRNService.create_bulk_grn(
                material=material,
                location=location,
                vendor=vendor,
                quantity=float(data['quantity']),
                plant=plant,
                cost=float(data.get('cost', 0)),  # Phase 56: Cost for avg calculation
                reference=data.get('reference', ""),
                granule_code_id=str(data.get('granule_code_id')) if data.get('granule_code_id') else None,
                granule_code=data.get('granule_code', ""),
                qty_uom=data.get('uom') or getattr(material, "base_uom", None),
            )
            return Response({"status": "success"}, status=status.HTTP_201_CREATED)
        except Exception as e:
            return Response({"error": str(getattr(e, "message", "") or e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'], url_path='roll')
    def create_roll(self, request):
        serializer = RollGRNSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        try:
            material = InventoryMaterial.objects.get(id=data['material_id'])
            location = InventoryLocation.objects.get(id=data['location_id'])
            vendor = Vendor.objects.get(id=data['vendor_id'])
            plant = data.get('plant_id')
            if plant:
                from apps.factory.models import Plant
                plant = Plant.objects.get(id=plant)
            else:
                plant = location.plant
            
            rolls_data = []
            for item in data['rolls']:
                # The serializer already validated these as Decimals or expected types
                rolls_data.append(dict(item))

            GRNService.create_roll_grn(
                material=material,
                location=location,
                vendor=vendor,
                plant=plant,
                rolls_data=rolls_data,
                reference=data.get('reference', "")
            )
            return Response({"status": "success"}, status=status.HTTP_201_CREATED)
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'], url_path='packaging')
    def create_packaging(self, request):
        serializer = PackagingGRNSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        try:
            material = InventoryMaterial.objects.get(id=data['material_id'])
            location = InventoryLocation.objects.get(id=data['location_id'])
            vendor = Vendor.objects.get(id=data['vendor_id'])

            if str(material.category or '').upper() != 'PACKAGING':
                return Response({"error": "Packaging GRN requires category=PACKAGING material."}, status=status.HTTP_400_BAD_REQUEST)

            PackagingService.add_packaging_stock(
                material_id=material.id,
                qty=data['quantity'],
                location_id=location.id,
                cost=data.get('cost', 0),
                vendor_id=vendor.id,
                reference=data.get('reference', '') or 'PACKAGING_GRN',
            )
            return Response({"status": "success"}, status=status.HTTP_201_CREATED)
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

class JobWorkOrderViewSet(viewsets.ModelViewSet):
    queryset = JobWorkOrder.objects.all().order_by('-created_at')
    serializer_class = JobWorkOrderSerializer
    permission_classes = [IsAuthenticated]

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        try:
            serializer.is_valid(raise_exception=True)
            self.perform_create(serializer)
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        except Exception as e:
            detail = getattr(e, "detail", None)
            if isinstance(detail, dict):
                return Response(
                    {"status": "error", "message": "Validation failed.", "field_errors": detail},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            return Response({"status": "error", "message": "Request failed."}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['get'], url_path='vendor-candidates')
    def vendor_candidates(self, request):
        process_code = str(request.query_params.get('process_code') or "").upper()
        plant_id = request.query_params.get('plant_id')
        production_job_id = request.query_params.get('production_job_id')
        plant = None
        if production_job_id:
            try:
                job = ProductionJob.objects.select_related(
                    'current_process',
                    'process',
                    'work_center__plant',
                    'from_location__plant',
                    'to_location__plant',
                ).get(id=production_job_id)
                process = job.current_process or job.process
                process_code = str(getattr(process, "code", process_code) or "").upper()
                if job.work_center and job.work_center.plant:
                    plant = job.work_center.plant
                elif job.from_location and job.from_location.plant:
                    plant = job.from_location.plant
                elif job.to_location and job.to_location.plant:
                    plant = job.to_location.plant
            except ProductionJob.DoesNotExist:
                return Response({"status": "error", "message": "production_job not found."}, status=status.HTTP_404_NOT_FOUND)
        elif plant_id:
            try:
                from apps.factory.models import Plant
                plant = Plant.objects.get(id=plant_id)
            except Exception:
                plant = None

        rows = []
        for vendor, verdict in JobWorkService.compatible_vendors(process_code=process_code, plant=plant):
            rows.append({
                "id": str(vendor.id),
                "name": vendor.name,
                "code": vendor.code,
                "type": vendor.type,
                "status": vendor.status,
                "turnaround_hours": vendor.turnaround_hours,
                "qc_required": vendor.qc_required,
                "jobwork_capabilities": vendor.jobwork_capabilities or [],
                "jobwork_plants": vendor.jobwork_plants or [],
                "vendor_capability_match": verdict["match"],
                "match_reasons": verdict["reasons"],
            })
        return Response({"results": rows})

    @action(detail=True, methods=['get'], url_path='eligible-rolls')
    def eligible_rolls(self, request, pk=None):
        try:
            order = self.get_object()
            qs = InventoryRoll.objects.select_related(
                'material',
                'location',
                'location__plant',
                'production_job',
            ).filter(location__plant=order.plant)
            if order.production_job_id:
                # Prefer source job rolls first.
                source_qs = qs.filter(
                    Q(production_job=order.production_job) | Q(created_by_job=order.production_job)
                )
                if source_qs.exists():
                    qs = source_qs
            qs = qs.filter(status__in=['AVAILABLE', 'RESERVED', 'IN_PROCESS', 'SENT_JOBWORK']).order_by('-created_at')[:200]
            rows = []
            for roll in qs:
                rows.append({
                    "id": str(roll.id),
                    "label_id": roll.label_id,
                    "material_name": getattr(roll.material, "name", None),
                    "status": roll.status,
                    "weight_kg": float(roll.weight_kg or 0),
                    "location_name": getattr(roll.location, "name", None),
                    "production_job_number": getattr(roll.production_job, "job_number", None),
                })
            return Response({"results": rows})
        except Exception as e:
            return Response({"status": "error", "message": "Request failed."}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='dispatch')
    def dispatch_order(self, request, pk=None):
        serializer = JobWorkDispatchSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        
        try:
            JobWorkService.dispatch_material(
                order_id=pk,
                roll_ids=[str(i) for i in data.get('roll_ids', [])],
                bulk_items=data.get('bulk_items')
            )
            return Response({"status": "dispatched"})
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'])
    def receive(self, request, pk=None):
        serializer = JobWorkReceiveSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        
        try:
            target_location = InventoryLocation.objects.get(id=data['target_location_id'])
            
            received_rolls = []
            if 'received_rolls' in data:
                for item in data['received_rolls']:
                    r_dict = dict(item)
                    if 'material_id' not in r_dict or not r_dict['material_id']:
                         return Response({"error": "material_id required for received rolls"}, status=400)
                    
                    r_dict['material_id'] = str(r_dict['material_id']) 
                    r_dict['thickness_micron'] = float(r_dict['thickness_micron'])
                    r_dict['width_mm'] = float(r_dict['width_mm'])
                    r_dict['weight_kg'] = float(r_dict['weight_kg'])
                    if r_dict.get('grade_id'):
                        r_dict['grade_id'] = str(r_dict['grade_id'])
                    if 'length_m' in r_dict:
                        r_dict['length_m'] = float(r_dict['length_m'])
                    received_rolls.append(r_dict)
            
            received_bulk = data.get('received_bulk', [])

            JobWorkService.receive_material(
                order_id=pk,
                target_location=target_location,
                received_rolls=received_rolls,
                received_bulk=received_bulk
            )
            
            return Response({"status": "received"})
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

class DeliveryChallanViewSet(viewsets.ModelViewSet):
    queryset = DeliveryChallan.objects.select_related(
        'from_plant',
        'to_plant',
        'source_job',
        'target_job',
    ).prefetch_related(
        'items__roll',
        'items__material',
        'items__from_location',
        'items__to_location',
    ).order_by('-created_at')
    serializer_class = DeliveryChallanSerializer
    permission_classes = [IsAuthenticated]

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        try:
            challan = InterPlantService.create_challan(
                from_plant_id=str(data["from_plant"].id),
                to_plant_id=str(data["to_plant"].id),
                source_job_id=str(data["source_job"].id) if data.get("source_job") else None,
                target_job_id=str(data["target_job"].id) if data.get("target_job") else None,
                is_system_generated=bool(data.get("is_system_generated", False)),
                vehicle_no=str(data.get("vehicle_no") or ""),
                driver_name=str(data.get("driver_name") or ""),
                driver_phone=str(data.get("driver_phone") or ""),
                transporter_name=str(data.get("transporter_name") or ""),
                lr_number=str(data.get("lr_number") or ""),
            )
            out = self.get_serializer(challan).data
            return Response(out, status=status.HTTP_201_CREATED)
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='dispatch')
    def dispatch_dc(self, request, pk=None):
        serializer = ChallanDispatchSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        
        try:
            InterPlantService.dispatch_challan(
                challan_id=pk,
                roll_ids=[str(i) for i in data.get('roll_ids', [])],
                bulk_items=data.get('bulk_items'),
                target_location_id=str(data['target_location_id']) if data.get('target_location_id') else None,
            )
            return Response({"status": "dispatched"})
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'])
    def receive_challan(self, request, pk=None):
        serializer = ChallanReceiveSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        
        try:
            InterPlantService.receive_challan(
                challan_id=pk,
                target_location_id=str(data['target_location_id']),
                roll_ids=[str(i) for i in data.get('roll_ids', [])],
                bulk_items=data.get('bulk_items', [])
            )
            return Response({"status": "received"})
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @xframe_options_exempt
    @action(detail=True, methods=['get'], url_path='print-pdf')
    def print_pdf(self, request, pk=None):
        try:
            challan = self.get_object()
            pdf_bytes = ChallanPDFService.generate_pdf_bytes(challan)
            file_name = f"{challan.dc_no or str(challan.id)}.pdf"
            return FileResponse(
                BytesIO(pdf_bytes),
                as_attachment=False,
                filename=file_name,
                content_type='application/pdf',
            )
        except RuntimeError as e:
            return Response({"error": "Request failed"}, status=status.HTTP_503_SERVICE_UNAVAILABLE)
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)



class PlantLocationListView(generics.ListAPIView):
    serializer_class = InventoryLocationSerializer
    permission_classes = [IsAuthenticated]
    
    def get_queryset(self):
        plant_id = self.kwargs['pk']
        return InventoryLocation.objects.filter(plant_id=plant_id)


# ============================================================================
# PHASE 54: ROLL TRACKING VIEWS
# ============================================================================

class RollViewSet(viewsets.ModelViewSet):
    """
    Phase 56: Roll management endpoints.
    Provides CRUD operations plus genealogy, movement, reserve, consume, and WCM filtering.
    """
    queryset = InventoryRoll.objects.all().select_related(
        'material', 'grade', 'plant', 'location', 'created_by_job', 'created_process'
    ).order_by('-created_at')
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['status', 'stage_index', 'location', 'plant', 'material', 'grade', 'is_fg']
    
    def get_serializer_class(self):
        if self.action == 'retrieve' or self.action == 'genealogy':
            return RollDetailSerializer
        return InventoryRollSerializer
    
    def get_queryset(self):
        """
        Phase 56: Enhanced filtering for WCM Terminal.
        Supports thickness range, grade, stage, location, status filters.
        """
        queryset = super().get_queryset()
        
        # Thickness range filters (for WCM)
        thickness_min = self.request.query_params.get('thickness_min')
        thickness_max = self.request.query_params.get('thickness_max')
        if thickness_min:
            queryset = queryset.filter(thickness_micron__gte=Decimal(thickness_min))
        if thickness_max:
            queryset = queryset.filter(thickness_micron__lte=Decimal(thickness_max))
        
        # Width range filters
        width_min = self.request.query_params.get('width_min')
        width_max = self.request.query_params.get('width_max')
        if width_min:
            queryset = queryset.filter(width_mm__gte=Decimal(width_min))
        if width_max:
            queryset = queryset.filter(width_mm__lte=Decimal(width_max))
        
        # Weight range filters
        weight_min = self.request.query_params.get('weight_min')
        weight_max = self.request.query_params.get('weight_max')
        if weight_min:
            queryset = queryset.filter(weight_kg__gte=Decimal(weight_min))
        if weight_max:
            queryset = queryset.filter(weight_kg__lte=Decimal(weight_max))
        
        return queryset

    @action(detail=False, methods=['get'], url_path='for-wcm')
    def for_wcm(self, request):
        """
        Phase 56: WCM Terminal roll selection endpoint.
        Returns available/reserved rolls with full filtering.
        
        Query params:
        - material: UUID (variant)
        - thickness_min, thickness_max: microns
        - grade: UUID
        - stage_index: 0-5
        - location: UUID
        - status: AVAILABLE, RESERVED
        """
        queryset = self.get_queryset().filter(
            status__in=['AVAILABLE', 'RESERVED']
        )
        
        # Apply additional filters from query params
        material_id = request.query_params.get('material')
        if material_id:
            queryset = queryset.filter(material_id=material_id)
        
        grade_id = request.query_params.get('grade')
        if grade_id:
            queryset = queryset.filter(grade_id=grade_id)
        
        stage_index = request.query_params.get('stage_index')
        if stage_index is not None:
            queryset = queryset.filter(stage_index=int(stage_index))
        
        location_id = request.query_params.get('location')
        if location_id:
            queryset = queryset.filter(location_id=location_id)
        
        plant_id = request.query_params.get('plant')
        if plant_id:
            queryset = queryset.filter(plant_id=plant_id)
        
        # Limit results for performance (after robust quarantine filtering).
        rows = [
            roll for roll in queryset
            if not bool(((getattr(roll, "meta_json", None) or {}).get("is_quarantined")))
        ][:100]

        return Response(InventoryRollSerializer(rows, many=True).data)

    @action(detail=True, methods=['post'], url_path='split')
    def split_for_claim(self, request, pk=None):
        try:
            roll = InventoryRoll.objects.select_related(
                "material",
                "grade",
                "plant",
                "location",
                "location__plant",
                "template",
                "sales_order_item",
                "created_by_job",
                "created_process",
                "production_job",
            ).get(id=pk)
        except InventoryRoll.DoesNotExist:
            return Response({"error": "Roll not found"}, status=status.HTTP_404_NOT_FOUND)

        try:
            child_weight = Decimal(str(request.data.get("child_weight_kg") or 0))
        except Exception:
            return Response({"error": "child_weight_kg must be numeric"}, status=status.HTTP_400_BAD_REQUEST)

        reason = str(request.data.get("reason") or "SALES_CLAIM").upper()
        if roll.status != "AVAILABLE":
            return Response({"error": f"Roll {roll.label_id} must be AVAILABLE before split."}, status=status.HTTP_400_BAD_REQUEST)
        if child_weight <= 0:
            return Response({"error": "child_weight_kg must be > 0"}, status=status.HTTP_400_BAD_REQUEST)
        if child_weight >= Decimal(str(roll.weight_kg or 0)):
            return Response(
                {"error": f"child_weight_kg must be less than parent weight ({roll.weight_kg} KG)."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if reason == "SALES_CLAIM" and roll.sales_order_item_id:
            return Response(
                {"error": f"Roll {roll.label_id} already has sales-order lineage. Split it before claim, not after claim."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        from apps.production.models import InventoryAllocation

        if InventoryAllocation.objects.filter(inventory_roll=roll, status="ACTIVE").exists():
            return Response(
                {"error": f"Roll {roll.label_id} already has an active allocation and cannot be split safely."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        balance_weight = Decimal(str(roll.weight_kg or 0)) - child_weight
        try:
            children = RollService.split_roll(
                parent_roll=roll,
                splits=[
                    {
                        "weight_kg": child_weight,
                        "label_suffix": "CLAIM",
                        "meta": {
                            "split_parent_id": str(roll.id),
                            "split_parent_label": roll.label_id,
                            "split_role": "CLAIM_CHILD",
                            "split_reason": reason,
                        },
                    },
                    {
                        "weight_kg": balance_weight,
                        "label_suffix": "BAL",
                        "meta": {
                            "split_parent_id": str(roll.id),
                            "split_parent_label": roll.label_id,
                            "split_role": "BALANCE_CHILD",
                            "split_reason": reason,
                        },
                    },
                ],
                reason=reason,
                user=request.user if request.user.is_authenticated else None,
            )
        except Exception as exc:
            return Response({"error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        split_children = {str((child.meta_json or {}).get("split_role") or ""): child for child in children}
        claim_child = split_children.get("CLAIM_CHILD")
        balance_child = split_children.get("BALANCE_CHILD")

        parent_meta = dict(roll.meta_json or {})
        history = list(parent_meta.get("split_history") or [])
        history.append(
            {
                "reason": reason,
                "child_weight_kg": float(child_weight),
                "balance_weight_kg": float(balance_weight),
                "claim_child_id": str(claim_child.id) if claim_child else None,
                "balance_child_id": str(balance_child.id) if balance_child else None,
            }
        )
        parent_meta["split_history"] = history
        parent_meta["last_split_reason"] = reason
        roll.meta_json = parent_meta
        roll.save(update_fields=["meta_json"])

        return Response(
            {
                "status": "split",
                "parent_roll_id": str(roll.id),
                "parent_roll_label": roll.label_id,
                "claim_roll": InventoryRollSerializer(claim_child).data if claim_child else None,
                "balance_roll": InventoryRollSerializer(balance_child).data if balance_child else None,
            },
            status=status.HTTP_201_CREATED,
        )

    @action(detail=False, methods=['get'], url_path='stock-form-operations')
    def stock_form_operations(self, request):
        """List physical roll stock-form conversions supported by inventory."""
        return Response({"operations": StockFormConversionService.operation_choices()})

    @action(detail=True, methods=['post'], url_path='convert-stock-form')
    def convert_stock_form(self, request, pk=None):
        """Convert/slit a roll while preserving genealogy and mass math."""
        roll = self.get_object()
        operation = request.data.get("operation")
        child_widths = request.data.get("child_widths_mm")
        if child_widths in (None, ""):
            child_widths = []
        if not isinstance(child_widths, list):
            return Response(
                {"error": "child_widths_mm must be a list of widths."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        try:
            result = StockFormConversionService.convert(
                roll,
                operation=operation,
                child_widths_mm=child_widths,
                trim_mm=request.data.get("trim_mm") or 0,
                reason=str(request.data.get("reason") or ""),
                user=request.user if request.user.is_authenticated else None,
            )
        except (ValueError, ValidationError) as exc:
            return Response({"error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        return Response(result, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'], url_path='availability')
    def roll_availability(self, request):
        """
        Return roll availability grouped by plant for a given spec.
        Used by WCM to request inter-plant roll transfers.
        """
        variant_id = request.query_params.get('variant_id')
        family_id = request.query_params.get('family_id')
        thickness = request.query_params.get('thickness_micron')
        grade_id = request.query_params.get('grade_id')
        min_width = request.query_params.get('min_width_mm')
        stock_form = request.query_params.get('stock_form')
        width_basis = request.query_params.get('width_basis')
        exclude_plant_id = request.query_params.get('exclude_plant')

        qs = InventoryRoll.objects.filter(
            status='AVAILABLE',
            material__isnull=False,
            thickness_micron__gt=0,
            width_mm__gt=0,
            grade__isnull=False
        ).select_related('plant', 'location', 'material', 'grade')
        if variant_id:
            qs = qs.filter(material_id=variant_id)
        elif family_id:
            qs = qs.filter(material__parent_family_id=family_id)
        if thickness:
            try:
                qs = qs.filter(Q(thickness_micron=Decimal(str(thickness))))
            except Exception:
                pass
        if grade_id:
            qs = qs.filter(Q(grade_id=grade_id))
        if min_width:
            try:
                qs = qs.filter(Q(width_mm__gte=Decimal(str(min_width))))
            except Exception:
                pass
        if stock_form:
            from apps.materials.stock_forms import normalize_stock_form, normalize_width_basis
            normalized_form = normalize_stock_form(stock_form)
            qs = qs.filter(stock_form=normalized_form)
            if width_basis:
                qs = qs.filter(width_basis=normalize_width_basis(width_basis, stock_form=normalized_form))
        if exclude_plant_id:
            qs = qs.exclude(location__plant_id=exclude_plant_id)

        grouped = {}
        for roll in qs:
            if bool(((getattr(roll, "meta_json", None) or {}).get("is_quarantined"))):
                continue
            plant_id = str(roll.location.plant_id) if roll.location else str(roll.plant_id)
            plant_name = roll.location.plant.name if roll.location and roll.location.plant else (roll.plant.name if roll.plant else None)
            if plant_id not in grouped:
                grouped[plant_id] = {
                    'plant_id': plant_id,
                    'plant_name': plant_name,
                    'total_rolls': 0,
                    'total_weight_kg': 0.0,
                    'rolls': []
                }
            grouped[plant_id]['total_rolls'] += 1
            grouped[plant_id]['total_weight_kg'] += float(roll.weight_kg or 0)
            grouped[plant_id]['rolls'].append({
                'id': str(roll.id),
                'label_id': roll.label_id,
                'material_id': str(roll.material_id) if roll.material_id else None,
                'material_name': roll.material.name if roll.material else None,
                'family_id': str(roll.material.parent_family_id) if roll.material and roll.material.parent_family_id else None,
                'weight_kg': float(roll.weight_kg or 0),
                'width_mm': float(roll.width_mm or 0),
                'stock_form': roll.stock_form,
                'width_basis': roll.width_basis,
                'thickness_micron': float(roll.thickness_micron or 0),
                'grade_id': str(roll.grade_id) if roll.grade_id else None,
                'grade_name': roll.grade.name if roll.grade else None,
                'location_id': str(roll.location_id) if roll.location_id else None,
                'location_name': roll.location.name if roll.location else None
            })

        return Response(list(grouped.values()))

    @action(detail=True, methods=['get'])
    def genealogy(self, request, pk=None):
        """Get full genealogy tree for a roll."""
        try:
            roll = self.get_object()
            genealogy = RollService.get_genealogy(roll)
            descendants = RollService.get_descendants(roll)
            return Response({
                'roll': RollDetailSerializer(roll).data,
                'ancestors': genealogy['ancestors'],
                'descendants': descendants
            })
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'])
    def move(self, request, pk=None):
        """Move a roll to a new location."""
        serializer = RollMoveSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        
        try:
            roll = self.get_object()
            to_location = InventoryLocation.objects.get(id=data['to_location_id'])
            job = None
            if data.get('job_id'):
                job = ProductionJob.objects.get(id=data['job_id'])
            
            movement = RollService.move_roll(
                roll=roll,
                to_location=to_location,
                reason=data['reason'],
                reason_note=data.get('reason_note', ''),
                job=job,
                user=request.user
            )
            return Response({
                "status": "moved",
                "movement": RollMovementSerializer(movement).data
            })
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'])
    def reserve(self, request):
        """Reserve multiple rolls for a production job."""
        serializer = RollReserveSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        
        try:
            rolls = InventoryRoll.objects.filter(id__in=data['roll_ids'])
            job = ProductionJob.objects.get(id=data['job_id'])
            
            reserved = RollService.reserve_rolls(list(rolls), job)
            return Response({
                "status": "reserved",
                "count": len(reserved),
                "rolls": InventoryRollSerializer(reserved, many=True).data
            })
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'])
    def release(self, request):
        """Release reserved rolls back to available."""
        serializer = RollReleaseSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        
        try:
            rolls = InventoryRoll.objects.filter(id__in=data['roll_ids'])
            released = RollService.release_rolls(list(rolls))
            return Response({
                "status": "released",
                "count": len(released)
            })
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='quarantine')
    def quarantine(self, request, pk=None):
        """Mark roll as quarantined (excluded from planning/auto-pick)."""
        try:
            roll = self.get_object()
            meta = dict(roll.meta_json or {})
            meta["is_quarantined"] = True
            reason = request.data.get("reason")
            if reason:
                meta["quarantine_reason"] = str(reason)
            roll.meta_json = meta
            roll.save(update_fields=["meta_json"])
            return Response({"status": "quarantined", "roll": InventoryRollSerializer(roll).data})
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='unquarantine')
    def unquarantine(self, request, pk=None):
        """Clear quarantine marker."""
        try:
            roll = self.get_object()
            meta = dict(roll.meta_json or {})
            meta["is_quarantined"] = False
            meta.pop("quarantine_reason", None)
            roll.meta_json = meta
            roll.save(update_fields=["meta_json"])
            return Response({"status": "unquarantined", "roll": InventoryRollSerializer(roll).data})
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'])
    def consume(self, request):
        """
        Consume rolls for a production job.
        Creates output, balance, and scrap rolls automatically.
        """
        serializer = RollConsumeSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        
        try:
            job = ProductionJob.objects.get(id=data['job_id'])
            process = Process.objects.get(id=data['process_id'])
            machine = Machine.objects.get(id=data['machine_id'])
            output_location = None
            if data.get('output_location_id'):
                output_location = InventoryLocation.objects.get(id=data['output_location_id'])
            
            results = []
            for input_data in data['inputs']:
                roll = InventoryRoll.objects.get(id=input_data['roll_id'])
                result = RollService.consume_roll(
                    input_roll=roll,
                    job=job,
                    process=process,
                    machine=machine,
                    used_kg=Decimal(str(input_data['used_kg'])),
                    scrap_kg=Decimal(str(input_data.get('scrap_kg', 0))),
                    output_location=output_location,
                    user=request.user,
                    notes=data.get('notes', '')
                )
                result_data = {
                    'input_roll': roll.label_id,
                    'output_roll': result['output_roll'].label_id if result['output_roll'] else None,
                    'balance_roll': result['balance_roll'].label_id if result['balance_roll'] else None,
                    'scrap_roll': result['scrap_roll'].label_id if result['scrap_roll'] else None,
                }
                results.append(result_data)
            
            return Response({
                "status": "consumed",
                "results": results
            })
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['get'], url_path='by-stage')
    def by_stage(self, request):
        """Get rolls grouped by resolved process stage for Roll Explorer tree view."""
        try:
            plant_id = request.query_params.get('plant')
            qs = InventoryRoll.objects.filter(
                status__in=['AVAILABLE', 'RESERVED']
            ).select_related('material', 'location', 'grade', 'plant', 'created_process', 'parent_roll__created_process')
            if plant_id:
                qs = qs.filter(location__plant_id=plant_id)

            stages = {
                'Raw Material': [],
                'Extruded': [],
                'Printed': [],
                'Laminated': [],
                'Slit': [],
                'Finished Good': [],
                'Remainder / Freed': [],
            }
            for roll in qs[:500]:
                stage_name = resolve_roll_stage_name(roll) or 'Raw Material'
                role = resolve_roll_role(roll)
                if role == 'REMAINDER':
                    stage_name = 'Remainder / Freed'
                if stage_name not in stages:
                    stages[stage_name] = []
                stages[stage_name].append(InventoryRollSerializer(roll).data)

            return Response(stages)
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['get'], url_path='explorer')
    def explorer(self, request):
        """
        Rich explorer payload with resolved stage + role grouping.
        Supports filters and returns both grouped buckets and flat rows.
        """
        try:
            params = request.query_params
            mode = (params.get('mode') or 'grouped').strip().lower()
            rows = _roll_query_rows(params)
            grouped = _group_rows_by_bucket(rows)

            payload = {
                'meta': {
                    'mode': mode,
                    'filters': {
                        'plant': params.get('plant'),
                        'stage': (params.get('stage') or '').strip().lower() or None,
                        'roll_role': (params.get('roll_role') or '').strip().upper() or None,
                        'origin_type': (params.get('origin_type') or '').strip().upper() or None,
                        'stock_strategy': (params.get('stock_strategy') or '').strip().upper() or None,
                        'status': params.get('status'),
                        'material': params.get('material'),
                        'grade': params.get('grade'),
                        'location': params.get('location'),
                        'family': params.get('family'),
                        'job_number': (params.get('job_number') or '').strip() or None,
                        'date_from': params.get('date_from'),
                        'date_to': params.get('date_to'),
                        'weight_min': params.get('weight_min'),
                        'weight_max': params.get('weight_max'),
                    }
                },
                'totals': grouped['totals'],
                'buckets': grouped['buckets'],
                'rows': rows if mode == 'table' else [],
            }
            return Response(payload)
        except Exception as e:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['get'], url_path='by-variant')
    def by_variant(self, request):
        """Variant-first stock view grouped as family -> variant -> rolls."""
        try:
            params = request.query_params
            rows = _roll_query_rows(params)
            grouped = _group_rows_by_bucket(rows)
            payload = {
                'meta': {
                    'filters': {
                        'plant': params.get('plant'),
                        'stage': (params.get('stage') or '').strip().lower() or None,
                        'roll_role': (params.get('roll_role') or '').strip().upper() or None,
                        'origin_type': (params.get('origin_type') or '').strip().upper() or None,
                        'stock_strategy': (params.get('stock_strategy') or '').strip().upper() or None,
                        'status': params.get('status'),
                        'material': params.get('material'),
                        'grade': params.get('grade'),
                        'location': params.get('location'),
                        'family': params.get('family'),
                        'job_number': (params.get('job_number') or '').strip() or None,
                        'date_from': params.get('date_from'),
                        'date_to': params.get('date_to'),
                        'weight_min': params.get('weight_min'),
                        'weight_max': params.get('weight_max'),
                    }
                },
                'totals': grouped['totals'],
                'families': _group_rows_by_variant(rows),
            }
            return Response(payload)
        except Exception:
            return Response({"error": "Request failed"}, status=status.HTTP_400_BAD_REQUEST)


class RollMovementListView(generics.ListAPIView):
    """List view for roll movements (movement log page)."""
    queryset = RollMovement.objects.all().select_related(
        'roll', 'from_location', 'to_location', 'job', 'moved_by'
    ).order_by('-timestamp')
    serializer_class = RollMovementSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['roll', 'from_location', 'to_location', 'reason']


class RollConsumptionListView(generics.ListAPIView):
    """List view for roll consumptions (scrap report support)."""
    queryset = RollConsumption.objects.all().select_related(
        'job', 'process', 'input_roll', 'output_roll', 'machine', 'operator'
    ).order_by('-timestamp')
    serializer_class = RollConsumptionSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['job', 'process', 'machine']


# ============================================================================
# PHASE 56: BULK INVENTORY VIEWS
# ============================================================================

from .models import InventoryBulk, BulkTransaction

class InventoryLedgerView(APIView):
    """
    Lightweight ledger endpoint (Phase 56+).

    InventoryLedger table was removed in Phase 58/0021; the UI still expects a
    unified "ledger" feed. We reconstruct it from:
    - BulkTransaction (bulk movements)
    - RollMovement (roll movements)

    Response shape matches `frontend_v2/src/app/(dashboard)/inventory/ledger/page.tsx`.
    """

    permission_classes = [IsAuthenticated]

    def get(self, request):
        tx_type_filter = (request.query_params.get("tx_type") or "ALL").upper()
        limit = int(request.query_params.get("limit") or 500)

        def include_tx(tx_type: str) -> bool:
            return tx_type_filter == "ALL" or tx_type == tx_type_filter

        items: list[dict] = []

        # ------------------------
        # Bulk transactions
        # ------------------------
        bulk_map = {
            "INWARD": "GRN",
            "CONSUME": "CONSUMPTION",
            "TRANSFER": "TRANSFER",
            "ADJUST": "ADJUST",
        }
        bulk_qs = BulkTransaction.objects.select_related(
            "material",
            "location__plant",
        ).order_by("-created_at")[:limit]

        for tx in bulk_qs:
            tx_type = bulk_map.get(tx.type, "ADJUST")
            if not include_tx(tx_type):
                continue

            items.append(
                {
                    "id": str(tx.id),
                    "tx_type": tx_type,
                    "quantity": float(tx.qty_kg or 0),
                    "uom": (tx.material.base_uom if tx.material else None) or "KG",
                    "material_name": tx.material.name if tx.material else None,
                    "item_label_id": None,
                    "from_location_name": None,
                    "to_location_name": tx.location.name if tx.location else None,
                    "plant_name": tx.location.plant.name if tx.location and tx.location.plant else None,
                    "reference": tx.reference or None,
                    "created_at": tx.created_at.isoformat() if tx.created_at else None,
                }
            )

        # ------------------------
        # Roll movements
        # ------------------------
        roll_map = {
            "GRN": "GRN",
            "PRODUCTION": "CONSUMPTION",
            "DISPATCH": "DISPATCH",
            "JOBWORK_OUT": "DISPATCH",
            "JOBWORK_IN": "RECEIVE",
            "INTER_PLANT": "TRANSFER",
            "WIP_TRANSFER": "TRANSFER",
            "FG_TRANSFER": "TRANSFER",
            "SCRAP": "CONSUMPTION",
            "ADJUSTMENT": "ADJUST",
        }
        roll_qs = RollMovement.objects.select_related(
            "roll__material",
            "from_location__plant",
            "to_location__plant",
            "job",
        ).order_by("-timestamp")[:limit]

        for mv in roll_qs:
            tx_type = roll_map.get(mv.reason, "TRANSFER")
            if not include_tx(tx_type):
                continue

            reference = mv.reason_note or ""
            if mv.job_id:
                reference = (reference + " " if reference else "") + f"JOB:{mv.job.job_number}"
            reference = reference or None

            items.append(
                {
                    "id": str(mv.id),
                    "tx_type": tx_type,
                    "quantity": float(mv.roll.weight_kg or 0) if mv.roll else 0,
                    "uom": "KG",
                    "material_name": mv.roll.material.name if mv.roll and mv.roll.material else None,
                    "item_label_id": mv.roll.label_id if mv.roll else None,
                    "from_location_name": mv.from_location.name if mv.from_location else None,
                    "to_location_name": mv.to_location.name if mv.to_location else None,
                    "plant_name": mv.to_location.plant.name if mv.to_location and mv.to_location.plant else None,
                    "reference": reference,
                    "created_at": mv.timestamp.isoformat() if mv.timestamp else None,
                }
            )

        # Newest first
        items.sort(key=lambda x: x.get("created_at") or "", reverse=True)
        return Response(items[:limit])

class BulkInventoryViewSet(viewsets.ReadOnlyModelViewSet):
    """
    Unified Bulk Inventory view (pooled).
    """
    queryset = InventoryBulk.objects.all().select_related(
        'material', 'granule_code', 'location', 'plant'
    ).order_by('material__name')
    serializer_class = InventoryBulkSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['plant', 'location', 'material', 'granule_code']

    @action(detail=False, methods=['post'])
    def inward(self, request):
        """Inward bulk material."""
        data = request.data
        try:
            tx = BulkService.add_bulk(
                material_id=data['material_id'],
                qty=Decimal(str(data['qty_kg'])),
                plant_id=data['plant_id'],
                location_id=data['location_id'],
                cost=Decimal(str(data.get('avg_cost', 0))),
                reference=data.get('reference', ''),
                granule_code_id=data.get('granule_code_id') or data.get('granule_code'),
                qty_uom=data.get('uom'),
            )
            return Response(BulkTransactionSerializer(tx).data, status=status.HTTP_201_CREATED)
        except Exception as e:
            return Response({"error": str(getattr(e, "message", "") or e)}, status=400)

    @action(detail=False, methods=['post'])
    def consume(self, request):
        """Consume bulk material."""
        data = request.data
        try:
            tx = BulkService.consume_bulk(
                material_id=data['material_id'],
                qty=Decimal(str(data['qty_kg'])),
                location_id=data['location_id'],
                job_id=data.get('job_id'),
                reference=data.get('reference', ''),
                granule_code_id=data.get('granule_code_id') or data.get('granule_code'),
                qty_uom=data.get('uom'),
            )
            return Response(BulkTransactionSerializer(tx).data)
        except Exception as e:
            return Response({"error": str(getattr(e, "message", "") or e)}, status=400)

    @action(detail=False, methods=['post'])
    def transfer(self, request):
        """Transfer bulk material between locations."""
        data = request.data
        try:
            BulkService.transfer_bulk(
                material_id=data['material_id'],
                qty=Decimal(str(data['qty_kg'])),
                from_location_id=data['from_location_id'],
                to_location_id=data['to_location_id'],
                reference=data.get('reference', ''),
                granule_code_id=data.get('granule_code_id') or data.get('granule_code'),
                qty_uom=data.get('uom'),
            )
            return Response({"status": "transferred"})
        except Exception as e:
            return Response({"error": str(getattr(e, "message", "") or e)}, status=400)

class BulkTransactionListView(generics.ListAPIView):
    """Audit log for Bulk Transactions."""
    queryset = BulkTransaction.objects.all().select_related(
        'material', 'granule_code', 'location', 'job'
    ).order_by('-created_at')
    serializer_class = BulkTransactionSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['material', 'granule_code', 'location', 'type']


class PackagingStockViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = PackagingStock.objects.all().select_related(
        'material', 'location', 'plant'
    ).order_by('material__name')
    serializer_class = PackagingStockSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['plant', 'location', 'material']


class PackagingTransactionListView(generics.ListAPIView):
    queryset = PackagingTransaction.objects.all().select_related(
        'material', 'location', 'vendor', 'job', 'sales_order_item', 'mts_order'
    ).order_by('-created_at')
    serializer_class = PackagingTransactionSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['material', 'location', 'type']


# ============================================================================
# PHASE 58: OBSERVABILITY & TRACEABILITY VIEWS
# ============================================================================

from .models import InventorySnapshot, InventoryAlert
from .serializers import InventorySnapshotSerializer, InventoryAlertSerializer, AlertResolveSerializer
from .services.inventory_audit_service import InventoryAuditService


class InventoryHealthView(generics.GenericAPIView):
    """
    GET /api/inventory/health/
    Returns current inventory health metrics, alerts summary, and shortages.
    """
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        plant_id = request.query_params.get('plant')
        plant = None
        if plant_id:
            from apps.factory.models import Plant
            try:
                plant = Plant.objects.get(id=plant_id)
            except Plant.DoesNotExist:
                return Response({"error": "Plant not found"}, status=404)
        
        try:
            health = InventoryAuditService.get_health_summary(plant)
            return Response(health)
        except Exception as e:
            return Response({"error": "Request failed"}, status=500)


class InventoryAlertViewSet(viewsets.ModelViewSet):
    """
    Inventory Alerts management.
    GET /api/inventory/alerts/
    POST /api/inventory/alerts/{id}/resolve/
    """
    queryset = InventoryAlert.objects.all().select_related(
        'material', 'roll', 'plant', 'resolved_by'
    ).order_by('-created_at')
    serializer_class = InventoryAlertSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['type', 'severity', 'resolved', 'plant']
    
    @action(detail=True, methods=['post'])
    def resolve(self, request, pk=None):
        """Mark an alert as resolved."""
        alert = self.get_object()
        if alert.resolved:
            return Response({"error": "Alert already resolved"}, status=400)
        
        serializer = AlertResolveSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        
        from django.utils import timezone
        alert.resolved = True
        alert.resolved_at = timezone.now()
        alert.resolved_by = request.user
        alert.resolution_note = serializer.validated_data.get('resolution_note', '')
        alert.save()
        
        return Response(InventoryAlertSerializer(alert).data)
    
    @action(detail=False, methods=['post'])
    def run_audit(self, request):
        """Manually trigger a reconciliation audit."""
        plant_id = request.data.get('plant_id')
        plant = None
        if plant_id:
            from apps.factory.models import Plant
            try:
                plant = Plant.objects.get(id=plant_id)
            except Plant.DoesNotExist:
                return Response({"error": "Plant not found"}, status=404)
        
        try:
            summary = InventoryAuditService.run_full_audit(plant)
            return Response(summary)
        except Exception as e:
            return Response({"error": "Request failed"}, status=500)


class InventorySnapshotViewSet(viewsets.ReadOnlyModelViewSet):
    """
    GET /api/inventory/snapshots/
    Historical inventory snapshots for trending and analysis.
    """
    queryset = InventorySnapshot.objects.all().select_related('plant').order_by('-created_at')
    serializer_class = InventorySnapshotSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['plant']

    def list(self, request, *args, **kwargs):
        queryset = self.filter_queryset(self.get_queryset())

        try:
            days = int(request.query_params.get("days", 180))
        except (TypeError, ValueError):
            days = 180
        days = max(7, min(days, 730))

        try:
            limit = int(request.query_params.get("limit", 180))
        except (TypeError, ValueError):
            limit = 180
        limit = max(1, min(limit, 500))

        cutoff = timezone.now() - timedelta(days=days)
        queryset = queryset.filter(created_at__gte=cutoff).order_by("-created_at")

        serializer = self.get_serializer(queryset[:limit], many=True)
        return Response(serializer.data)
    
    @action(detail=False, methods=['post'])
    def create_now(self, request):
        """Create a snapshot immediately."""
        plant_id = request.data.get('plant_id')
        if not plant_id:
            return Response({"error": "plant_id required"}, status=400)
        
        from apps.factory.models import Plant
        try:
            plant = Plant.objects.get(id=plant_id)
        except Plant.DoesNotExist:
            return Response({"error": "Plant not found"}, status=404)
        
        try:
            snapshot = InventoryAuditService.create_snapshot(plant)
            return Response(InventorySnapshotSerializer(snapshot).data, status=201)
        except Exception as e:
            return Response({"error": "Request failed"}, status=500)


class RollGenealogyView(generics.GenericAPIView):
    """
    GET /api/inventory/rolls/{id}/genealogy/
    Returns full genealogy tree for a roll.
    """
    permission_classes = [IsAuthenticated]
    
    def get(self, request, pk):
        try:
            genealogy = InventoryAuditService.get_roll_genealogy(str(pk))
            if 'error' in genealogy:
                return Response(genealogy, status=404)
            return Response(genealogy)
        except Exception as e:
            return Response({"error": "Request failed"}, status=500)


class RollTraceLookupView(generics.GenericAPIView):
    """
    GET /api/inventory/roll-trace/?q=<roll-id-or-label>
    Resolve by UUID or label and return rich genealogy payload.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        query = (request.query_params.get("q") or "").strip()
        if not query:
            return Response({"error": "q query param is required"}, status=400)

        roll = None
        matched_by = None

        try:
            roll_uuid = uuid.UUID(query)
            roll = InventoryRoll.objects.select_related(
                "material", "grade", "location", "location__plant", "production_job", "created_by_job", "created_process"
            ).filter(id=roll_uuid).first()
            matched_by = "id" if roll else None
        except Exception:
            roll = None

        if not roll:
            roll = InventoryRoll.objects.select_related(
                "material", "grade", "location", "location__plant", "production_job", "created_by_job", "created_process"
            ).filter(label_id__iexact=query).first()
            matched_by = "label_exact" if roll else matched_by

        if not roll:
            roll = InventoryRoll.objects.select_related(
                "material", "grade", "location", "location__plant", "production_job", "created_by_job", "created_process"
            ).filter(label_id__icontains=query).order_by("-created_at").first()
            matched_by = "label_contains" if roll else matched_by

        if not roll:
            return Response({"error": "Roll not found for provided query"}, status=404)

        genealogy = InventoryAuditService.get_roll_genealogy(str(roll.id))
        if "error" in genealogy:
            return Response(genealogy, status=404)

        stage_name = resolve_roll_stage_name(roll)
        roll_role = resolve_roll_role(roll)
        recent_movements = list(
            RollMovement.objects.filter(roll=roll).select_related("from_location", "to_location").order_by("-timestamp")[:8]
        )

        payload = {
            "query": query,
            "matched_by": matched_by,
            "roll": {
                "id": str(roll.id),
                "label_id": roll.label_id,
                "material_name": roll.material.name if roll.material else None,
                "material_code": roll.material.code if roll.material else None,
                "grade_name": roll.grade.name if roll.grade else None,
                "weight_kg": float(roll.weight_kg or 0),
                "original_weight_kg": float(roll.original_weight_kg or 0),
                "width_mm": float(roll.width_mm or 0),
                "stock_form": roll.stock_form,
                "width_basis": roll.width_basis,
                "thickness_micron": float(roll.thickness_micron or 0),
                "status": roll.status,
                "stage_name": stage_name,
                "roll_role": roll_role,
                "location_name": roll.location.name if roll.location else None,
                "plant_name": roll.location.plant.name if roll.location and roll.location.plant else None,
                "job_number": (
                    roll.production_job.job_number if roll.production_job else
                    (roll.created_by_job.job_number if roll.created_by_job else None)
                ),
                "created_at": roll.created_at.isoformat() if roll.created_at else None,
            },
            "genealogy": genealogy,
            "recent_movements": [
                {
                    "timestamp": m.timestamp.isoformat() if m.timestamp else None,
                    "from_location_name": m.from_location.name if m.from_location else None,
                    "to_location_name": m.to_location.name if m.to_location else None,
                    "reason": m.reason,
                    "reason_note": m.reason_note,
                }
                for m in recent_movements
            ],
        }
        return Response(payload)
