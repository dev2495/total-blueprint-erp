from __future__ import annotations

import hashlib
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Iterable

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from apps.analytics.models import ReportDispatchRun, ReportDistributionProfile
from apps.analytics.reports_service import ReportService
from apps.factory.models import Plant
from apps.inventory.models import InventoryBulk, InventoryRoll, InventorySnapshot, PackagingStock
from apps.inventory.serializers import resolve_roll_role, resolve_roll_stage_name
from apps.inventory.services.inventory_audit_service import InventoryAuditService
from apps.inventory.services.roll_naming import build_roll_naming_payload, build_variant_key
from apps.platformops.models import OperationalAlert
from apps.users.models import User
from apps.users.services.email_service import EmailDeliveryService

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except Exception:  # pragma: no cover
    Workbook = None
    Alignment = None
    Font = None
    PatternFill = None

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
except Exception:  # pragma: no cover
    colors = None
    A4 = None
    ParagraphStyle = None
    getSampleStyleSheet = None
    mm = None
    canvas = None
    Drawing = None
    Rect = None
    String = None
    PageBreak = None
    Paragraph = None
    SimpleDocTemplate = None
    Spacer = None
    Table = None
    TableStyle = None


DEFAULT_REPORT_PROFILES = {
    ReportDistributionProfile.ReportCode.PRODUCTION_DAILY: {
        "target_roles": ["OWNER", "ADMIN", "PLANNER", "PLANT_MANAGER", "WORK_CENTER_MANAGER"],
        "schedule_hour": 8,
        "schedule_minute": 0,
        "email_subject_template": "Daily Production Report - {report_date}",
        "email_body_template": "Attached is the daily production report for {report_date}.",
    },
    ReportDistributionProfile.ReportCode.STOCK_STANDING_DAILY: {
        "target_roles": ["OWNER", "ADMIN", "STORE", "PLANT_MANAGER"],
        "schedule_hour": 8,
        "schedule_minute": 0,
        "email_subject_template": "Daily Stock Standing Report - {report_date}",
        "email_body_template": "Attached is the daily stock standing report for {report_date}.",
    },
}


@dataclass
class RenderedAttachment:
    file_name: str
    content: bytes
    content_type: str
    checksum_sha1: str


@dataclass
class RenderedReport:
    report_code: str
    report_date: datetime.date
    pdf: bytes
    file_name: str
    checksum_sha1: str
    summary_text: str
    warning_text: str
    window_start: datetime
    window_end: datetime
    detail_attachments: list[RenderedAttachment]


def _company_name() -> str:
    return "TOTAL POLY PRINT ERP"


def _start_of_day(report_date):
    tz = timezone.get_current_timezone()
    return timezone.make_aware(datetime.combine(report_date, time.min), tz)


def _end_of_day(report_date):
    tz = timezone.get_current_timezone()
    return timezone.make_aware(datetime.combine(report_date, time.max), tz)


def _fmt_dt(value):
    if not value:
        return "-"
    local = timezone.localtime(value) if timezone.is_aware(value) else value
    return local.strftime("%d-%b-%Y %H:%M")


def _safe_number(value, default=0.0):
    try:
        return float(value or 0)
    except Exception:
        return default


def _sha1(content: bytes) -> str:
    return hashlib.sha1(content).hexdigest()


def _to_decimal(value, default: str = "0") -> Decimal:
    try:
        return Decimal(str(value if value not in (None, "") else default))
    except Exception:
        return Decimal(default)


class ReportDistributionService:
    @staticmethod
    def artifact_root() -> Path:
        base_dir = Path(getattr(settings, "BASE_DIR", "."))
        return base_dir / ".runtime" / "report-smoke"

    @staticmethod
    def persist_rendered_artifact(rendered: RenderedReport, *, folder: str | None = None) -> Path:
        paths = ReportDistributionService.persist_rendered_artifacts(rendered, folder=folder)
        return paths["pdf"]

    @staticmethod
    def persist_rendered_artifacts(rendered: RenderedReport, *, folder: str | None = None) -> dict[str, Path]:
        root = ReportDistributionService.artifact_root()
        if folder:
            root = root / folder
        else:
            root = root / rendered.report_date.isoformat()
        root.mkdir(parents=True, exist_ok=True)
        artifact_path = root / rendered.file_name
        artifact_path.write_bytes(rendered.pdf)
        paths: dict[str, Path] = {"pdf": artifact_path}
        for attachment in rendered.detail_attachments or []:
            attachment_path = root / attachment.file_name
            attachment_path.write_bytes(attachment.content)
            paths[attachment.file_name] = attachment_path
        return paths

    @staticmethod
    def ensure_defaults():
        created = []
        for report_code, defaults in DEFAULT_REPORT_PROFILES.items():
            profile, was_created = ReportDistributionProfile.objects.get_or_create(
                report_code=report_code,
                defaults=defaults,
            )
            if was_created:
                created.append(profile.report_code)
        return created

    @staticmethod
    def list_profiles():
        ReportDistributionService.ensure_defaults()
        return list(ReportDistributionProfile.objects.order_by("report_code"))

    @staticmethod
    def update_profiles(rows: Iterable[dict], updated_by=None):
        ReportDistributionService.ensure_defaults()
        updated = []
        with transaction.atomic():
            for row in rows:
                report_code = str(row.get("report_code") or "").strip()
                if report_code not in DEFAULT_REPORT_PROFILES:
                    continue
                profile = ReportDistributionProfile.objects.filter(report_code=report_code).first()
                if not profile:
                    profile = ReportDistributionProfile(report_code=report_code)
                target_roles = [str(role or "").upper() for role in (row.get("target_roles") or []) if str(role or "").strip()]
                extra_recipients = [str(email or "").strip() for email in (row.get("extra_recipients") or []) if str(email or "").strip()]
                profile.active = bool(row.get("active", True))
                profile.target_roles = target_roles
                profile.extra_recipients = extra_recipients
                profile.schedule_hour = max(0, min(23, int(row.get("schedule_hour", profile.schedule_hour or 8) or 8)))
                profile.schedule_minute = max(0, min(59, int(row.get("schedule_minute", profile.schedule_minute or 0) or 0)))
                profile.email_subject_template = str(
                    row.get("email_subject_template")
                    or profile.email_subject_template
                    or DEFAULT_REPORT_PROFILES[report_code]["email_subject_template"]
                )
                profile.email_body_template = str(
                    row.get("email_body_template")
                    or profile.email_body_template
                    or DEFAULT_REPORT_PROFILES[report_code]["email_body_template"]
                )
                profile.updated_by = updated_by
                profile.save()
                updated.append(profile)
        return updated

    @staticmethod
    def list_runs(limit=30):
        return list(ReportDispatchRun.objects.select_related("profile", "triggered_by").order_by("-created_at")[:limit])

    @staticmethod
    def get_run(run_id: str):
        return (
            ReportDispatchRun.objects.select_related("profile", "triggered_by")
            .filter(id=run_id)
            .first()
        )

    @staticmethod
    def serialize_profile(profile: ReportDistributionProfile) -> dict:
        defaults = DEFAULT_REPORT_PROFILES.get(profile.report_code, {})
        return {
            "report_code": profile.report_code,
            "label": dict(ReportDistributionProfile.ReportCode.choices).get(profile.report_code, profile.report_code),
            "active": bool(profile.active),
            "target_roles": list(profile.target_roles or []),
            "extra_recipients": list(profile.extra_recipients or []),
            "schedule_hour": int(profile.schedule_hour or 0),
            "schedule_minute": int(profile.schedule_minute or 0),
            "email_subject_template": profile.email_subject_template or defaults.get("email_subject_template", ""),
            "email_body_template": profile.email_body_template or defaults.get("email_body_template", ""),
            "updated_at": profile.updated_at.isoformat() if profile.updated_at else None,
            "updated_by": getattr(profile.updated_by, "username", None),
        }

    @staticmethod
    def serialize_run(run: ReportDispatchRun) -> dict:
        return {
            "id": str(run.id),
            "report_code": run.report_code,
            "report_date": run.report_date.isoformat() if run.report_date else None,
            "status": run.status,
            "recipients": list(run.recipients or []),
            "recipient_count": int(run.recipient_count or 0),
            "warning_text": run.warning_text or "",
            "error_text": run.error_text or "",
            "provider": run.provider or "",
            "provider_message_id": run.provider_message_id or "",
            "pdf_file_name": run.pdf_file_name or "",
            "pdf_checksum_sha1": run.pdf_checksum_sha1 or "",
            "pdf_size_bytes": int(run.pdf_size_bytes or 0),
            "detail_file_name": run.detail_file_name or "",
            "detail_checksum_sha1": run.detail_checksum_sha1 or "",
            "detail_size_bytes": int(run.detail_size_bytes or 0),
            "triggered_manually": bool(run.triggered_manually),
            "triggered_by": getattr(run.triggered_by, "username", None),
            "window_start": run.window_start.isoformat() if run.window_start else None,
            "window_end": run.window_end.isoformat() if run.window_end else None,
            "created_at": run.created_at.isoformat() if run.created_at else None,
            "sent_at": run.sent_at.isoformat() if run.sent_at else None,
        }

    @staticmethod
    def recipients_for_profile(profile: ReportDistributionProfile):
        role_recipients = list(
            User.objects.filter(role__code__in=profile.target_roles)
            .exclude(email="")
            .values_list("email", flat=True)
        )
        deduped = []
        seen = set()
        for email in [*role_recipients, *(profile.extra_recipients or [])]:
            normalized = str(email or "").strip().lower()
            if not normalized or normalized in seen:
                continue
            seen.add(normalized)
            deduped.append(normalized)
        return deduped

    @staticmethod
    def report_date_for_run(explicit_date=None):
        if explicit_date:
            return explicit_date
        return timezone.localdate() - timedelta(days=1)

    @staticmethod
    def render_report(report_code: str, report_date=None) -> RenderedReport:
        if canvas is None:
            raise RuntimeError("PDF engine unavailable: reportlab is not installed.")
        report_date = ReportDistributionService.report_date_for_run(report_date)
        if report_code == ReportDistributionProfile.ReportCode.PRODUCTION_DAILY:
            return _DailyProductionPDFRenderer.render(report_date)
        if report_code == ReportDistributionProfile.ReportCode.STOCK_STANDING_DAILY:
            return _StockStandingPDFRenderer.render(report_date)
        raise ValueError("Unsupported report_code.")

    @staticmethod
    def send_profile(profile: ReportDistributionProfile, *, report_date=None, triggered_by=None, triggered_manually=False):
        rendered = ReportDistributionService.render_report(profile.report_code, report_date=report_date)
        recipients = ReportDistributionService.recipients_for_profile(profile)
        warning_text = rendered.warning_text or ""
        detail_attachment = rendered.detail_attachments[0] if rendered.detail_attachments else None
        run = ReportDispatchRun.objects.create(
            profile=profile,
            report_code=profile.report_code,
            report_date=rendered.report_date,
            window_start=rendered.window_start,
            window_end=rendered.window_end,
            recipients=recipients,
            recipient_count=len(recipients),
            warning_text=warning_text,
            pdf_file_name=rendered.file_name,
            pdf_checksum_sha1=rendered.checksum_sha1,
            pdf_size_bytes=len(rendered.pdf),
            detail_file_name=detail_attachment.file_name if detail_attachment else "",
            detail_checksum_sha1=detail_attachment.checksum_sha1 if detail_attachment else "",
            detail_size_bytes=len(detail_attachment.content) if detail_attachment else 0,
            triggered_by=triggered_by,
            triggered_manually=triggered_manually,
        )
        ReportDistributionService.persist_rendered_artifacts(rendered, folder=f"runs/{run.id}")
        if not recipients:
            run.status = ReportDispatchRun.Status.SKIPPED
            run.warning_text = (warning_text + "\nNo recipients configured.").strip()
            run.save(update_fields=["status", "warning_text"])
            return run

        email_ready, email_warning = EmailDeliveryService.configuration_status()
        if not email_ready:
            run.status = ReportDispatchRun.Status.SKIPPED_EMAIL
            run.warning_text = "\n".join(filter(None, [warning_text, email_warning])).strip()
            run.provider = str(getattr(settings, "EMAIL_PROVIDER", "") or "")
            run.sent_at = timezone.now()
            run.save(update_fields=["status", "warning_text", "provider", "sent_at"])
            from apps.users.services.notification_service import NotificationService

            NotificationService.emit_event(
                event_key="reports.daily_pack_email_skipped",
                title=f"Report generated without email: {profile.report_code}",
                message=f"{profile.report_code} for {rendered.report_date.isoformat()} generated successfully, but email delivery was skipped. {email_warning}",
                notification_type="SYSTEM",
                related_object_type="ReportDispatchRun",
                related_object_id=str(run.id),
                priority="NORMAL",
                idempotency_key=f"report-email-skipped:{profile.report_code}:{rendered.report_date.isoformat()}",
            )
            return run

        subject_template = profile.email_subject_template or DEFAULT_REPORT_PROFILES[profile.report_code]["email_subject_template"]
        body_template = profile.email_body_template or DEFAULT_REPORT_PROFILES[profile.report_code]["email_body_template"]
        context = {
            "report_date": rendered.report_date.isoformat(),
            "company_name": _company_name(),
            "summary_text": rendered.summary_text,
        }
        subject = subject_template.format(**context)
        body = (
            body_template.format(**context)
            + "<br/><br/>"
            + rendered.summary_text.replace("\n", "<br/>")
        )
        try:
            result = EmailDeliveryService.send_email(
                subject=subject,
                body=body,
                recipients=recipients,
                idempotency_key=f"{profile.report_code}:{rendered.report_date.isoformat()}:{'manual' if triggered_manually else 'scheduled'}",
                attachments=[
                    {
                        "filename": rendered.file_name,
                        "content": rendered.pdf,
                        "content_type": "application/pdf",
                    },
                    *[
                        {
                            "filename": attachment.file_name,
                            "content": attachment.content,
                            "content_type": attachment.content_type,
                        }
                        for attachment in rendered.detail_attachments
                    ],
                ],
            )
        except Exception as exc:
            run.status = ReportDispatchRun.Status.FAILED
            run.error_text = str(exc)
            run.save(update_fields=["status", "error_text"])
            OperationalAlert.objects.create(
                category="REPORTING",
                severity=OperationalAlert.Severity.WARNING,
                message="Daily report delivery failed",
                details={
                    "report_code": profile.report_code,
                    "report_date": rendered.report_date.isoformat(),
                    "error": str(exc),
                },
            )
            from apps.users.services.notification_service import NotificationService

            NotificationService.emit_event(
                event_key="reports.daily_pack_failed",
                title=f"Report delivery failed: {profile.report_code}",
                message=f"{profile.report_code} for {rendered.report_date.isoformat()} failed to send. Review reporting audit logs.",
                notification_type="SYSTEM",
                related_object_type="ReportDispatchRun",
                related_object_id=str(run.id),
                priority="HIGH",
                idempotency_key=f"report-failed:{profile.report_code}:{rendered.report_date.isoformat()}",
            )
            raise

        run.status = ReportDispatchRun.Status.SUCCEEDED
        run.provider = str(result.get("provider") or "")
        run.provider_message_id = str(result.get("provider_message_id") or "")
        run.sent_at = timezone.now()
        run.save(update_fields=["status", "provider", "provider_message_id", "sent_at"])
        from apps.users.services.notification_service import NotificationService

        NotificationService.emit_event(
            event_key="reports.daily_pack_sent",
            title=f"Report sent: {profile.report_code}",
            message=f"{profile.report_code} for {rendered.report_date.isoformat()} was sent to {len(recipients)} recipients.",
            notification_type="SYSTEM",
            related_object_type="ReportDispatchRun",
            related_object_id=str(run.id),
            priority="LOW",
            idempotency_key=f"report-sent:{profile.report_code}:{rendered.report_date.isoformat()}",
        )
        return run


class _BaseDailyPDFRenderer:
    report_title = ""
    report_code = ""

    @classmethod
    def _new_pdf(cls):
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=A4)
        return buffer, pdf

    @classmethod
    def _draw_title(cls, pdf, report_date, subtitle):
        width, height = A4
        pdf.setFillColor(colors.HexColor("#0F172A"))
        pdf.rect(10 * mm, height - 24 * mm, 190 * mm, 16 * mm, fill=1, stroke=0)
        pdf.setFillColor(colors.white)
        pdf.setFont("Helvetica-Bold", 15)
        pdf.drawString(14 * mm, height - 16 * mm, _company_name())
        pdf.setFont("Helvetica-Bold", 11)
        pdf.drawRightString(196 * mm, height - 16 * mm, cls.report_title)
        pdf.setFillColor(colors.HexColor("#334155"))
        pdf.setFont("Helvetica", 8.5)
        pdf.drawString(14 * mm, height - 29 * mm, f"Report Date: {report_date.strftime('%d-%b-%Y')}")
        pdf.drawRightString(196 * mm, height - 29 * mm, f"Generated: {_fmt_dt(timezone.now())}")
        pdf.setFont("Helvetica", 8)
        pdf.drawString(14 * mm, height - 34 * mm, subtitle)
        return height - 42 * mm

    @classmethod
    def _draw_summary_cards(cls, pdf, y, cards: list[tuple[str, str]]):
        width = A4[0]
        x = 12 * mm
        card_width = (width - 24 * mm - (len(cards) - 1) * 4 * mm) / max(len(cards), 1)
        for label, value in cards:
            pdf.setFillColor(colors.HexColor("#F8FAFC"))
            pdf.setStrokeColor(colors.HexColor("#CBD5E1"))
            pdf.roundRect(x, y - 14 * mm, card_width, 12 * mm, 2 * mm, fill=1, stroke=1)
            pdf.setFillColor(colors.HexColor("#64748B"))
            pdf.setFont("Helvetica-Bold", 7)
            pdf.drawString(x + 3 * mm, y - 5 * mm, label.upper())
            pdf.setFillColor(colors.HexColor("#0F172A"))
            pdf.setFont("Helvetica-Bold", 12)
            pdf.drawString(x + 3 * mm, y - 10 * mm, value)
            x += card_width + 4 * mm
        return y - 18 * mm

    @classmethod
    def _draw_table(cls, pdf, y, title, headers, rows):
        width, height = A4
        pdf.setFillColor(colors.HexColor("#0F172A"))
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawString(12 * mm, y, title)
        y -= 4 * mm
        pdf.setFillColor(colors.HexColor("#E2E8F0"))
        pdf.rect(10 * mm, y - 6 * mm, 190 * mm, 7 * mm, fill=1, stroke=0)
        pdf.setFillColor(colors.HexColor("#0F172A"))
        pdf.setFont("Helvetica-Bold", 7.2)
        positions = [12 * mm, 68 * mm, 108 * mm, 142 * mm, 172 * mm]
        for index, header in enumerate(headers[: len(positions)]):
            pdf.drawString(positions[index], y - 2 * mm, header)
        y -= 9 * mm
        pdf.setFont("Helvetica", 7.4)
        for row in rows:
            if y < 20 * mm:
                pdf.showPage()
                y = height - 18 * mm
            for index, value in enumerate(row[: len(positions)]):
                pdf.drawString(positions[index], y, str(value)[:30])
            y -= 4.5 * mm
        return y - 2 * mm


class _DailyProductionPDFRenderer(_BaseDailyPDFRenderer):
    report_title = "DAILY PRODUCTION REPORT"
    report_code = ReportDistributionProfile.ReportCode.PRODUCTION_DAILY

    @classmethod
    def render(cls, report_date):
        filters = {"date_from": report_date.isoformat(), "date_to": report_date.isoformat()}
        payload = ReportService.get_report_tab("production", filters)
        summary = payload.get("summary") or {}
        machine_rows = payload.get("breakdown") or payload.get("rows") or []
        process_rows = payload.get("by_process") or []
        window_start = _start_of_day(report_date)
        window_end = _end_of_day(report_date)

        buffer, pdf = cls._new_pdf()
        y = cls._draw_title(
            pdf,
            report_date,
            "Plant-wise previous-day output, yield, scrap, downtime, and release risk highlights.",
        )
        y = cls._draw_summary_cards(
            pdf,
            y,
            [
                ("Total Output", f"{_safe_number(summary.get('total_output_kg')):,.1f} kg"),
                ("Yield", f"{_safe_number(summary.get('yield_pct')):,.1f}%"),
                ("Scrap", f"{_safe_number(summary.get('scrap_rate')):,.1f}%"),
                ("Completion", f"{_safe_number(summary.get('completion_rate')):,.1f}%"),
            ],
        )
        y = cls._draw_table(
            pdf,
            y,
            "Top Machines",
            ["Machine", "Output", "Scrap", "Yield", "Earned Hrs"],
            [
                [
                    row.get("machine", "-"),
                    f"{_safe_number(row.get('actual_output')):,.1f} kg",
                    f"{_safe_number(row.get('scrap_kg')):,.1f} kg",
                    f"{_safe_number(row.get('yield_pct')):,.1f}%",
                    f"{_safe_number(row.get('earned_hours')):,.1f}",
                ]
                for row in machine_rows[:12]
            ] or [["No machine activity", "-", "-", "-", "-"]],
        )
        y = cls._draw_table(
            pdf,
            y,
            "Top Processes",
            ["Process", "Output", "", "", ""],
            [
                [row.get("name", "-"), f"{_safe_number(row.get('value')):,.1f} kg", "", "", ""]
                for row in process_rows[:10]
            ] or [["No process data", "-", "", "", ""]],
        )
        pdf.showPage()
        pdf.save()
        payload_bytes = buffer.getvalue()
        checksum = hashlib.sha1(payload_bytes).hexdigest()
        return RenderedReport(
            report_code=cls.report_code,
            report_date=report_date,
            pdf=payload_bytes,
            file_name=f"production-daily-{report_date.isoformat()}.pdf",
            checksum_sha1=checksum,
            summary_text=(
                f"Output: {_safe_number(summary.get('total_output_kg')):,.1f} kg\n"
                f"Yield: {_safe_number(summary.get('yield_pct')):,.1f}%\n"
                f"Scrap: {_safe_number(summary.get('scrap_rate')):,.1f}%\n"
                f"Completion: {_safe_number(summary.get('completion_rate')):,.1f}%"
            ),
            warning_text="",
            window_start=window_start,
            window_end=window_end,
            detail_attachments=[],
        )


class _StockStandingPDFRenderer(_BaseDailyPDFRenderer):
    report_title = "DAILY STOCK STANDING REPORT"
    report_code = ReportDistributionProfile.ReportCode.STOCK_STANDING_DAILY

    SECTION_ORDER = [
        "In-house FG rolls",
        "Intermediate pool / semi-FG rolls",
        "Printed rolls",
        "Laminated rolls",
        "Plain / extruded rolls",
        "Purchased / external inbound rolls",
        "Blocked / aged / exception stock",
    ]

    @classmethod
    def _official_plants(cls):
        return list(Plant.objects.filter(include_in_official_reports=True).order_by("name"))

    @classmethod
    def _latest_snapshots(cls, report_date):
        plant_ids = [plant.id for plant in cls._official_plants()]
        if not plant_ids:
            return []
        snapshot_qs = (
            InventorySnapshot.objects.filter(created_at__date=report_date, plant_id__in=plant_ids)
            .select_related("plant")
            .order_by("plant__name", "-created_at")
        )
        if snapshot_qs.exists():
            latest_by_plant = {}
            for snapshot in snapshot_qs:
                latest_by_plant.setdefault(str(snapshot.plant_id), snapshot)
            return list(latest_by_plant.values())

        created = []
        for plant in cls._official_plants():
            created.append(InventoryAuditService.create_snapshot(plant))
        return created

    @classmethod
    def _print_status(cls, stage_name: str) -> str:
        stage = str(stage_name or "").lower()
        return "Printed" if stage in {"printed", "laminated", "slit", "finished good"} else "Plain"

    @classmethod
    def _lamination_status(cls, stage_name: str) -> str:
        stage = str(stage_name or "").lower()
        return "Laminated" if stage in {"laminated", "slit", "finished good"} else "Not laminated"

    @classmethod
    def _stock_strategy(cls, roll: InventoryRoll) -> str:
        meta = roll.meta_json or {}
        strategy = str(
            meta.get("stock_strategy")
            or getattr(getattr(roll.created_by_job, "mts_order", None), "stock_strategy", "")
            or ("PACKAGING_STOCK" if str(getattr(roll.material, "category", "") or "").upper() == "PACKAGING" else "")
            or ("FINAL_STOCK" if bool(getattr(roll, "is_fg", False)) else "INTERMEDIATE_POOL")
        ).upper()
        if strategy not in {"FINAL_STOCK", "INTERMEDIATE_POOL", "PACKAGING_STOCK"}:
            return "FINAL_STOCK" if bool(getattr(roll, "is_fg", False)) else "INTERMEDIATE_POOL"
        return strategy

    @classmethod
    def _origin_type(cls, roll: InventoryRoll, role: str) -> str:
        meta = roll.meta_json or {}
        explicit = str(meta.get("origin_type") or "").upper()
        if explicit in {"IN_HOUSE", "PURCHASED", "JOBWORK_RETURN", "INTERPLANT_IN", "REMAINDER"}:
            return explicit
        if role == "REMAINDER":
            return "REMAINDER"
        if str(getattr(getattr(roll, "created_by_job", None), "source_type", "") or "").upper() == "JOBWORK_RETURN":
            return "JOBWORK_RETURN"
        if meta.get("interplant_challan_id") or meta.get("from_plant_id") or meta.get("received_from_challan"):
            return "INTERPLANT_IN"
        if roll.created_by_job_id or roll.production_job_id:
            return "IN_HOUSE"
        if str(roll.status or "").upper() == "SENT_JOBWORK":
            return "JOBWORK_RETURN"
        return "PURCHASED"

    @classmethod
    def _roll_entries(cls, report_date):
        plants = cls._official_plants()
        plant_ids = [plant.id for plant in plants]
        if not plant_ids:
            return []
        qs = (
            InventoryRoll.objects.filter(location__plant_id__in=plant_ids)
            .exclude(status__in=["CONSUMED", "SCRAPPED"])
            .select_related(
                "material",
                "material__parent_family",
                "grade",
                "plant",
                "location",
                "location__plant",
                "template",
                "created_by_job",
                "production_job",
                "created_process",
                "parent_roll",
            )
            .order_by("location__plant__name", "location__name", "label_id")
        )
        rows = []
        for roll in qs:
            role = resolve_roll_role(roll) or ""
            stage_name = resolve_roll_stage_name(roll) or "Raw Material"
            naming = build_roll_naming_payload(roll, role=role, stage_name=stage_name)
            width_mm = _safe_number(roll.width_mm)
            thickness_micron = _safe_number(roll.thickness_micron)
            plant_name = (
                getattr(getattr(roll.location, "plant", None), "name", "")
                or getattr(getattr(roll, "plant", None), "name", "")
                or "-"
            )
            location_name = getattr(getattr(roll, "location", None), "name", "") or "-"
            age_days = max((report_date - roll.created_at.date()).days, 0) if roll.created_at else 0
            available_kg = _safe_number(roll.weight_kg if roll.status == "AVAILABLE" else 0)
            reserved_kg = _safe_number(roll.weight_kg if roll.status == "RESERVED" else 0)
            pouch_style = str(getattr(getattr(roll, "template", None), "pouch_style", "") or "")
            rows.append(
                {
                    "roll_id": str(roll.id),
                    "label_id": roll.label_id,
                    "family_display_name": naming["family_display_name"],
                    "variant_display_name": naming["variant_display_name"],
                    "size_line": naming["size_line"],
                    "form_label": naming["form_label"],
                    "process_state_label": naming["process_state_label"],
                    "pouch_or_roll_form": pouch_style or naming["form_label"] or (getattr(getattr(roll, "template", None), "fg_type", "") or "ROLL"),
                    "stage": stage_name,
                    "width_mm": width_mm,
                    "thickness_micron": thickness_micron,
                    "print_status": naming["print_status"],
                    "lamination_status": naming["lamination_status"],
                    "stock_strategy": naming["stock_strategy"],
                    "stock_strategy_label": naming["stock_strategy_label"],
                    "plant": plant_name,
                    "location": location_name,
                    "status": str(roll.status or "").upper(),
                    "origin_type": naming["origin_type"],
                    "origin_label": naming["origin_label"],
                    "reporting_group": naming["reporting_group"],
                    "roll_count": 1,
                    "available_kg": available_kg,
                    "reserved_kg": reserved_kg,
                    "blocked_kg": _safe_number(roll.weight_kg if str(roll.status or "").upper() not in {"AVAILABLE", "RESERVED"} else 0),
                    "weight_kg": _safe_number(roll.weight_kg),
                    "oldest_age_days": age_days,
                    "is_exception": bool((roll.meta_json or {}).get("is_quarantined")) or str(roll.status or "").upper() not in {"AVAILABLE", "RESERVED"},
                }
            )
        return rows

    @classmethod
    def _variant_rows(cls, roll_entries: list[dict]) -> list[dict]:
        grouped: dict[tuple, dict] = {}
        for row in roll_entries:
            key = build_variant_key(row)
            bucket = grouped.setdefault(
                key,
                {
                    "family_display_name": row["family_display_name"],
                    "variant_display_name": row["variant_display_name"],
                    "size_line": row["size_line"],
                    "form_label": row["form_label"],
                    "process_state_label": row["process_state_label"],
                    "pouch_or_roll_form": row["pouch_or_roll_form"],
                    "stage": row["stage"],
                    "width_mm": row["width_mm"],
                    "thickness_micron": row["thickness_micron"],
                    "print_status": row["print_status"],
                    "lamination_status": row["lamination_status"],
                    "stock_strategy": row["stock_strategy"],
                    "stock_strategy_label": row["stock_strategy_label"],
                    "reporting_group": row["reporting_group"],
                    "roll_count": 0,
                    "available_kg": 0.0,
                    "reserved_kg": 0.0,
                    "blocked_kg": 0.0,
                    "oldest_age_days": 0,
                    "plant_summary": defaultdict(lambda: {"plant": "", "locations": set(), "available_kg": 0.0, "reserved_kg": 0.0, "blocked_kg": 0.0}),
                },
            )
            bucket["roll_count"] += 1
            bucket["available_kg"] += row["available_kg"]
            bucket["reserved_kg"] += row["reserved_kg"]
            bucket["blocked_kg"] += row.get("blocked_kg", 0)
            bucket["oldest_age_days"] = max(bucket["oldest_age_days"], row["oldest_age_days"])
            plant_bucket = bucket["plant_summary"][row["plant"]]
            plant_bucket["plant"] = row["plant"]
            if row["location"]:
                plant_bucket["locations"].add(row["location"])
            plant_bucket["available_kg"] += row["available_kg"]
            plant_bucket["reserved_kg"] += row["reserved_kg"]
            plant_bucket["blocked_kg"] += row.get("blocked_kg", 0)

        variant_rows = []
        for bucket in grouped.values():
            plant_summary = []
            for plant in bucket["plant_summary"].values():
                plant_summary.append(
                    {
                        "plant": plant["plant"],
                        "locations": sorted(plant["locations"]),
                        "available_kg": round(plant["available_kg"], 3),
                        "reserved_kg": round(plant["reserved_kg"], 3),
                        "blocked_kg": round(plant["blocked_kg"], 3),
                    }
                )
            bucket["plant_summary"] = plant_summary
            bucket["plant_location_summary"] = " • ".join(
                f"{plant['plant']}: {', '.join(plant['locations']) if plant['locations'] else '-'}"
                for plant in plant_summary
            )
            variant_rows.append(bucket)

        return sorted(variant_rows, key=lambda row: (row["family_display_name"], row["variant_display_name"], row["stage"], row["stock_strategy"]))

    @classmethod
    def _family_rows(cls, variant_rows: list[dict]) -> list[dict]:
        grouped: dict[tuple, dict] = {}
        for row in variant_rows:
            key = (
                row["family_display_name"],
                row["form_label"],
                row["reporting_group"],
            )
            bucket = grouped.setdefault(
                key,
                {
                    "family_display_name": row["family_display_name"],
                    "form_label": row["form_label"],
                    "reporting_group": row["reporting_group"],
                    "variant_count": 0,
                    "roll_count": 0,
                    "available_kg": 0.0,
                    "reserved_kg": 0.0,
                    "blocked_kg": 0.0,
                    "oldest_age_days": 0,
                },
            )
            bucket["variant_count"] += 1
            bucket["roll_count"] += row["roll_count"]
            bucket["available_kg"] += row["available_kg"]
            bucket["reserved_kg"] += row["reserved_kg"]
            bucket["blocked_kg"] += row["blocked_kg"]
            bucket["oldest_age_days"] = max(bucket["oldest_age_days"], row["oldest_age_days"])
        return sorted(grouped.values(), key=lambda row: (row["family_display_name"], row["reporting_group"]))

    @classmethod
    def _stage_rows(cls, roll_entries: list[dict]) -> list[dict]:
        grouped: dict[str, dict] = defaultdict(lambda: {"stage": "", "weight_kg": 0.0, "roll_count": 0})
        for row in roll_entries:
            grouped[row["stage"]]["stage"] = row["stage"]
            grouped[row["stage"]]["weight_kg"] += row["weight_kg"]
            grouped[row["stage"]]["roll_count"] += 1
        return sorted(grouped.values(), key=lambda row: row["stage"])

    @classmethod
    def _strategy_rows(cls, roll_entries: list[dict]) -> list[dict]:
        grouped: dict[str, dict] = defaultdict(lambda: {"stock_strategy": "", "weight_kg": 0.0, "roll_count": 0})
        for row in roll_entries:
            grouped[row["stock_strategy"]]["stock_strategy"] = row["stock_strategy"]
            grouped[row["stock_strategy"]]["weight_kg"] += row["weight_kg"]
            grouped[row["stock_strategy"]]["roll_count"] += 1
        return sorted(grouped.values(), key=lambda row: row["stock_strategy"])

    @classmethod
    def _bulk_rows(cls, model):
        plants = cls._official_plants()
        plant_ids = [plant.id for plant in plants]
        if not plant_ids:
            return []
        rows = []
        qs = (
            model.objects.filter(plant_id__in=plant_ids)
            .exclude(**({"qty_kg__lte": 0} if model is InventoryBulk else {"qty__lte": 0}))
            .select_related("material", "plant", "location")
            .order_by("plant__name", "location__name", "material__name")
        )
        for row in qs:
            qty_value = _safe_number(row.qty_kg if model is InventoryBulk else row.qty)
            avg_cost = _safe_number(getattr(row, "avg_cost", 0))
            rows.append(
                {
                    "material": getattr(row.material, "name", "-"),
                    "category": getattr(row.material, "category", "-"),
                    "plant": getattr(row.plant, "name", "-"),
                    "location": getattr(row.location, "name", "-"),
                    "qty": qty_value,
                    "uom": "KG" if model is InventoryBulk else getattr(row.material, "base_uom", "QTY"),
                    "avg_cost": avg_cost,
                    "value": qty_value * avg_cost,
                }
            )
        return rows

    @classmethod
    def _section_rows(cls, roll_entries: list[dict]) -> dict[str, list[dict]]:
        sections = {label: [] for label in cls.SECTION_ORDER}
        for row in roll_entries:
            stage = str(row["stage"]).lower()
            origin = str(row["origin_type"]).upper()
            strategy = str(row["stock_strategy"]).upper()
            is_exception = bool(row["is_exception"]) or row["oldest_age_days"] >= 90
            if strategy == "FINAL_STOCK" and origin == "IN_HOUSE":
                sections["In-house FG rolls"].append(row)
            if strategy == "INTERMEDIATE_POOL":
                sections["Intermediate pool / semi-FG rolls"].append(row)
            if stage == "printed":
                sections["Printed rolls"].append(row)
            if stage == "laminated":
                sections["Laminated rolls"].append(row)
            if stage in {"raw material", "extruded", "slit"} and origin == "IN_HOUSE":
                sections["Plain / extruded rolls"].append(row)
            if origin in {"PURCHASED", "JOBWORK_RETURN", "INTERPLANT_IN"}:
                sections["Purchased / external inbound rolls"].append(row)
            if is_exception:
                sections["Blocked / aged / exception stock"].append(row)
        return {key: cls._variant_rows(value) for key, value in sections.items() if value}

    @classmethod
    def _location_rows(cls, roll_entries: list[dict]) -> list[dict]:
        grouped: dict[tuple, dict] = {}
        for row in roll_entries:
            key = (row["plant"], row["location"])
            bucket = grouped.setdefault(key, {"plant": row["plant"], "location": row["location"], "weight_kg": 0.0, "roll_count": 0})
            bucket["weight_kg"] += row["weight_kg"]
            bucket["roll_count"] += 1
        return sorted(grouped.values(), key=lambda row: (-row["weight_kg"], row["plant"], row["location"]))

    @classmethod
    def _risk_callouts(cls, roll_entries: list[dict], snapshots: list[InventorySnapshot]) -> list[str]:
        callouts = []
        aged_90 = sum(row["weight_kg"] for row in roll_entries if row["oldest_age_days"] >= 90)
        blocked = sum(row["weight_kg"] for row in roll_entries if row["is_exception"])
        reserved = sum(_safe_number(snapshot.reserved_roll_kg) for snapshot in snapshots)
        if aged_90 > 0:
            callouts.append(f"Aged stock above 90 days: {aged_90:,.1f} kg")
        if blocked > 0:
            callouts.append(f"Blocked or exception stock: {blocked:,.1f} kg")
        if reserved > 0:
            callouts.append(f"Reserved roll stock: {reserved:,.1f} kg")
        return callouts or ["No material stock risks exceeded the configured operational thresholds."]

    @classmethod
    def _build_detail_workbook(cls, report_date, snapshots, family_rows, roll_entries, variant_rows, bulk_rows, packaging_rows, exception_rows) -> RenderedAttachment | None:
        if Workbook is None:
            return None

        def add_sheet_rows(ws, rows, headers):
            ws.append(headers)
            for header_cell in ws[1]:
                if Font and PatternFill and Alignment:
                    header_cell.font = Font(bold=True, color="FFFFFF")
                    header_cell.fill = PatternFill(fill_type="solid", fgColor="0F172A")
                    header_cell.alignment = Alignment(horizontal="center")
            ws.freeze_panes = "A2"
            for row in rows:
                ws.append([row.get(header, "") for header in headers])
            for col in ws.columns:
                width = max(len(str(cell.value or "")) for cell in col[: min(len(col), 25)]) + 2
                ws.column_dimensions[col[0].column_letter].width = min(max(width, 12), 28)

        workbook = Workbook()
        summary_ws = workbook.active
        summary_ws.title = "Summary"
        summary_rows = [
            {
                "Plant": snapshot.plant.name,
                "Bulk KG": round(_safe_number(snapshot.total_bulk_kg), 3),
                "Roll KG": round(_safe_number(snapshot.total_roll_kg), 3),
                "FG KG": round(_safe_number(snapshot.total_fg_kg), 3),
                "WIP KG": round(_safe_number(snapshot.total_wip_kg), 3),
                "Reserved KG": round(_safe_number(snapshot.reserved_roll_kg), 3),
            }
            for snapshot in snapshots
        ]
        add_sheet_rows(summary_ws, summary_rows, ["Plant", "Bulk KG", "Roll KG", "FG KG", "WIP KG", "Reserved KG"])

        family_ws = workbook.create_sheet("Family Summary")
        add_sheet_rows(
            family_ws,
            [
                {
                    "Business Family": row["family_display_name"],
                    "Form": row["form_label"],
                    "Reporting Group": row["reporting_group"],
                    "Variant Count": row["variant_count"],
                    "Roll Count": row["roll_count"],
                    "Available KG": round(row["available_kg"], 3),
                    "Reserved KG": round(row["reserved_kg"], 3),
                    "Blocked KG": round(row["blocked_kg"], 3),
                    "Oldest Age Days": row["oldest_age_days"],
                }
                for row in family_rows
            ],
            ["Business Family", "Form", "Reporting Group", "Variant Count", "Roll Count", "Available KG", "Reserved KG", "Blocked KG", "Oldest Age Days"],
        )

        variants_ws = workbook.create_sheet("Variant Summary")
        add_sheet_rows(
            variants_ws,
            [
                {
                    "Business Family": row["family_display_name"],
                    "Variant": row["variant_display_name"],
                    "Size": row["size_line"],
                    "Form": row["form_label"],
                    "Stage": row["stage"],
                    "Width MM": round(row["width_mm"], 2),
                    "Thickness Micron": round(row["thickness_micron"], 2),
                    "Print": row["print_status"],
                    "Lamination": row["lamination_status"],
                    "Strategy": row["stock_strategy_label"],
                    "Roll Count": row["roll_count"],
                    "Available KG": round(row["available_kg"], 3),
                    "Reserved KG": round(row["reserved_kg"], 3),
                    "Blocked KG": round(row["blocked_kg"], 3),
                    "Where": row["plant_location_summary"],
                    "Oldest Age Days": row["oldest_age_days"],
                }
                for row in variant_rows
            ],
            ["Business Family", "Variant", "Size", "Form", "Stage", "Width MM", "Thickness Micron", "Print", "Lamination", "Strategy", "Roll Count", "Available KG", "Reserved KG", "Blocked KG", "Where", "Oldest Age Days"],
        )

        detail_ws = workbook.create_sheet("Roll Detail")
        add_sheet_rows(
            detail_ws,
            [
                {
                    "Label ID": row["label_id"],
                    "Business Family": row["family_display_name"],
                    "Variant": row["variant_display_name"],
                    "Size": row["size_line"],
                    "Form": row["form_label"],
                    "Stage": row["stage"],
                    "Origin": row["origin_label"],
                    "Strategy": row["stock_strategy_label"],
                    "Plant": row["plant"],
                    "Location": row["location"],
                    "Status": row["status"],
                    "Weight KG": round(row["weight_kg"], 3),
                    "Available KG": round(row["available_kg"], 3),
                    "Reserved KG": round(row["reserved_kg"], 3),
                    "Blocked KG": round(row["blocked_kg"], 3),
                    "Age Days": row["oldest_age_days"],
                }
                for row in roll_entries
            ],
            ["Label ID", "Business Family", "Variant", "Size", "Form", "Stage", "Origin", "Strategy", "Plant", "Location", "Status", "Weight KG", "Available KG", "Reserved KG", "Blocked KG", "Age Days"],
        )

        bulk_ws = workbook.create_sheet("Bulk RM")
        add_sheet_rows(
            bulk_ws,
            [
                {
                    "Material": row["material"],
                    "Category": row["category"],
                    "Plant": row["plant"],
                    "Location": row["location"],
                    "Qty": round(row["qty"], 3),
                    "UOM": row["uom"],
                    "Avg Cost": round(row["avg_cost"], 3),
                    "Value": round(row["value"], 3),
                }
                for row in [*bulk_rows, *packaging_rows]
            ],
            ["Material", "Category", "Plant", "Location", "Qty", "UOM", "Avg Cost", "Value"],
        )

        exceptions_ws = workbook.create_sheet("Exceptions")
        add_sheet_rows(
            exceptions_ws,
            [
                {
                    "Label ID": row["label_id"],
                    "Variant": row["variant_display_name"],
                    "Business Family": row["family_display_name"],
                    "Stage": row["stage"],
                    "Plant": row["plant"],
                    "Location": row["location"],
                    "Status": row["status"],
                    "Weight KG": round(row["weight_kg"], 3),
                    "Age Days": row["oldest_age_days"],
                }
                for row in exception_rows
            ],
            ["Label ID", "Business Family", "Variant", "Stage", "Plant", "Location", "Status", "Weight KG", "Age Days"],
        )

        buffer = BytesIO()
        workbook.save(buffer)
        content = buffer.getvalue()
        return RenderedAttachment(
            file_name=f"stock-standing-detail-{report_date.isoformat()}.xlsx",
            content=content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            checksum_sha1=_sha1(content),
        )

    @classmethod
    def _build_pdf(cls, report_date, snapshots, stage_rows, strategy_rows, family_rows, variant_rows, location_rows, summary_cards, risk_callouts):
        if SimpleDocTemplate is None:
            raise RuntimeError("PDF engine unavailable: reportlab platypus is not installed.")
        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=10 * mm,
            rightMargin=10 * mm,
            topMargin=12 * mm,
            bottomMargin=10 * mm,
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("ReportTitle", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=22, textColor=colors.HexColor("#FFFFFF"), leading=26)
        section_style = ParagraphStyle("SectionTitle", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, textColor=colors.HexColor("#0F172A"), spaceAfter=6)
        body_style = ParagraphStyle("Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.5, textColor=colors.HexColor("#334155"), leading=11)
        small_style = ParagraphStyle("Small", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.5, textColor=colors.HexColor("#475569"), leading=9)
        tiny_style = ParagraphStyle("Tiny", parent=styles["BodyText"], fontName="Helvetica", fontSize=6.8, textColor=colors.HexColor("#475569"), leading=8.3)
        card_label_style = ParagraphStyle("CardLabel", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=7.5, textColor=colors.HexColor("#64748B"), leading=9)
        card_value_style = ParagraphStyle("CardValue", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=15, textColor=colors.HexColor("#0F172A"), leading=17)
        hero_title_style = ParagraphStyle("HeroTitle", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=18, textColor=colors.HexColor("#0F172A"), leading=20)
        chip_style = ParagraphStyle("ChipStyle", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=7, textColor=colors.HexColor("#0F172A"), leading=8)

        def styled_table(rows, col_widths, header_bg="#E2E8F0", body_font_size=7.2, repeat_rows=1):
            table = Table(rows, colWidths=col_widths, repeatRows=repeat_rows)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(header_bg)),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), body_font_size),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            return table

        def metric_card(label: str, value: str, accent: str):
            card = Table(
                [
                    [Paragraph(label.upper(), card_label_style)],
                    [Paragraph(value, card_value_style)],
                ],
                colWidths=[44 * mm],
            )
            card.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#CBD5E1")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 8),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                        ("LINEABOVE", (0, 0), (-1, 0), 4, colors.HexColor(accent)),
                    ]
                )
            )
            return card

        def label_chip(text: str, fill: str, text_color: str = "#0F172A"):
            chip = Table([[Paragraph(text.upper(), chip_style)]], colWidths=[None])
            chip.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(fill)),
                        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor(text_color)),
                        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor(fill)),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                    ]
                )
            )
            return chip

        def horizontal_bar_chart(title: str, rows: list[dict], label_key: str, value_key: str, color_hex: str, suffix: str = "kg"):
            if not rows or Drawing is None or Rect is None or String is None:
                return styled_table(
                    [[title, "Value"]] + [[str(row.get(label_key, "-")), f"{_safe_number(row.get(value_key)):,.1f} {suffix}"] for row in rows],
                    [80 * mm, 28 * mm],
                    repeat_rows=1,
                )

            limited_rows = rows[:6]
            max_value = max((_safe_number(row.get(value_key)) for row in limited_rows), default=1.0) or 1.0
            drawing = Drawing(180, 18 + (len(limited_rows) * 20))
            y = drawing.height - 20
            for row in limited_rows:
                label = str(row.get(label_key, "-"))[:24]
                value = _safe_number(row.get(value_key))
                width = max(6, (value / max_value) * 74) if value > 0 else 0
                drawing.add(String(0, y + 3, label, fontName="Helvetica-Bold", fontSize=7.2, fillColor=colors.HexColor("#334155")))
                drawing.add(Rect(72, y, 74, 10, fillColor=colors.HexColor("#E2E8F0"), strokeColor=colors.HexColor("#CBD5E1"), strokeWidth=0.4))
                if width:
                    drawing.add(Rect(72, y, width, 10, fillColor=colors.HexColor(color_hex), strokeColor=colors.HexColor(color_hex), strokeWidth=0))
                drawing.add(String(151, y + 2, f"{value:,.1f} {suffix}", fontName="Helvetica-Bold", fontSize=7.2, fillColor=colors.HexColor("#0F172A")))
                y -= 18
            return Table(
                [
                    [Paragraph(title, section_style)],
                    [drawing],
                ],
                colWidths=[87 * mm],
                style=TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#E2E8F0")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 8),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                        ("TOPPADDING", (0, 0), (-1, -1), 8),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                    ]
                ),
            )

        def family_hero(family: dict, rows: list[dict]):
            chips = [
                label_chip(family["reporting_group"], "#E2E8F0"),
                label_chip(f"{int(family['variant_count'])} variants", "#DBEAFE", "#1D4ED8"),
                label_chip(f"{int(family['roll_count'])} rolls", "#E0F2FE", "#0369A1"),
            ]
            if any(str(row.get("stock_strategy", "")).upper() == "FINAL_STOCK" for row in rows):
                chips.append(label_chip("final stock", "#DCFCE7", "#166534"))
            if any(str(row.get("stock_strategy", "")).upper() == "INTERMEDIATE_POOL" for row in rows):
                chips.append(label_chip("intermediate", "#FEF3C7", "#92400E"))
            if any("PRINT" in str(row.get("print_status", "")).upper() for row in rows):
                chips.append(label_chip("printed", "#DBEAFE", "#1D4ED8"))
            if any("LAMIN" in str(row.get("lamination_status", "")).upper() for row in rows):
                chips.append(label_chip("laminated", "#EDE9FE", "#6D28D9"))

            hero_table = Table(
                [
                    [
                        Table(
                            [
                                [Paragraph(family["family_display_name"], hero_title_style)],
                                [Paragraph(f"Available {family['available_kg']:,.1f} kg • Reserved {family['reserved_kg']:,.1f} kg • Blocked {family['blocked_kg']:,.1f} kg", body_style)],
                            ],
                            colWidths=[92 * mm],
                            style=TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]),
                        ),
                        Table(
                            [
                                [Paragraph("Oldest stock", card_label_style), Paragraph(f"{int(family['oldest_age_days'])} days", ParagraphStyle("HeroValue", parent=card_value_style, fontSize=13, leading=15))],
                                [Paragraph("Reporting group", card_label_style), Paragraph(str(family["reporting_group"]), body_style)],
                            ],
                            colWidths=[25 * mm, 28 * mm],
                            style=TableStyle(
                                [
                                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                                ]
                            ),
                        ),
                    ],
                    [Table([chips], style=TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)])), ""],
                ],
                colWidths=[112 * mm, 53 * mm],
            )
            hero_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                        ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#CBD5E1")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 10),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                        ("TOPPADDING", (0, 0), (-1, -1), 8),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                        ("SPAN", (0, 1), (1, 1)),
                    ]
                )
            )
            return hero_table

        def variant_card(row: dict):
            chips = [
                label_chip(str(row["stage"]), "#E2E8F0"),
                label_chip(str(row["print_status"]), "#DBEAFE", "#1D4ED8"),
                label_chip(str(row["lamination_status"]), "#EDE9FE", "#6D28D9"),
                label_chip(str(row["stock_strategy_label"]), "#DCFCE7" if str(row["stock_strategy"]).upper() == "FINAL_STOCK" else "#FEF3C7", "#166534" if str(row["stock_strategy"]).upper() == "FINAL_STOCK" else "#92400E"),
            ]
            location_text = row["plant_location_summary"] or "-"
            card = Table(
                [
                    [Paragraph(row["variant_display_name"], ParagraphStyle("VariantTitle", parent=body_style, fontName="Helvetica-Bold", fontSize=9.4, textColor=colors.HexColor("#0F172A"), leading=11))],
                    [Paragraph(row["size_line"], small_style)],
                    [Table([chips[:2], chips[2:]], style=TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))],
                    [styled_table(
                        [
                            ["Rolls", "Avail", "Reserved", "Blocked", "Oldest"],
                            [str(int(row["roll_count"])), f"{row['available_kg']:,.1f}", f"{row['reserved_kg']:,.1f}", f"{row['blocked_kg']:,.1f}", f"{int(row['oldest_age_days'])}d"],
                        ],
                        [12 * mm, 18 * mm, 18 * mm, 18 * mm, 14 * mm],
                        header_bg="#F8FAFC",
                        body_font_size=7.2,
                        repeat_rows=0,
                    )],
                    [Paragraph(f"Where: {location_text[:120]}", tiny_style)],
                ],
                colWidths=[82 * mm],
            )
            card.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 7),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            return card

        def family_location_strip(rows: list[dict]):
            plant_map: dict[str, dict] = {}
            for row in rows:
                for plant_summary in row.get("plant_summary", []):
                    bucket = plant_map.setdefault(
                        plant_summary["plant"],
                        {"plant": plant_summary["plant"], "available_kg": 0.0, "reserved_kg": 0.0, "blocked_kg": 0.0, "locations": set()},
                    )
                    bucket["available_kg"] += _safe_number(plant_summary["available_kg"])
                    bucket["reserved_kg"] += _safe_number(plant_summary["reserved_kg"])
                    bucket["blocked_kg"] += _safe_number(plant_summary["blocked_kg"])
                    bucket["locations"].update(plant_summary.get("locations") or [])
            strip_rows = sorted(plant_map.values(), key=lambda item: (-item["available_kg"], item["plant"]))[:6]
            if not strip_rows:
                return Paragraph("No plant/location detail available.", small_style)
            return styled_table(
                [["Plant", "Locations", "Avail KG", "Reserved KG", "Blocked KG"]]
                + [
                    [
                        row["plant"],
                        ", ".join(sorted(row["locations"]))[:42] or "-",
                        f"{row['available_kg']:,.1f}",
                        f"{row['reserved_kg']:,.1f}",
                        f"{row['blocked_kg']:,.1f}",
                    ]
                    for row in strip_rows
                ],
                [33 * mm, 66 * mm, 19 * mm, 21 * mm, 20 * mm],
                header_bg="#F8FAFC",
                body_font_size=6.8,
            )

        elements = [
            styled_table(
                [[Paragraph(_company_name(), title_style), Paragraph(cls.report_title, title_style)]],
                [95 * mm, 95 * mm],
                header_bg="#0F172A",
                body_font_size=11,
                repeat_rows=0,
            ),
            Spacer(1, 4 * mm),
            Paragraph(f"Report Date: {report_date.strftime('%d-%b-%Y')}<br/>Generated: {_fmt_dt(timezone.now())}", body_style),
            Paragraph("Official plant stock position with visual family pages, stock mix charts, and a compact audit appendix.", body_style),
            Spacer(1, 4 * mm),
            Table(
                [[metric_card(label, value, accent) for (label, value), accent in zip(summary_cards, ["#2563EB", "#10B981", "#D97706", "#7C3AED"], strict=False)]],
                colWidths=[47 * mm, 47 * mm, 47 * mm, 47 * mm],
                style=TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]),
            ),
            Spacer(1, 4 * mm),
            Paragraph("Official Plant Snapshot", section_style),
            styled_table(
                [["Plant", "Bulk KG", "Roll KG", "FG KG", "WIP KG", "Reserved KG"]]
                + [
                    [
                        snapshot.plant.name,
                        f"{_safe_number(snapshot.total_bulk_kg):,.1f}",
                        f"{_safe_number(snapshot.total_roll_kg):,.1f}",
                        f"{_safe_number(snapshot.total_fg_kg):,.1f}",
                        f"{_safe_number(snapshot.total_wip_kg):,.1f}",
                        f"{_safe_number(snapshot.reserved_roll_kg):,.1f}",
                    ]
                    for snapshot in snapshots
                ],
                [45 * mm, 24 * mm, 24 * mm, 24 * mm, 24 * mm, 24 * mm],
            ),
            Spacer(1, 4 * mm),
            Paragraph("Risk Callouts", section_style),
            *[Paragraph(f"- {item}", body_style) for item in risk_callouts],
            PageBreak(),
            Table(
                [[
                    horizontal_bar_chart("Stage mix", stage_rows, "stage", "weight_kg", "#2563EB"),
                    horizontal_bar_chart("Strategy mix", strategy_rows, "stock_strategy", "weight_kg", "#D97706"),
                ]],
                colWidths=[90 * mm, 90 * mm],
                style=TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]),
            ),
            Spacer(1, 4 * mm),
            Table(
                [[
                    horizontal_bar_chart("Family contribution", family_rows, "family_display_name", "available_kg", "#7C3AED"),
                    horizontal_bar_chart("Blocked / aging watch", [{"label": "Blocked kg", "value": sum(_safe_number(row.get("blocked_kg")) for row in family_rows)}, {"label": "Aged 90+ kg", "value": sum(_safe_number(row.get("weight_kg")) for row in variant_rows if row.get("oldest_age_days", 0) >= 90)}], "label", "value", "#DC2626"),
                ]],
                colWidths=[90 * mm, 90 * mm],
                style=TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]),
            ),
            Spacer(1, 4 * mm),
            Paragraph("Top Locations", section_style),
            styled_table(
                [["Plant", "Location", "Weight KG", "Roll Count"]]
                + [[row["plant"], row["location"], f"{row['weight_kg']:,.1f}", str(int(row["roll_count"]))] for row in location_rows[:20]],
                [50 * mm, 70 * mm, 30 * mm, 30 * mm],
            ),
        ]

        for family in family_rows:
            rows = [
                row for row in variant_rows
                if row["family_display_name"] == family["family_display_name"]
                and row["reporting_group"] == family["reporting_group"]
            ]
            if not rows:
                continue
            elements.append(PageBreak())
            elements.append(family_hero(family, rows))
            elements.append(Spacer(1, 3 * mm))
            elements.append(Paragraph("Size / variant view", section_style))
            card_rows = []
            pending_cards = [variant_card(row) for row in rows]
            for index in range(0, len(pending_cards), 2):
                left = pending_cards[index]
                right = pending_cards[index + 1] if index + 1 < len(pending_cards) else ""
                card_rows.append([left, right])
            elements.append(
                Table(
                    card_rows,
                    colWidths=[87 * mm, 87 * mm],
                    style=TableStyle(
                        [
                            ("LEFTPADDING", (0, 0), (-1, -1), 0),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                            ("TOPPADDING", (0, 0), (-1, -1), 0),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ]
                    ),
                )
            )
            elements.append(Spacer(1, 2 * mm))
            elements.append(Paragraph("Plant / location strip", section_style))
            elements.append(family_location_strip(rows))

        elements.append(PageBreak())
        elements.append(Paragraph("Audit appendix", section_style))
        appendix_rows = [[
            "Family",
            "Variant",
            "Size",
            "Stage",
            "Strategy",
            "Rolls",
            "Avail KG",
            "Res KG",
            "Blocked KG",
            "Where",
        ]]
        for row in variant_rows:
            appendix_rows.append(
                [
                    row["family_display_name"],
                    row["variant_display_name"][:24],
                    row["size_line"][:22],
                    row["stage"],
                    row["stock_strategy_label"],
                    str(int(row["roll_count"])),
                    f"{row['available_kg']:.1f}",
                    f"{row['reserved_kg']:.1f}",
                    f"{row['blocked_kg']:.1f}",
                    row["plant_location_summary"][:42],
                ]
            )
        elements.append(
            styled_table(
                appendix_rows,
                [28 * mm, 28 * mm, 20 * mm, 16 * mm, 18 * mm, 10 * mm, 14 * mm, 14 * mm, 16 * mm, 34 * mm],
                body_font_size=6.3,
            )
        )

        doc.build(elements)
        return buffer.getvalue()

    @classmethod
    def render(cls, report_date):
        snapshots = cls._latest_snapshots(report_date)
        roll_entries = cls._roll_entries(report_date)
        variant_rows = cls._variant_rows(roll_entries)
        family_rows = cls._family_rows(variant_rows)
        stage_rows = cls._stage_rows(roll_entries)
        strategy_rows = cls._strategy_rows(roll_entries)
        location_rows = cls._location_rows(roll_entries)
        bulk_rows = cls._bulk_rows(InventoryBulk)
        packaging_rows = cls._bulk_rows(PackagingStock)
        exception_rows = [row for row in roll_entries if row["is_exception"] or row["oldest_age_days"] >= 90]
        warning_text = "" if snapshots else "No inventory snapshots were available."
        if not cls._official_plants():
            warning_text = "\n".join(filter(None, [warning_text, "No plants are marked for official reports."]))
        window_start = _start_of_day(report_date)
        window_end = _end_of_day(report_date)
        total_roll_kg = sum(_safe_number(snapshot.total_roll_kg) for snapshot in snapshots)
        total_bulk_kg = sum(_safe_number(snapshot.total_bulk_kg) for snapshot in snapshots)
        total_reserved_kg = sum(_safe_number(snapshot.reserved_roll_kg) for snapshot in snapshots)
        aged_90_kg = sum(row["weight_kg"] for row in roll_entries if row["oldest_age_days"] >= 90)
        summary_cards = [
            ("Roll stock", f"{total_roll_kg:,.1f} kg"),
            ("Bulk stock", f"{total_bulk_kg:,.1f} kg"),
            ("Reserved", f"{total_reserved_kg:,.1f} kg"),
            ("Aged 90+", f"{aged_90_kg:,.1f} kg"),
        ]
        pdf_bytes = cls._build_pdf(
            report_date,
            snapshots,
            stage_rows,
            strategy_rows,
            family_rows,
            variant_rows,
            location_rows,
            summary_cards,
            cls._risk_callouts(roll_entries, snapshots),
        )
        checksum = _sha1(pdf_bytes)
        detail_attachment = cls._build_detail_workbook(
            report_date,
            snapshots,
            family_rows,
            roll_entries,
            variant_rows,
            bulk_rows,
            packaging_rows,
            exception_rows,
        )
        return RenderedReport(
            report_code=cls.report_code,
            report_date=report_date,
            pdf=pdf_bytes,
            file_name=f"stock-standing-daily-{report_date.isoformat()}.pdf",
            checksum_sha1=checksum,
            summary_text=(
                f"Roll stock: {total_roll_kg:,.1f} kg\n"
                f"Bulk stock: {total_bulk_kg:,.1f} kg\n"
                f"Reserved rolls: {total_reserved_kg:,.1f} kg\n"
                f"Aged 90+ stock: {aged_90_kg:,.1f} kg"
            ),
            warning_text=warning_text,
            window_start=window_start,
            window_end=window_end,
            detail_attachments=[detail_attachment] if detail_attachment else [],
        )
