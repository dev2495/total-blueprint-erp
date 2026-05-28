import tempfile
from datetime import date
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from apps.analytics.report_delivery import _StockStandingPDFRenderer
from apps.analytics.models import ReportDispatchRun
from apps.analytics.report_delivery import RenderedAttachment, RenderedReport, ReportDistributionService
from apps.factory.models import Plant
from apps.users.models import Role, User

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


class ReportDeliveryTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.sales_role = Role.objects.create(
            code="SALES",
            name="Sales",
            default_permissions=["analytics.view"],
        )
        self.owner = User.objects.create_user(
            username="report-owner",
            password="pass1234",
            email="owner@example.com",
            is_owner=True,
            is_staff=True,
        )
        self.admin = User.objects.create_user(
            username="report-admin",
            password="pass1234",
            email="admin@example.com",
            is_staff=True,
        )
        self.sales_user = User.objects.create_user(
            username="report-sales",
            password="pass1234",
            email="sales@example.com",
            role=self.sales_role,
        )
        self.client.force_authenticate(self.owner)
        self.production_profile = next(
            profile
            for profile in ReportDistributionService.list_profiles()
            if profile.report_code == "production_daily"
        )

    def test_list_profiles_exposes_all_five_daily_packs(self):
        profiles = ReportDistributionService.list_profiles()
        self.assertEqual(
            {profile.report_code for profile in profiles},
            {
                "owner_executive_daily",
                "production_daily",
                "dispatch_daily",
                "packing_dispatch_summary_daily",
                "stock_standing_daily",
            },
        )

    def _rendered(self, report_code: str = "production_daily") -> RenderedReport:
        window_end = timezone.now()
        window_start = window_end - timezone.timedelta(days=1)
        return RenderedReport(
            report_code=report_code,
            report_date=date(2026, 3, 5),
            pdf=b"%PDF-1.4 mock report",
            file_name=f"{report_code}-2026-03-05.pdf",
            checksum_sha1="abc123",
            summary_text="Plant A stable\nPlant B stable",
            warning_text="",
            window_start=window_start,
            window_end=window_end,
            detail_attachments=[],
        )

    def test_persist_rendered_artifact_writes_pdf_for_local_review(self):
        rendered = self._rendered()
        with tempfile.TemporaryDirectory() as tmp_dir:
            with patch.object(ReportDistributionService, "artifact_root", return_value=Path(tmp_dir)):
                artifact_path = ReportDistributionService.persist_rendered_artifact(rendered, folder="manual-preview")
                self.assertTrue(str(artifact_path).endswith(rendered.file_name))
                self.assertEqual(artifact_path.read_bytes(), rendered.pdf)

    @patch("apps.analytics.report_delivery.ReportDistributionService.render_report")
    @patch("apps.users.services.notification_service.NotificationService.emit_event")
    def test_send_profile_creates_audited_run_and_owner_notification(self, emit_event, render_report):
        render_report.return_value = self._rendered()
        self.production_profile.target_roles = ["OWNER", "ADMIN"]
        self.production_profile.extra_recipients = []
        self.production_profile.save(update_fields=["target_roles", "extra_recipients"])

        with tempfile.TemporaryDirectory() as tmp_dir:
            with patch.object(ReportDistributionService, "artifact_root", return_value=Path(tmp_dir)):
                run = ReportDistributionService.send_profile(
                    self.production_profile,
                    report_date=date(2026, 3, 5),
                    triggered_by=self.owner,
                    triggered_manually=True,
                )

        self.assertEqual(run.status, ReportDispatchRun.Status.SUCCEEDED)
        self.assertEqual(run.recipient_count, 2)
        self.assertEqual(run.recipients, ["OWNER", "ADMIN"])
        self.assertEqual(run.provider, "")
        self.assertEqual(run.provider_message_id, "")
        emit_event.assert_called_once()
        self.assertEqual(emit_event.call_args.kwargs["event_key"], "reports.daily_pack_generated")
        self.assertEqual(emit_event.call_args.kwargs["related_object_id"], str(run.id))

    @patch("apps.analytics.report_delivery.ReportDistributionService.render_report")
    @patch("apps.users.services.notification_service.NotificationService.emit_event")
    def test_send_profile_keeps_detail_attachment_and_succeeds_without_email_provider(self, emit_event, render_report):
        rendered = self._rendered("stock_standing_daily")
        rendered.detail_attachments = [
            RenderedAttachment(
                file_name="stock-standing-detail-2026-03-05.xlsx",
                content=b"detail workbook",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                checksum_sha1="detail123",
            )
        ]
        render_report.return_value = rendered

        with tempfile.TemporaryDirectory() as tmp_dir:
            with patch.object(ReportDistributionService, "artifact_root", return_value=Path(tmp_dir)):
                run = ReportDistributionService.send_profile(self.production_profile, report_date=date(2026, 3, 5))

        self.assertEqual(run.status, ReportDispatchRun.Status.SUCCEEDED)
        self.assertEqual(run.detail_file_name, "stock-standing-detail-2026-03-05.xlsx")
        emit_event.assert_called_once()

    @patch("apps.analytics.views.ReportDistributionService.send_profile")
    @patch("apps.analytics.views.ReportDistributionService.render_report")
    def test_manual_send_and_preview_endpoints_share_report_pipeline(self, render_report, send_profile):
        rendered_preview = self._rendered()
        rendered_detail = self._rendered("stock_standing_daily")
        rendered_detail.detail_attachments = [
            RenderedAttachment(
                file_name="stock-standing-detail-2026-03-05.xlsx",
                content=b"detail workbook",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                checksum_sha1="detail123",
            )
        ]
        render_report.side_effect = [rendered_preview, rendered_detail]
        run = ReportDispatchRun.objects.create(
            profile=self.production_profile,
            report_code=self.production_profile.report_code,
            report_date=date(2026, 3, 5),
            window_start=timezone.now() - timezone.timedelta(days=1),
            window_end=timezone.now(),
            recipients=["OWNER", "ADMIN"],
            recipient_count=1,
            pdf_file_name="production_daily-2026-03-05.pdf",
            pdf_checksum_sha1="abc123",
            pdf_size_bytes=18,
            detail_file_name="stock-standing-detail-2026-03-05.xlsx",
            detail_checksum_sha1="detail123",
            detail_size_bytes=12,
            status=ReportDispatchRun.Status.SUCCEEDED,
            triggered_by=self.owner,
            triggered_manually=True,
        )
        send_profile.return_value = run

        send_response = self.client.post("/api/analytics/report-distributions/production_daily/send/", {"report_date": "2026-03-05"}, format="json")
        self.assertEqual(send_response.status_code, 200)
        self.assertEqual(send_response.json()["run"]["id"], str(run.id))

        preview_response = self.client.get(f"/api/analytics/report-runs/{run.id}/preview-pdf/")
        self.assertEqual(preview_response.status_code, 200)
        self.assertEqual(preview_response["Content-Type"], "application/pdf")
        self.assertIn("production_daily-2026-03-05.pdf", preview_response["Content-Disposition"])

        run.report_code = "stock_standing_daily"
        run.detail_file_name = "stock-standing-detail-2026-03-05.xlsx"
        run.save(update_fields=["report_code", "detail_file_name"])
        detail_response = self.client.get(f"/api/analytics/report-runs/{run.id}/download-detail/")
        self.assertEqual(detail_response.status_code, 200)
        self.assertIn("stock-standing-detail-2026-03-05.xlsx", detail_response["Content-Disposition"])

    def test_non_admin_users_cannot_read_report_delivery_endpoints(self):
        run = ReportDispatchRun.objects.create(
            profile=self.production_profile,
            report_code=self.production_profile.report_code,
            report_date=date(2026, 3, 5),
            window_start=timezone.now() - timezone.timedelta(days=1),
            window_end=timezone.now(),
            recipients=["OWNER", "ADMIN"],
            recipient_count=1,
            pdf_file_name="production_daily-2026-03-05.pdf",
            pdf_checksum_sha1="abc123",
            pdf_size_bytes=18,
            detail_file_name="stock-standing-detail-2026-03-05.xlsx",
            detail_checksum_sha1="detail123",
            detail_size_bytes=12,
            status=ReportDispatchRun.Status.SUCCEEDED,
            triggered_by=self.owner,
            triggered_manually=True,
        )
        self.client.force_authenticate(self.sales_user)

        distributions_response = self.client.get("/api/analytics/report-distributions/")
        self.assertEqual(distributions_response.status_code, 403, distributions_response.content)

        runs_response = self.client.get("/api/analytics/report-runs/")
        self.assertEqual(runs_response.status_code, 403, runs_response.content)

        with patch("apps.analytics.views.ReportDistributionService.render_report") as render_report:
            preview_response = self.client.get(f"/api/analytics/report-runs/{run.id}/preview-pdf/")
        self.assertEqual(preview_response.status_code, 403, preview_response.content)
        render_report.assert_not_called()

    def test_owner_can_read_report_delivery_lists(self):
        run = ReportDispatchRun.objects.create(
            profile=self.production_profile,
            report_code=self.production_profile.report_code,
            report_date=date(2026, 3, 5),
            window_start=timezone.now() - timezone.timedelta(days=1),
            window_end=timezone.now(),
            recipients=["owner@example.com"],
            recipient_count=1,
            pdf_file_name="production_daily-2026-03-05.pdf",
            pdf_checksum_sha1="abc123",
            pdf_size_bytes=18,
            status=ReportDispatchRun.Status.SUCCEEDED,
            triggered_by=self.owner,
            triggered_manually=True,
        )
        self.client.force_authenticate(self.owner)

        distributions_response = self.client.get("/api/analytics/report-distributions/")
        self.assertEqual(distributions_response.status_code, 200, distributions_response.content)
        self.assertGreaterEqual(len(distributions_response.json().get("profiles", [])), 1)

        runs_response = self.client.get("/api/analytics/report-runs/")
        self.assertEqual(runs_response.status_code, 200, runs_response.content)
        self.assertEqual(runs_response.json()["runs"][0]["id"], str(run.id))

    def test_capability_matrix_endpoint_exposes_registry_sections(self):
        response = self.client.get("/api/analytics/capability-matrix/")

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertIn("sections", payload)
        self.assertEqual(
            {section["label"] for section in payload["sections"]},
            {"Supported Today", "Config Only Extensions", "Needs New Physical Logic"},
        )

    @patch("apps.analytics.report_delivery.InventoryAuditService.create_snapshot")
    def test_stock_standing_snapshot_bootstrap_uses_only_official_plants(self, create_snapshot):
        plant_a = Plant.objects.create(code="PLANT_A", name="Plant A", include_in_official_reports=True)
        Plant.objects.create(code="PLANT_B", name="UI E2E Plant", include_in_official_reports=False)
        create_snapshot.side_effect = [
            SimpleNamespace(plant=plant_a),
        ]

        snapshots = _StockStandingPDFRenderer._latest_snapshots(date(2026, 3, 5))

        self.assertEqual(len(snapshots), 1)
        self.assertEqual(create_snapshot.call_count, 1)

    def test_stock_standing_detail_workbook_contains_operational_tabs(self):
        if load_workbook is None:
            self.skipTest("openpyxl not available")

        attachment = _StockStandingPDFRenderer._build_detail_workbook(
            date(2026, 3, 5),
            snapshots=[
                SimpleNamespace(
                    plant=SimpleNamespace(name="Plant A"),
                    total_bulk_kg=120,
                    total_roll_kg=80,
                    total_fg_kg=40,
                    total_wip_kg=20,
                    reserved_roll_kg=5,
                )
            ],
            roll_entries=[
                {
                    "label_id": "ROLL-001",
                    "family_display_name": "PET Printed Laminate",
                    "variant_display_name": "PET/PE · 12x50",
                    "size_line": "420 mm x 62 micron",
                    "form_label": "Pouch",
                    "pouch_or_roll_form": "STAND_UP",
                    "stage": "Laminated",
                    "origin_label": "In-house made",
                    "stock_strategy_label": "Intermediate pool",
                    "plant": "Plant A",
                    "location": "LAM Store",
                    "status": "AVAILABLE",
                    "weight_kg": 12.5,
                    "available_kg": 12.5,
                    "reserved_kg": 0,
                    "blocked_kg": 0,
                    "oldest_age_days": 8,
                }
            ],
            family_rows=[
                {
                    "family_display_name": "PET Printed Laminate",
                    "form_label": "Pouch",
                    "reporting_group": "LAMINATED",
                    "variant_count": 1,
                    "roll_count": 1,
                    "available_kg": 12.5,
                    "reserved_kg": 0,
                    "blocked_kg": 0,
                    "oldest_age_days": 8,
                }
            ],
            variant_rows=[
                {
                    "family_display_name": "PET Printed Laminate",
                    "variant_display_name": "PET/PE · 12x50",
                    "size_line": "420 mm x 62 micron",
                    "form_label": "Pouch",
                    "stage": "Laminated",
                    "width_mm": 420,
                    "thickness_micron": 62,
                    "print_status": "Printed",
                    "lamination_status": "Laminated",
                    "stock_strategy_label": "Intermediate pool",
                    "roll_count": 1,
                    "available_kg": 12.5,
                    "reserved_kg": 0,
                    "blocked_kg": 0,
                    "plant_location_summary": "Plant A / LAM Store",
                    "oldest_age_days": 8,
                }
            ],
            bulk_rows=[
                {
                    "material": "LDPE Granule",
                    "category": "GRANULE",
                    "plant": "Plant A",
                    "location": "RM Store",
                    "qty": 100,
                    "uom": "KG",
                    "avg_cost": 90,
                    "value": 9000,
                }
            ],
            packaging_rows=[],
            exception_rows=[],
        )

        self.assertIsNotNone(attachment)
        workbook = load_workbook(BytesIO(attachment.content), data_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Summary", "Family Summary", "Variant Summary", "Roll Detail", "Bulk RM", "Exceptions"],
        )
        summary_headers = [cell.value for cell in workbook["Summary"][1]]
        self.assertEqual(summary_headers, ["Plant", "Bulk KG", "Roll KG", "FG KG", "WIP KG", "Reserved KG"])
        family_headers = [cell.value for cell in workbook["Family Summary"][1]]
        self.assertIn("Business Family", family_headers)
        self.assertEqual(workbook["Family Summary"]["A2"].value, "PET Printed Laminate")
        variant_headers = [cell.value for cell in workbook["Variant Summary"][1]]
        self.assertIn("Strategy", variant_headers)
        self.assertEqual(workbook["Variant Summary"]["A2"].value, "PET Printed Laminate")
        self.assertEqual(workbook["Roll Detail"]["A2"].value, "ROLL-001")
        self.assertEqual(workbook["Bulk RM"]["A2"].value, "LDPE Granule")
